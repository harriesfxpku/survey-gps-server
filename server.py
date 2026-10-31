from datetime import datetime, timezone, timedelta
import os
import json
import secrets
import hashlib
import re
import unicodedata
from fastapi import UploadFile, File, Form

import psycopg
import uvicorn

import uuid
import urllib.request
import urllib.error
from urllib.parse import quote
from fastapi import Request, Header

from psycopg.rows import dict_row
from psycopg.types.json import Jsonb
from fastapi import FastAPI, HTTPException, Depends
from fastapi.security import APIKeyHeader, HTTPBasic, HTTPBasicCredentials
from fastapi import Security
from fastapi.responses import HTMLResponse, RedirectResponse, Response
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field
from html import escape
from fastapi.responses import FileResponse
from pathlib import Path

# =========================================================
# KONFIGURASI
# =========================================================

DATABASE_URL = os.environ.get("DATABASE_URL", "").strip()

# Normalize common provider URL schemes for psycopg 3.
if DATABASE_URL.startswith("//"):
    DATABASE_URL = "postgresql:" + DATABASE_URL
elif DATABASE_URL.startswith("postgres://"):
    DATABASE_URL = "postgresql://" + DATABASE_URL[len("postgres://"):]
elif DATABASE_URL.startswith("postgresql+psycopg://"):
    DATABASE_URL = "postgresql://" + DATABASE_URL[len("postgresql+psycopg://"):]

API_TOKEN = os.environ.get("API_TOKEN", "")

api_key_header = APIKeyHeader(
    name="X-API-Key",
    auto_error=False,
)

app = FastAPI(
    title="Survei GPS API",
    description="Server pusat aplikasi Survei GPS. Pengelolaan akun menggunakan istilah DPS (Kecamatan) dan DPSS (Desa/Kelurahan).",
    version="3.11.0-member-id-directory",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


# Tema global Theme 2 untuk seluruh halaman HTML, termasuk halaman submenu.
# CSS disisipkan setelah CSS halaman masing-masing agar tampilan konsisten tanpa
# mengubah endpoint, query database, formulir, maupun alur kerja yang sudah ada.
_THEME2_GLOBAL_CSS = r"""
<style id="theme2-global">
:root{color-scheme:dark;--t2-bg:#100709;--t2-panel:#211316;--t2-panel2:#2b191d;--t2-maroon:#650b18;--t2-gold:#e0ad35;--t2-text:#f7eee5;--t2-muted:#b8a8a0;--t2-line:#5b3939}
html,body{background:radial-gradient(ellipse at top right,#35151b 0,#16090c 56%,#100709 100%) !important;color:var(--t2-text) !important;font-family:Inter,system-ui,Arial,sans-serif !important}
body{min-height:100vh}
header,.hero,.modalhead{background:linear-gradient(110deg,#19080c,#500813 55%,#26090e) !important;color:#fff !important;border-bottom:2px solid var(--t2-gold) !important}
header h1,header h2,h1,h2,h3,h4,.section-title{color:#f1cb72 !important}
main,.wrap,.container,.content{color:var(--t2-text)}
.panel,.card,.metric,.stat,.member-card,.report-card,.box,form:not(.inline-form),.notice,.empty,.modalbody,.modalfoot{background:linear-gradient(145deg,#26191b,#1c1114) !important;color:var(--t2-text) !important;border-color:var(--t2-line) !important;box-shadow:0 8px 20px #0003}
label,b,strong,th{color:#e8cc8b}
p,td,li,span,small,.muted,.subtle,.count{color:inherit}
a{color:#f0c85e}
a.btn,button,.btn,input[type=submit]{background:linear-gradient(110deg,#e9bd51,#b87a16) !important;color:#24130a !important;border:1px solid #d8a83d !important;border-radius:9px !important;font-weight:750 !important;text-decoration:none !important}
a.btn.light,.btn.light,button.secondary,.btn.secondary{background:#321b20 !important;color:#f3dfb1 !important;border:1px solid #644349 !important}
button.danger,.danger{background:#8f202b !important;color:#fff !important;border-color:#b3444c !important}
input:not([type=checkbox]):not([type=radio]),select,textarea{background:#180e11 !important;color:#f8eee3 !important;border:1px solid #624449 !important;border-radius:8px !important}
input::placeholder,textarea::placeholder{color:#a99790 !important}
table{border-color:#493235 !important}th{background:#301b20 !important;color:#e7c46e !important}td{border-color:#493235 !important;color:#eee3d9 !important}
tr:hover td{background:#ffffff08 !important}
hr{border-color:#503337 !important}
footer,.foot{color:#a99b93 !important}
@media(max-width:700px){header{padding-left:14px !important;padding-right:14px !important}}
</style>
"""

@app.middleware("http")
async def apply_theme2_to_all_html(request: Request, call_next):
    response = await call_next(request)
    content_type = response.headers.get("content-type", "")
    if "text/html" not in content_type.lower():
        return response
    body_parts = []
    async for part in response.body_iterator:
        body_parts.append(part)
    body = b"".join(body_parts)
    try:
        html = body.decode("utf-8")
        if request.url.path.rstrip("/") != "/" and "id=\"theme2-global\"" not in html:
            if "</head>" in html:
                html = html.replace("</head>", _THEME2_GLOBAL_CSS + "</head>", 1)
            else:
                html = _THEME2_GLOBAL_CSS + html
        # Perbaikan khusus halaman Data Survei Admin (/admin):
        # latar sel dibuat terang dan teks marun agar terbaca jelas.
        if request.url.path.rstrip("/") == "/admin":
            admin_survey_css = r"""<style id="admin-survey-table-contrast">
.admin-survey-table td, .tablewrap td, table tbody td {
    color: #5b1018 !important;
    background-color: #fff8ef !important;
}
.admin-survey-table tbody tr:hover td, .tablewrap tbody tr:hover td, table tbody tr:hover td {
    background-color: #f3e5d2 !important;
}
.admin-survey-table td a, .tablewrap td a, table tbody td a {
    color: #5b1018 !important;
    font-weight: 700 !important;
}
</style>"""
            if "id=\"admin-survey-table-contrast\"" not in html:
                html = html.replace("</head>", admin_survey_css + "</head>", 1) if "</head>" in html else admin_survey_css + html
        headers = dict(response.headers)
        headers.pop("content-length", None)
        headers.pop("content-encoding", None)
        return Response(content=html.encode("utf-8"), status_code=response.status_code,
                        headers=headers, media_type="text/html")
    except UnicodeDecodeError:
        return response

@app.get("/logo.png")
def get_logo():
    # Logo utama berada di static/logo.png pada repository.
    base_dir = Path(__file__).parent
    candidates = [base_dir / "static" / "logo.png", base_dir / "logo.png"]
    logo_path = next((candidate for candidate in candidates if candidate.is_file()), None)
    if logo_path is None:
        raise HTTPException(status_code=404, detail="Logo tidak ditemukan di static/logo.png")
    return FileResponse(logo_path, media_type="image/png", headers={"Cache-Control": "public, max-age=3600"})

# =========================================================
# KONEKSI DATABASE
# =========================================================

def get_db():
    if not DATABASE_URL:
        raise RuntimeError(
            "DATABASE_URL belum dikonfigurasi."
        )

    return psycopg.connect(
        DATABASE_URL,
        row_factory=dict_row,
        connect_timeout=15,
    )


# =========================================================
# ID ANGGOTA OTOMATIS BERDASARKAN WILAYAH
# =========================================================

def _member_short_code(value: str, aliases: dict[str, str] | None = None) -> str:
    """Bentuk kode wilayah pendek, tanpa spasi atau tanda baca."""
    normalized = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode("ascii")
    words = re.findall(r"[A-Za-z0-9]+", normalized.upper())
    # Hilangkan awalan administratif agar kode tidak ikut memuat KEC/KAB/DESA.
    words = [word for word in words if word not in {"KEC", "KECAMATAN", "KAB", "KABUPATEN", "DESA", "KEL", "KELURAHAN"}]
    if not words:
        return "WLY"
    key = " ".join(words)
    aliases = aliases or {}
    if key in aliases:
        return aliases[key]
    # Nama dua kata atau lebih memakai inisial; nama satu kata memakai tiga huruf awal.
    if len(words) > 1:
        return "".join(word[0] for word in words)[:4]
    return words[0][:3]


def _member_code_prefix(org_level: str, region: str, village_name: str = "") -> str:
    """Prefix ringkas tanpa pemisah, misalnya DPSTBG001 dan DPSS TBG AST 001 (tanpa spasi pada ID)."""
    region_code = _member_short_code(region, {
        "TAMBANG": "TBG",
    })
    if str(org_level).upper() == "DPSS":
        village_code = _member_short_code(village_name, {
            "AUR SATI": "AST",
            "BALAM JAYA": "BLJ",
        })
        return f"DPSS{region_code}{village_code}"
    return f"DPS{region_code}"


def _allocate_member_code(connection, org_level: str, region: str, village_name: str = "") -> str:
    """Nomor urut tiga digit per wilayah; advisory lock mencegah nomor ganda saat pendaftaran bersamaan."""
    prefix = _member_code_prefix(org_level, region, village_name)
    connection.execute("SELECT pg_advisory_xact_lock(hashtext(%s))", (prefix,))
    existing = connection.execute(
        "SELECT member_code FROM organization_member_registrations WHERE member_code LIKE %s",
        (prefix + "%",),
    ).fetchall()
    used = []
    pattern = re.compile(r"^" + re.escape(prefix) + r"(\d+)$", re.IGNORECASE)
    for row in existing:
        match = pattern.fullmatch(str(row.get("member_code") or "").strip())
        if match:
            used.append(int(match.group(1)))
    number = max(used, default=0) + 1
    candidate = f"{prefix}{number:03d}"
    while connection.execute("""SELECT 1 FROM officers WHERE member_id=%s
            UNION ALL SELECT 1 FROM members WHERE member_id=%s LIMIT 1""", (candidate, candidate)).fetchone():
        number += 1
        candidate = f"{prefix}{number:03d}"
    return candidate


# =========================================================
# MEMBUAT TABEL DAN MEMASTIKAN ANGGOTA TES AKTIF
# =========================================================

def initialize_database():
    with get_db() as connection:

        connection.execute("""
            CREATE TABLE IF NOT EXISTS members (
                id BIGSERIAL PRIMARY KEY,
                member_id VARCHAR(50) UNIQUE NOT NULL,
                member_name VARCHAR(100) NOT NULL,
                active INTEGER NOT NULL DEFAULT 1,
                created_at TEXT NOT NULL
            )
        """)

        connection.execute("""
            CREATE TABLE IF NOT EXISTS surveys (
                id BIGSERIAL PRIMARY KEY,
                member_id VARCHAR(50) NOT NULL,
                latitude DOUBLE PRECISION NOT NULL,
                longitude DOUBLE PRECISION NOT NULL,
                accuracy DOUBLE PRECISION,
                notes TEXT NOT NULL DEFAULT '',
                surveyed_at TEXT NOT NULL,
                received_at TEXT NOT NULL,
                applicant_name TEXT NOT NULL DEFAULT '',
                applicant_nik TEXT NOT NULL DEFAULT '',
                applicant_phone TEXT NOT NULL DEFAULT '',
                applicant_address TEXT NOT NULL DEFAULT '',
                village TEXT NOT NULL DEFAULT '',
                business_name TEXT NOT NULL DEFAULT '',
                business_type TEXT NOT NULL DEFAULT '',
                business_since TEXT NOT NULL DEFAULT '',
                business_location TEXT NOT NULL DEFAULT '',
                ownership_status TEXT NOT NULL DEFAULT '',
                nib TEXT NOT NULL DEFAULT '',
                capital_amount DOUBLE PRECISION,
                turnover_amount DOUBLE PRECISION,
                net_profit_amount DOUBLE PRECISION,
                worker_count INTEGER,
                business_active TEXT NOT NULL DEFAULT '',
                business_issues TEXT NOT NULL DEFAULT '',
                assistance_types TEXT NOT NULL DEFAULT '',
                assistance_amount DOUBLE PRECISION,
                assistance_reason TEXT NOT NULL DEFAULT '',
                business_verified TEXT NOT NULL DEFAULT '',
                verification_status TEXT NOT NULL DEFAULT '',
                verification_notes TEXT NOT NULL DEFAULT ''
            )
        """)

        # Migrasi aman untuk database yang sudah berisi data survei lama.
        migration_columns = {
            "applicant_name": "TEXT NOT NULL DEFAULT ''",
            "applicant_nik": "TEXT NOT NULL DEFAULT ''",
            "applicant_phone": "TEXT NOT NULL DEFAULT ''",
            "applicant_address": "TEXT NOT NULL DEFAULT ''",
            "village": "TEXT NOT NULL DEFAULT ''",
            "business_name": "TEXT NOT NULL DEFAULT ''",
            "business_type": "TEXT NOT NULL DEFAULT ''",
            "business_since": "TEXT NOT NULL DEFAULT ''",
            "business_location": "TEXT NOT NULL DEFAULT ''",
            "ownership_status": "TEXT NOT NULL DEFAULT ''",
            "nib": "TEXT NOT NULL DEFAULT ''",
            "capital_amount": "DOUBLE PRECISION",
            "turnover_amount": "DOUBLE PRECISION",
            "net_profit_amount": "DOUBLE PRECISION",
            "worker_count": "INTEGER",
            "business_active": "TEXT NOT NULL DEFAULT ''",
            "business_issues": "TEXT NOT NULL DEFAULT ''",
            "assistance_types": "TEXT NOT NULL DEFAULT ''",
            "assistance_amount": "DOUBLE PRECISION",
            "assistance_reason": "TEXT NOT NULL DEFAULT ''",
            "business_verified": "TEXT NOT NULL DEFAULT ''",
            "verification_status": "TEXT NOT NULL DEFAULT ''",
            "verification_notes": "TEXT NOT NULL DEFAULT ''",
        }
        for column_name, column_type in migration_columns.items():
            connection.execute(
                f"ALTER TABLE surveys ADD COLUMN IF NOT EXISTS {column_name} {column_type}"
            )

        connection.execute("""
            CREATE TABLE IF NOT EXISTS survey_media (
                id BIGSERIAL PRIMARY KEY,
                survey_id BIGINT NOT NULL REFERENCES surveys(id) ON DELETE CASCADE,
                member_id VARCHAR(50) NOT NULL,
                object_path TEXT NOT NULL UNIQUE,
                content_type VARCHAR(100) NOT NULL,
                original_filename TEXT,
                file_size BIGINT NOT NULL,
                uploaded_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                media_category VARCHAR(40) NOT NULL DEFAULT 'business_photo'
            )
        """)
        connection.execute(
            "ALTER TABLE survey_media ADD COLUMN IF NOT EXISTS media_category VARCHAR(40) NOT NULL DEFAULT 'business_photo'"
        )
        connection.execute(
            "ALTER TABLE survey_media ADD COLUMN IF NOT EXISTS storage_bucket VARCHAR(100) NOT NULL DEFAULT 'survey-media'"
        )

        # Akun petugas dan sesi login ditambahkan tanpa mengubah tabel survei/media lama.
        connection.execute("""
            CREATE TABLE IF NOT EXISTS officers (
                id BIGSERIAL PRIMARY KEY,
                member_id VARCHAR(50) UNIQUE NOT NULL REFERENCES members(member_id),
                member_name VARCHAR(100) NOT NULL,
                password_salt VARCHAR(64) NOT NULL,
                password_hash VARCHAR(128) NOT NULL,
                role VARCHAR(10) NOT NULL CHECK (role IN ('DPSS', 'DPS')),
                region VARCHAR(150) NOT NULL DEFAULT '',
                active INTEGER NOT NULL DEFAULT 1,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
        """)
        connection.execute("""
            CREATE TABLE IF NOT EXISTS officer_sessions (
                id BIGSERIAL PRIMARY KEY,
                officer_id BIGINT NOT NULL REFERENCES officers(id) ON DELETE CASCADE,
                token_hash VARCHAR(64) UNIQUE NOT NULL,
                expires_at TIMESTAMPTZ NOT NULL,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
        """)
        # Migrate officer columns used by login/admin on older databases.
        connection.execute("ALTER TABLE officers ADD COLUMN IF NOT EXISTS password_salt VARCHAR(64)")
        connection.execute("ALTER TABLE officers ADD COLUMN IF NOT EXISTS password_hash VARCHAR(128)")
        connection.execute("ALTER TABLE officers ADD COLUMN IF NOT EXISTS role VARCHAR(10) NOT NULL DEFAULT 'DPSS'")
        connection.execute("ALTER TABLE officers ADD COLUMN IF NOT EXISTS region VARCHAR(150) NOT NULL DEFAULT ''")
        connection.execute("ALTER TABLE officers ADD COLUMN IF NOT EXISTS active INTEGER NOT NULL DEFAULT 1")
        connection.execute("ALTER TABLE officers ADD COLUMN IF NOT EXISTS updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()")
        # Data verifikasi dipisahkan dari survei awal agar data DPSS tidak tertimpa.
        connection.execute("""
            CREATE TABLE IF NOT EXISTS survey_verifications (
                id BIGSERIAL PRIMARY KEY,
                survey_id BIGINT NOT NULL UNIQUE REFERENCES surveys(id) ON DELETE CASCADE,
                verifier_id VARCHAR(50) NOT NULL,
                verifier_name VARCHAR(100) NOT NULL,
                verifier_region VARCHAR(150) NOT NULL DEFAULT '',
                decision VARCHAR(40) NOT NULL CHECK (decision IN ('Sudah diverifikasi', 'Ditolak')),
                business_condition TEXT NOT NULL DEFAULT '',
                actual_business_activity TEXT NOT NULL DEFAULT '',
                existing_assets TEXT NOT NULL DEFAULT '',
                needs JSONB NOT NULL DEFAULT '[]'::jsonb,
                working_capital_details TEXT NOT NULL DEFAULT '',
                verified_turnover DOUBLE PRECISION,
                verified_profit DOUBLE PRECISION,
                verified_worker_count INTEGER,
                data_match VARCHAR(40) NOT NULL DEFAULT '',
                discrepancy_notes TEXT NOT NULL DEFAULT '',
                verification_notes TEXT NOT NULL DEFAULT '',
                rejection_reason TEXT NOT NULL DEFAULT '',
                latitude DOUBLE PRECISION NOT NULL,
                longitude DOUBLE PRECISION NOT NULL,
                accuracy DOUBLE PRECISION,
                verified_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                total_needs_amount DOUBLE PRECISION NOT NULL DEFAULT 0
            )
        """)
        # Ensure verification columns exist when upgrading an older database.
        verification_migrations = {
            "verifier_region": "VARCHAR(150) NOT NULL DEFAULT ''",
            "business_condition": "TEXT NOT NULL DEFAULT ''",
            "actual_business_activity": "TEXT NOT NULL DEFAULT ''",
            "existing_assets": "TEXT NOT NULL DEFAULT ''",
            "needs": "JSONB NOT NULL DEFAULT '[]'::jsonb",
            "working_capital_details": "TEXT NOT NULL DEFAULT ''",
            "verified_turnover": "DOUBLE PRECISION",
            "verified_profit": "DOUBLE PRECISION",
            "verified_worker_count": "INTEGER",
            "data_match": "VARCHAR(40) NOT NULL DEFAULT ''",
            "discrepancy_notes": "TEXT NOT NULL DEFAULT ''",
            "verification_notes": "TEXT NOT NULL DEFAULT ''",
            "rejection_reason": "TEXT NOT NULL DEFAULT ''",
            "latitude": "DOUBLE PRECISION",
            "longitude": "DOUBLE PRECISION",
            "accuracy": "DOUBLE PRECISION",
            "verified_at": "TIMESTAMPTZ NOT NULL DEFAULT NOW()",
            "total_needs_amount": "DOUBLE PRECISION NOT NULL DEFAULT 0",
        }
        for column_name, column_type in verification_migrations.items():
            connection.execute(
                f"ALTER TABLE survey_verifications ADD COLUMN IF NOT EXISTS {column_name} {column_type}"
            )

        connection.execute("""
            CREATE TABLE IF NOT EXISTS verification_media (
                id BIGSERIAL PRIMARY KEY,
                survey_id BIGINT NOT NULL REFERENCES surveys(id) ON DELETE CASCADE,
                verifier_id VARCHAR(50) NOT NULL,
                object_path TEXT NOT NULL UNIQUE,
                content_type VARCHAR(100) NOT NULL,
                original_filename TEXT,
                file_size BIGINT NOT NULL,
                uploaded_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
        """)
        # Konfigurasi desa dan penugasan petugas, ditambahkan tanpa menghapus data lama.
        connection.execute("""CREATE TABLE IF NOT EXISTS villages (
            id BIGSERIAL PRIMARY KEY, village_name VARCHAR(150) UNIQUE NOT NULL,
            district VARCHAR(150) NOT NULL DEFAULT 'Tambang', bucket_name VARCHAR(100) NOT NULL UNIQUE,
            active INTEGER NOT NULL DEFAULT 1, created_at TIMESTAMPTZ NOT NULL DEFAULT NOW())""")
        connection.execute("""CREATE TABLE IF NOT EXISTS officer_villages (
            officer_id BIGINT NOT NULL REFERENCES officers(id) ON DELETE CASCADE,
            village_id BIGINT NOT NULL REFERENCES villages(id) ON DELETE CASCADE,
            created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(), PRIMARY KEY (officer_id, village_id))""")
        # Profil personel terpisah dari kredensial akun. Dokumen disimpan di bucket privat.
        connection.execute("""
            CREATE TABLE IF NOT EXISTS officer_profiles (
                officer_id BIGINT PRIMARY KEY REFERENCES officers(id) ON DELETE CASCADE,
                nik VARCHAR(30) NOT NULL DEFAULT '',
                phone VARCHAR(30) NOT NULL DEFAULT '',
                address TEXT NOT NULL DEFAULT '',
                birth_place VARCHAR(100) NOT NULL DEFAULT '',
                birth_date DATE,
                profile_photo_path TEXT NOT NULL DEFAULT '',
                ktp_path TEXT NOT NULL DEFAULT '',
                updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
        """)
        # Pendaftaran mandiri anggota organisasi: terpisah dari akun login petugas.
        connection.execute("""
            CREATE TABLE IF NOT EXISTS organization_member_registrations (
                id BIGSERIAL PRIMARY KEY,
                registration_code VARCHAR(40) UNIQUE NOT NULL,
                member_code VARCHAR(80) NOT NULL DEFAULT '',
                full_name VARCHAR(150) NOT NULL,
                nik VARCHAR(30) NOT NULL,
                birth_place VARCHAR(100) NOT NULL DEFAULT '',
                birth_date DATE,
                gender VARCHAR(20) NOT NULL DEFAULT '',
                address TEXT NOT NULL,
                phone VARCHAR(30) NOT NULL,
                email VARCHAR(150) NOT NULL DEFAULT '',
                org_level VARCHAR(10) NOT NULL CHECK (org_level IN ('DPS','DPSS')),
                org_position VARCHAR(150) NOT NULL,
                region VARCHAR(150) NOT NULL DEFAULT '',
                village_id BIGINT REFERENCES villages(id),
                profile_photo_path TEXT NOT NULL DEFAULT '',
                ktp_path TEXT NOT NULL DEFAULT '',
                status VARCHAR(20) NOT NULL DEFAULT 'pending'
                    CHECK (status IN ('pending','revision','approved','rejected')),
                admin_note TEXT NOT NULL DEFAULT '',
                linked_member_id VARCHAR(50) NOT NULL DEFAULT '',
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                reviewed_at TIMESTAMPTZ
            )
        """)
        connection.execute("CREATE INDEX IF NOT EXISTS idx_org_reg_status_created ON organization_member_registrations(status, created_at DESC)")
        connection.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_org_reg_nik_unique ON organization_member_registrations(nik)")
        connection.execute("ALTER TABLE surveys ADD COLUMN IF NOT EXISTS village_id BIGINT REFERENCES villages(id)")
        connection.execute("ALTER TABLE survey_media ADD COLUMN IF NOT EXISTS storage_bucket VARCHAR(100) NOT NULL DEFAULT 'survey-media'")
        connection.execute("ALTER TABLE verification_media ADD COLUMN IF NOT EXISTS storage_bucket VARCHAR(100) NOT NULL DEFAULT 'survey-media'")
        connection.execute("""INSERT INTO villages (village_name, district, bucket_name) VALUES
            ('Aur Sati', 'Tambang', 'Aur-Sati-Kec-Tambang'),
            ('Balam Jaya', 'Tambang', 'Balam-Jaya-Kec-Tambang')
            ON CONFLICT (village_name)
            DO UPDATE SET district = EXCLUDED.district""")
        connection.execute("ALTER TABLE organization_member_registrations ADD COLUMN IF NOT EXISTS member_code VARCHAR(80) NOT NULL DEFAULT ''")
        connection.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_org_reg_member_code_unique ON organization_member_registrations(member_code) WHERE member_code <> ''")
        # Backfill ID untuk pendaftaran lama. ID login yang sudah dipakai dipertahankan agar akun aktif tidak berubah.
        old_members = connection.execute("""SELECT r.id,r.member_code,r.linked_member_id,r.org_level,r.region,
            COALESCE(v.village_name,'') AS village_name
            FROM organization_member_registrations r LEFT JOIN villages v ON v.id=r.village_id
            WHERE COALESCE(r.member_code,'')='' ORDER BY r.created_at,r.id""").fetchall()
        for old in old_members:
            old_code = (old.get('linked_member_id') or '').strip().upper()
            if not old_code:
                old_code = _allocate_member_code(connection, old['org_level'], old['region'], old['village_name'])
            connection.execute("UPDATE organization_member_registrations SET member_code=%s WHERE id=%s", (old_code, old['id']))
        # Hanya tautkan survei lama jika nama desa cocok persis; media lama tetap di bucket survey-media.
        connection.execute("""UPDATE surveys s SET village_id = v.id FROM villages v
            WHERE s.village_id IS NULL AND LOWER(TRIM(s.village)) = LOWER(v.village_name)""")

        # Tidak membuat atau mengaktifkan ulang akun anggota secara otomatis saat startup.


def get_village_bucket(connection, village_id):
    if village_id is None:
        return os.environ.get("SUPABASE_DEFAULT_BUCKET", "survey-media")
    row = connection.execute("SELECT bucket_name FROM villages WHERE id = %s AND active = 1", (village_id,)).fetchone()
    if row is None:
        raise HTTPException(status_code=400, detail="Desa survei tidak dikenal atau tidak aktif.")
    return row["bucket_name"]


def ensure_dps_survey_access(connection, identity, survey_id):
    # DPS (Kecamatan) dapat membuka survei dari desa di kecamatannya.
    # Penugasan desa eksplisit tetap didukung untuk menjaga kompatibilitas.
    row = connection.execute("""SELECT s.id
        FROM surveys s
        JOIN officers o ON o.member_id = %s AND o.active = 1
        WHERE s.id = %s
          AND (
            EXISTS (
                SELECT 1 FROM officer_villages ov
                WHERE ov.officer_id = o.id AND ov.village_id = s.village_id
            )
            OR EXISTS (
                SELECT 1 FROM villages v
                WHERE v.active = 1
                  AND lower(trim(regexp_replace(coalesce(v.district, ''), '^(kecamatan|kec\\.?)[[:space:]]+', '', 'i'))) =
                      lower(trim(regexp_replace(regexp_replace(coalesce(o.region, ''), '^(kecamatan|kec\\.?)[[:space:]]+', '', 'i'), '[,;].*$', '')))
                  AND (
                      v.id = s.village_id
                      OR lower(trim(regexp_replace(coalesce(v.village_name, ''), '^(desa|kelurahan)[[:space:]]+', '', 'i'))) =
                         lower(trim(regexp_replace(coalesce(s.village, ''), '^(desa|kelurahan)[[:space:]]+', '', 'i')))
                  )
            )
          )""", (identity["member_id"], survey_id)).fetchone()
    if row is None:
        raise HTTPException(status_code=403, detail="Survei ini berada di luar wilayah kecamatan akun DPS.")


@app.on_event("startup")
def startup_event():
    initialize_database()


# =========================================================
# AUTENTIKASI
# =========================================================

def verify_api_key(
    api_key: str = Depends(api_key_header),
):
    if (
        not API_TOKEN
        or not api_key
        or not secrets.compare_digest(api_key, API_TOKEN)
    ):
        raise HTTPException(
            status_code=401,
            detail="Kunci akses tidak valid",
        )

    return True


# =========================================================
# AUTENTIKASI PETUGAS
# =========================================================

def _password_hash(password: str, salt_hex: str) -> str:
    return hashlib.pbkdf2_hmac(
        "sha256", password.encode("utf-8"), bytes.fromhex(salt_hex), 310_000
    ).hex()


def _token_hash(token: str) -> str:
    return hashlib.sha256(token.encode("utf-8")).hexdigest()


def get_request_identity(
    authorization: str | None = Header(default=None),
    api_key: str | None = Depends(api_key_header),
):
    # API key lama tetap diterima untuk kompatibilitas dashboard/versi aplikasi lama.
    if API_TOKEN and api_key and secrets.compare_digest(api_key, API_TOKEN):
        return {"auth_type": "api_key", "role": "LEGACY_ADMIN"}

    if authorization and authorization.lower().startswith("bearer "):
        token = authorization[7:].strip()
        if token:
            with get_db() as connection:
                row = connection.execute(
                    """
                    SELECT o.member_id, o.member_name, o.role, o.region, o.active,
                           s.expires_at
                    FROM officer_sessions s
                    JOIN officers o ON o.id = s.officer_id
                    WHERE s.token_hash = %s AND s.expires_at > NOW()
                    """,
                    (_token_hash(token),),
                ).fetchone()
            if row and row["active"] == 1:
                return {
                    "auth_type": "officer",
                    "member_id": row["member_id"],
                    "member_name": row["member_name"],
                    "role": row["role"],
                    "region": row["region"],
                }

    raise HTTPException(status_code=401, detail="Sesi tidak valid atau sudah berakhir. Silakan login kembali.")


def get_officer_identity(
    authorization: str | None = Header(default=None),
):
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(status_code=401, detail="Silakan login menggunakan akun petugas.")
    token = authorization[7:].strip()
    if not token:
        raise HTTPException(status_code=401, detail="Token login tidak valid.")
    with get_db() as connection:
        row = connection.execute(
            """
            SELECT o.member_id, o.member_name, o.role, o.region, o.active, s.expires_at
            FROM officer_sessions s
            JOIN officers o ON o.id = s.officer_id
            WHERE s.token_hash = %s AND s.expires_at > NOW()
            """,
            (_token_hash(token),),
        ).fetchone()
    if not row or row["active"] != 1:
        raise HTTPException(status_code=401, detail="Sesi tidak valid atau sudah berakhir. Silakan login kembali.")
    return {
        "auth_type": "officer",
        "member_id": row["member_id"],
        "member_name": row["member_name"],
        "role": row["role"],
        "region": row["region"],
    }


class OfficerLogin(BaseModel):
    member_id: str = Field(min_length=1, max_length=50)
    password: str = Field(min_length=1, max_length=200)


class OfficerCreate(BaseModel):
    member_id: str = Field(min_length=2, max_length=50)
    member_name: str = Field(min_length=2, max_length=100)
    password: str = Field(min_length=10, max_length=200)
    role: str
    region: str = Field(min_length=1, max_length=150)


@app.post("/auth/login")
def officer_login(payload: OfficerLogin):
    member_id = payload.member_id.strip()
    with get_db() as connection:
        officer = connection.execute(
            "SELECT * FROM officers WHERE member_id = %s", (member_id,)
        ).fetchone()
        if (
            not officer
            or officer["active"] != 1
            or not officer.get("password_salt")
            or not officer.get("password_hash")
            or not secrets.compare_digest(
                _password_hash(payload.password, officer["password_salt"]),
                officer["password_hash"],
            )
        ):
            raise HTTPException(status_code=401, detail="ID petugas atau kata sandi salah, atau akun tidak aktif.")
        token = secrets.token_urlsafe(48)
        expires_at = datetime.now(timezone.utc) + timedelta(hours=12)
        connection.execute(
            "DELETE FROM officer_sessions WHERE expires_at <= NOW()"
        )
        connection.execute(
            "INSERT INTO officer_sessions (officer_id, token_hash, expires_at) VALUES (%s, %s, %s)",
            (officer["id"], _token_hash(token), expires_at),
        )
    return {
        "access_token": token,
        "token_type": "Bearer",
        "expires_at": expires_at.isoformat(),
        "member_id": officer["member_id"],
        "member_name": officer["member_name"],
        "role": officer["role"],
        "region": officer["region"],
    }


@app.get("/auth/me")
def officer_me(identity: dict = Depends(get_officer_identity)):
    return {key: identity[key] for key in ("member_id", "member_name", "role", "region")}


@app.post("/auth/logout")
def officer_logout(authorization: str | None = Header(default=None)):
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(status_code=401, detail="Sesi login tidak ditemukan.")
    token = authorization[7:].strip()
    with get_db() as connection:
        connection.execute("DELETE FROM officer_sessions WHERE token_hash = %s", (_token_hash(token),))
    return {"message": "Berhasil keluar."}


@app.post("/admin/officers", include_in_schema=False)
@app.post("/admin/akun-petugas", tags=["Pengelolaan Akun"], summary="Pembuatan akun melalui anggota terverifikasi")
def create_officer(
    payload: OfficerCreate,
    _: bool = Depends(verify_api_key),
):
    raise HTTPException(
        status_code=403,
        detail="Akun aplikasi hanya dapat diberikan kepada anggota yang sudah disetujui melalui Kelola Akses Aplikasi. Pembuatan akun manual dinonaktifkan.",
    )

@app.get("/admin/officers", include_in_schema=False)
@app.get("/admin/akun-petugas", tags=["Pengelolaan Akun"], summary="Lihat daftar akun DPS dan DPSS")
def list_officers(_: bool = Depends(verify_api_key)):
    with get_db() as connection:
        rows = connection.execute(
            """SELECT o.member_id, o.member_name, o.role, o.region, o.active,
                      o.created_at, o.updated_at,
                      ARRAY(SELECT v.village_name
                            FROM officer_villages ov
                            JOIN villages v ON v.id = ov.village_id
                            WHERE ov.officer_id = o.id
                            ORDER BY v.village_name) AS villages
               FROM officers o ORDER BY o.role, o.region, o.member_id"""
        ).fetchall()
    return [dict(row) for row in rows]


class VillageCreate(BaseModel):
    village_name: str = Field(min_length=2, max_length=150)
    district: str = Field(default="Tambang", min_length=2, max_length=150)
    bucket_name: str | None = Field(default=None, max_length=100)


def _storage_slug(value: str) -> str:
    slug = re.sub(r"[^a-z0-9-]+", "-", value.strip().lower()).strip("-")
    return slug[:100] or "survey-media"


@app.get("/admin/villages", include_in_schema=False)
@app.get("/admin/data-desa", tags=["Data Wilayah"], summary="Lihat daftar desa dan kecamatan")
def list_villages(_: bool = Depends(verify_api_key)):
    with get_db() as connection:
        rows = connection.execute(
            "SELECT id, village_name, district, bucket_name, active, created_at FROM villages ORDER BY district, village_name"
        ).fetchall()
    return [dict(row) for row in rows]


@app.post("/admin/villages", include_in_schema=False)
@app.post("/admin/data-desa", tags=["Data Wilayah"], summary="Tambah data desa dan kecamatan")
def create_village(payload: VillageCreate, _: bool = Depends(verify_api_key)):
    village_name = payload.village_name.strip()
    district = payload.district.strip()
    bucket_name = (payload.bucket_name or _storage_slug(village_name)).strip()
    if not village_name or not district or not bucket_name:
        raise HTTPException(status_code=422, detail="Nama desa, kecamatan, dan nama bucket wajib diisi.")
    if not re.fullmatch(r"[a-z0-9][a-z0-9-]{1,99}", bucket_name):
        raise HTTPException(status_code=422, detail="Nama bucket hanya boleh berisi huruf kecil, angka, dan tanda hubung; panjang 2-100 karakter.")
    try:
        with get_db() as connection:
            row = connection.execute(
                """INSERT INTO villages (village_name, district, bucket_name)
                   VALUES (%s, %s, %s)
                   RETURNING id, village_name, district, bucket_name, active, created_at""",
                (village_name, district, bucket_name),
            ).fetchone()
    except psycopg.errors.UniqueViolation as exc:
        raise HTTPException(status_code=409, detail="Nama desa atau nama bucket sudah terdaftar.") from exc
    return {"message": "Desa berhasil didaftarkan. Pastikan bucket dengan nama tersebut sudah dibuat sebagai PRIVATE di Supabase Storage.", "village": dict(row)}


class OfficerVillageAssignment(BaseModel):
    villages: list[str] = Field(min_length=1, max_length=20)


@app.post("/admin/officers/{member_id}/villages", include_in_schema=False)
@app.post("/admin/akun-petugas/{member_id}/penugasan-desa", tags=["Pengelolaan Akun"], summary="Atur desa tugas akun petugas")
def assign_officer_villages(
    member_id: str,
    payload: OfficerVillageAssignment,
    _: bool = Depends(verify_api_key),
):
    """Ganti daftar desa tugas petugas. villages berisi nama desa yang terdaftar."""
    normalized_id = member_id.strip().upper()
    requested = list(dict.fromkeys(name.strip() for name in payload.villages if name.strip()))
    if not requested:
        raise HTTPException(status_code=422, detail="Isi minimal satu nama desa.")
    with get_db() as connection:
        officer = connection.execute("SELECT id, role FROM officers WHERE member_id=%s AND active=1", (normalized_id,)).fetchone()
        if not officer:
            raise HTTPException(status_code=404, detail="Akun petugas aktif tidak ditemukan.")
        villages = connection.execute("SELECT id, village_name FROM villages WHERE active=1 AND LOWER(village_name)=ANY(%s)", ([v.lower() for v in requested],)).fetchall()
        found = {v["village_name"].lower() for v in villages}
        missing = [v for v in requested if v.lower() not in found]
        if missing:
            raise HTTPException(status_code=400, detail="Desa tidak dikenal: " + ", ".join(missing))
        connection.execute("DELETE FROM officer_villages WHERE officer_id=%s", (officer["id"],))
        for village in villages:
            connection.execute("INSERT INTO officer_villages (officer_id, village_id) VALUES (%s, %s) ON CONFLICT DO NOTHING", (officer["id"], village["id"]))
    return {"message": "Penugasan desa berhasil diperbarui.", "member_id": normalized_id, "villages": [v["village_name"] for v in villages]}


@app.post("/admin/officers/{member_id}/deactivate", include_in_schema=False)
@app.post("/admin/akun-petugas/{member_id}/nonaktifkan", tags=["Pengelolaan Akun"], summary="Nonaktifkan akun petugas")
def deactivate_officer(member_id: str, _: bool = Depends(verify_api_key)):
    with get_db() as connection:
        result = connection.execute(
            "UPDATE officers SET active = 0, updated_at = NOW() WHERE member_id = %s",
            (member_id.strip().upper(),),
        )
        connection.execute(
            "DELETE FROM officer_sessions WHERE officer_id IN (SELECT id FROM officers WHERE member_id = %s)",
            (member_id.strip().upper(),),
        )
        if result.rowcount == 0:
            raise HTTPException(status_code=404, detail="Akun petugas tidak ditemukan.")
        connection.execute("UPDATE members SET active = 0 WHERE member_id = %s", (member_id.strip().upper(),))
    return {"message": "Akun petugas dinonaktifkan.", "member_id": member_id.strip().upper()}


@app.delete("/admin/akun-petugas/{member_id}", tags=["Pengelolaan Akun"], summary="Hapus akun petugas")
def delete_officer(member_id: str, _: bool = Depends(verify_api_key)):
    """Hapus akses akun petugas tanpa menghapus riwayat survei yang sudah dikirim."""
    normalized_id = member_id.strip().upper()
    with get_db() as connection:
        officer = connection.execute(
            "SELECT id FROM officers WHERE member_id = %s", (normalized_id,)
        ).fetchone()
        if officer is None:
            raise HTTPException(status_code=404, detail="Akun petugas tidak ditemukan.")
        connection.execute(
            "DELETE FROM officer_sessions WHERE officer_id = %s", (officer["id"],)
        )
        connection.execute(
            "DELETE FROM officers WHERE member_id = %s", (normalized_id,)
        )
        # Pertahankan baris anggota nonaktif untuk menjaga riwayat dan kompatibilitas data lama.
        connection.execute(
            "UPDATE members SET active = 0 WHERE member_id = %s", (normalized_id,)
        )
    return {"message": "Akun petugas berhasil dihapus. Riwayat survei tetap tersimpan.", "member_id": normalized_id}


# =========================================================
# MODEL DATA
# =========================================================

class MemberCreate(BaseModel):
    member_id: str = Field(
        min_length=1,
        max_length=50,
    )

    member_name: str = Field(
        min_length=1,
        max_length=100,
    )


class AssistanceNeedItem(BaseModel):
    item_name: str = Field(min_length=1, max_length=200)
    specification: str = Field(default="", max_length=1000)
    quantity: float = Field(gt=0, le=1000000)
    unit: str = Field(default="unit", max_length=40)
    unit_price: float = Field(ge=0, le=1000000000000)
    notes: str = Field(default="", max_length=1000)


class VerificationCreate(BaseModel):
    decision: str
    business_condition: str = Field(default="", max_length=3000)
    actual_business_activity: str = Field(default="", max_length=3000)
    existing_assets: str = Field(default="", max_length=3000)
    needs: list[AssistanceNeedItem] = Field(default_factory=list, max_length=50)
    working_capital_details: str = Field(default="", max_length=3000)
    verified_turnover: float | None = Field(default=None, ge=0)
    verified_profit: float | None = Field(default=None, ge=0)
    verified_worker_count: int | None = Field(default=None, ge=0, le=100000)
    data_match: str = Field(default="", max_length=40)
    discrepancy_notes: str = Field(default="", max_length=3000)
    verification_notes: str = Field(default="", max_length=3000)
    rejection_reason: str = Field(default="", max_length=3000)
    latitude: float = Field(ge=-90, le=90)
    longitude: float = Field(ge=-180, le=180)
    accuracy: float | None = Field(default=None, ge=0)


class SurveyCreate(BaseModel):
    member_id: str = Field(
        min_length=1,
        max_length=50,
    )

    latitude: float = Field(
        ge=-90,
        le=90,
    )

    longitude: float = Field(
        ge=-180,
        le=180,
    )

    accuracy: float | None = Field(
        default=None,
        ge=0,
    )

    notes: str = Field(
        default="",
        max_length=5000,
    )

    surveyed_at: datetime
    applicant_name: str = Field(default="", max_length=150)
    applicant_nik: str = Field(default="", max_length=30)
    applicant_phone: str = Field(default="", max_length=40)
    applicant_address: str = Field(default="", max_length=1000)
    village: str = Field(default="", max_length=120)
    business_name: str = Field(default="", max_length=200)
    business_type: str = Field(default="", max_length=150)
    business_since: str = Field(default="", max_length=30)
    business_location: str = Field(default="", max_length=1000)
    ownership_status: str = Field(default="", max_length=100)
    nib: str = Field(default="", max_length=100)
    capital_amount: float | None = Field(default=None, ge=0)
    turnover_amount: float | None = Field(default=None, ge=0)
    net_profit_amount: float | None = Field(default=None, ge=0)
    worker_count: int | None = Field(default=None, ge=0, le=100000)
    business_active: str = Field(default="", max_length=50)
    business_issues: str = Field(default="", max_length=3000)
    assistance_types: str = Field(default="", max_length=1000)
    assistance_amount: float | None = Field(default=None, ge=0)
    assistance_reason: str = Field(default="", max_length=3000)
    business_verified: str = Field(default="", max_length=50)
    verification_status: str = Field(default="", max_length=50)
    verification_notes: str = Field(default="", max_length=3000)


# =========================================================
# HALAMAN UTAMA
# =========================================================

@app.get("/", response_class=HTMLResponse)
def home():
    """Landing publik Theme 2; CTA tetap menuju formulir pendaftaran yang sudah ada."""
    html = """<!doctype html>
<html lang="id">
<head>
<meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="theme-color" content="#160609">
<title>555 | Pendaftaran Anggota</title>
<style>
:root{color-scheme:dark;--gold:#d8b15a;--gold2:#f6df9e}
*{box-sizing:border-box}body{margin:0;min-height:100vh;overflow-x:hidden;background:radial-gradient(ellipse at 50% 42%,#57131d 0,#260910 42%,#090609 100%);color:#fff5e9;font-family:Inter,system-ui,-apple-system,Segoe UI,Arial,sans-serif}
.top-glow{height:4px;background:linear-gradient(90deg,#6b0c18,#f4d47e,#9b1b25,#d8b15a)}
header{max-width:1240px;margin:auto;padding:20px 28px;display:flex;align-items:center}
.brand{display:flex;align-items:center;gap:13px;text-decoration:none}.brand img{width:62px;height:62px;object-fit:contain;border-radius:50%;border:1px solid #d8b15a;padding:3px;background:#24090e;box-shadow:0 0 25px #d8b15a22}.brand strong{display:block;color:var(--gold2);font-size:16px;letter-spacing:.12em}.brand small{color:#c9b5a8;font-size:10px;letter-spacing:.2em;text-transform:uppercase}
main{min-height:calc(100vh - 110px);display:flex;flex-direction:column;align-items:center;justify-content:center;padding:10px 20px 58px;text-align:center}.art{position:relative;width:min(90vw,570px);display:grid;place-items:center;margin:0 auto 18px}.halo{position:absolute;width:92%;aspect-ratio:1;border-radius:50%;background:radial-gradient(circle,#d8a34238 0,#8d1b2633 45%,transparent 72%);filter:blur(6px)}.ring{position:absolute;width:82%;aspect-ratio:1;border:1px solid #f2d17b66;border-radius:50%;box-shadow:0 0 0 18px #e5b9500b,0 0 70px #d8a34222,inset 0 0 55px #d8a34215}.ring:before,.ring:after{content:"";position:absolute;inset:9%;border:1px solid #e7c16c44;border-radius:50%}.ring:after{inset:20%;border-color:#e7c16c33}.garuda{position:relative;z-index:1;width:100%;height:auto;filter:drop-shadow(0 18px 24px #0009)}
.cta{position:relative;z-index:2;display:inline-flex;align-items:center;justify-content:center;gap:13px;min-height:58px;min-width:245px;padding:15px 30px;border:1px solid #ffe7a8;border-radius:12px;background:linear-gradient(110deg,#ffe69a,#d6a43c 62%,#f5d77f);color:#27130b;text-decoration:none;font-size:15px;font-weight:900;letter-spacing:.09em;box-shadow:0 10px 32px #08030799;transition:transform .18s,filter .18s}.cta:hover{transform:translateY(-2px);filter:brightness(1.08)}
@media(max-width:520px){header{padding:15px 17px}.brand img{width:49px;height:49px}.brand strong{font-size:13px}.brand small{font-size:8px}.art{width:min(96vw,390px);margin-bottom:8px}.cta{width:min(90vw,330px);min-height:54px;font-size:13px}}
@media(prefers-reduced-motion:reduce){.cta{transition:none}}
</style>
</head>
<body><div class="top-glow"></div>
<header><a class="brand" href="/" aria-label="Logo Ikatan Persaudaraan Triple 555"><img src="/logo.png" alt="Logo resmi 555"><span><strong>TRIPLE 555</strong><small>Ikatan Persaudaraan</small></span></a></header>
<main>
  <div class="art" role="img" aria-label="Ornamen Garuda berwarna emas dengan sayap terbentang">
    <div class="halo"></div><div class="ring"></div>
    <svg class="garuda" viewBox="0 0 520 390" xmlns="http://www.w3.org/2000/svg" aria-hidden="true">
      <defs><linearGradient id="gGold" x1="0" y1="0" x2="1" y2="1"><stop stop-color="#fff0b0"/><stop offset=".42" stop-color="#e5bd62"/><stop offset="1" stop-color="#9d651d"/></linearGradient><linearGradient id="gLight" x1="0" y1="0" x2="0" y2="1"><stop stop-color="#ffedac"/><stop offset="1" stop-color="#bd852e"/></linearGradient></defs>
      <g fill="url(#gGold)" stroke="#8d5c20" stroke-width="2.2" stroke-linejoin="round">
        <path d="M252 151 C207 106 133 66 16 61 C58 92 76 119 87 148 C52 135 28 137 8 149 C68 170 104 199 133 235 C96 224 68 229 45 245 C112 255 163 279 211 310 L260 278 L309 310 C357 279 408 255 475 245 C452 229 424 224 387 235 C416 199 452 170 512 149 C492 137 468 135 433 148 C444 119 462 92 504 61 C387 66 313 106 268 151 Z"/>
        <path d="M251 159 C213 127 161 101 91 91 C133 125 159 158 177 193 C143 181 116 181 91 190 C147 209 184 238 224 272 L260 253 L296 272 C336 238 373 209 429 190 C404 181 377 181 343 193 C361 158 387 125 429 91 C359 101 307 127 269 159 Z" fill="url(#gLight)"/>
        <path d="M260 119 C240 89 235 58 247 30 C256 44 263 53 270 57 C278 39 291 28 309 22 C302 44 302 61 311 76 C329 70 343 72 357 82 C335 88 322 103 313 126 L297 164 L223 164 L207 126 C198 103 185 88 163 82 C177 72 191 70 209 76 C218 61 218 44 211 22 C229 28 242 39 250 57 C257 53 264 44 273 30 C285 58 280 89 260 119 Z"/>
        <path d="M228 151 Q260 129 292 151 L306 226 Q289 263 260 287 Q231 263 214 226 Z" fill="url(#gLight)"/>
        <path d="M238 167 L260 182 L282 167 L291 224 L260 260 L229 224 Z" fill="#5b1018" stroke="#f1d17a" stroke-width="3"/>
        <path d="M250 284 L260 317 L270 284 L291 334 L260 358 L229 334 Z"/>
        <path d="M225 205 L190 227 L220 238 L244 226 M295 205 L330 227 L300 238 L276 226" fill="none" stroke="#8d5c20" stroke-width="5"/>
        <path d="M245 119 Q260 104 275 119 L269 143 L251 143 Z" fill="#fff0b0"/><path d="M260 111 Q276 92 294 103 L280 120 Z"/><path d="M285 119 L318 126 L288 137 Z" fill="#f5d77f"/>
      </g><g fill="none" stroke="#fff0b0" stroke-width="2" opacity=".8"><path d="M34 83 Q111 103 168 171 M67 157 Q133 178 188 222 M486 83 Q409 103 352 171 M453 157 Q387 178 332 222"/><path d="M260 153 V279"/></g>
    </svg>
  </div>
  <a class="cta" href="/pendaftaran-anggota">DAFTAR SEKARANG <span aria-hidden="true">→</span></a>
</main></body></html>"""
    return HTMLResponse(html)


@app.get("/health")
def health():
    with get_db() as connection:
        connection.execute("SELECT 1")

    return {
        "status": "ok",
        "database": "connected",
    }


# =========================================================
# MANAJEMEN ANGGOTA
# =========================================================

@app.post("/admin/members")
def create_member(
    member: MemberCreate,
    _: bool = Depends(verify_api_key),
):
    member_id = member.member_id.strip()
    member_name = member.member_name.strip()

    if not member_id or not member_name:
        raise HTTPException(
            status_code=400,
            detail="ID dan nama anggota wajib diisi",
        )

    try:
        with get_db() as connection:
            connection.execute(
                """
                INSERT INTO members (
                    member_id,
                    member_name,
                    created_at
                )
                VALUES (%s, %s, %s)
                """,
                (
                    member_id,
                    member_name,
                    datetime.now(timezone.utc).isoformat(),
                ),
            )

    except psycopg.errors.UniqueViolation:
        raise HTTPException(
            status_code=409,
            detail="ID anggota sudah terdaftar",
        )

    return {
        "message": "Anggota berhasil didaftarkan",
        "member_id": member_id,
        "member_name": member_name,
    }


@app.get("/admin/members")
def list_members(
    _: bool = Depends(verify_api_key),
):
    with get_db() as connection:
        rows = connection.execute(
            """
            SELECT
                member_id,
                member_name,
                active,
                created_at
            FROM members
            ORDER BY member_id
            """
        ).fetchall()

    return [dict(row) for row in rows]


# =========================================================
# PENGIRIMAN DATA SURVEI GPS
# =========================================================

@app.post("/surveys")
def create_survey(
    survey: SurveyCreate,
    identity: dict = Depends(get_request_identity),
):
    if identity["auth_type"] == "officer" and identity.get("role") != "DPSS":
        raise HTTPException(status_code=403, detail="Akun DPS menggunakan menu verifikasi, bukan membuat survei awal.")
    member_id = (identity["member_id"] if identity["auth_type"] == "officer" else survey.member_id).strip()

    with get_db() as connection:
        member = connection.execute(
            """
            SELECT member_id
            FROM members
            WHERE member_id = %s
              AND active = 1
            """,
            (member_id,),
        ).fetchone()

        if member is None:
            raise HTTPException(
                status_code=403,
                detail=(
                    "ID anggota tidak terdaftar "
                    "atau tidak aktif"
                ),
            )

        row = connection.execute(
            """
            INSERT INTO surveys (
                member_id, latitude, longitude, accuracy, notes,
                surveyed_at, received_at, applicant_name, applicant_nik,
                applicant_phone, applicant_address, village, business_name,
                business_type, business_since, business_location, ownership_status,
                nib, capital_amount, turnover_amount, net_profit_amount,
                worker_count, business_active, business_issues, assistance_types,
                assistance_amount, assistance_reason, business_verified,
                verification_status, verification_notes
            )
            VALUES (
                %s, %s, %s, %s, %s, %s, %s, %s, %s, %s,
                %s, %s, %s, %s, %s, %s, %s, %s, %s, %s,
                %s, %s, %s, %s, %s, %s, %s, %s, %s, %s
            )
            RETURNING id
            """,
            (
                member_id, survey.latitude, survey.longitude, survey.accuracy,
                survey.notes, survey.surveyed_at.isoformat(),
                datetime.now(timezone.utc).isoformat(), survey.applicant_name.strip(),
                survey.applicant_nik.strip(), survey.applicant_phone.strip(),
                survey.applicant_address.strip(), survey.village.strip(),
                survey.business_name.strip(), survey.business_type.strip(),
                survey.business_since.strip(), survey.business_location.strip(),
                survey.ownership_status.strip(), survey.nib.strip(),
                survey.capital_amount, survey.turnover_amount, survey.net_profit_amount,
                survey.worker_count, survey.business_active.strip(),
                survey.business_issues.strip(), survey.assistance_types.strip(),
                survey.assistance_amount, survey.assistance_reason.strip(),
                "Belum diverifikasi",
                "Perlu ditinjau",
                "",
            ),
        ).fetchone()

        survey_id = row["id"]
        village_name = survey.village.strip()
        village_row = None
        if village_name:
            village_row = connection.execute(
                "SELECT id, village_name FROM villages WHERE LOWER(TRIM(village_name)) = LOWER(%s) AND active = 1",
                (village_name,),
            ).fetchone()
        if identity["auth_type"] == "officer":
            if village_row is None:
                raise HTTPException(status_code=400, detail="Pilih desa yang terdaftar: Aur Sati atau Balam Jaya.")
            assigned = connection.execute("""SELECT 1 FROM officer_villages ov JOIN officers o ON o.id=ov.officer_id
                WHERE o.member_id=%s AND o.active=1 AND ov.village_id=%s""", (member_id, village_row["id"])).fetchone()
            if not assigned:
                raise HTTPException(status_code=403, detail="Akun DPSS belum ditugaskan ke desa survei ini.")
            connection.execute("UPDATE surveys SET village_id=%s, village=%s WHERE id=%s", (village_row["id"], village_row["village_name"], survey_id))
        elif village_row:
            connection.execute("UPDATE surveys SET village_id=%s, village=%s WHERE id=%s", (village_row["id"], village_row["village_name"], survey_id))

    return {
        "message": "Data survei berhasil disimpan",
        "survey_id": survey_id,
        "member_id": member_id,
    }


# =========================================================
# MELIHAT HASIL SURVEI
# =========================================================

@app.get("/admin/surveys")
def list_surveys(
    member_id: str | None = None,
    _: bool = Depends(verify_api_key),
):
    with get_db() as connection:
        if member_id:
            rows = connection.execute(
                """
                SELECT *
                FROM surveys
                WHERE member_id = %s
                ORDER BY id DESC
                """,
                (member_id,),
            ).fetchall()
        else:
            rows = connection.execute(
                """
                SELECT *
                FROM surveys
                ORDER BY id DESC
                """
            ).fetchall()

    return [dict(row) for row in rows]



# =========================================================
# UNGGAH FOTO DAN VIDEO SURVEI
# =========================================================

def _validate_file_signature(content: bytes, content_type: str) -> bool:
    """Validasi sederhana signature berkas, bukan hanya Content-Type dari klien."""
    if content_type == "image/jpeg":
        return content.startswith(b"\xff\xd8\xff")
    if content_type == "image/png":
        return content.startswith(b"\x89PNG\r\n\x1a\n")
    if content_type == "image/webp":
        return len(content) >= 12 and content.startswith(b"RIFF") and content[8:12] == b"WEBP"
    if content_type in {"video/mp4", "video/quicktime"}:
        return len(content) >= 12 and content[4:8] == b"ftyp"
    return False


def require_dps(identity: dict = Depends(get_officer_identity)):
    if identity.get("role") != "DPS":
        raise HTTPException(status_code=403, detail="Fitur ini khusus akun petugas DPS.")
    return identity


@app.get("/dps/surveys")
def list_surveys_for_verification(
    include_completed: bool = False,
    identity: dict = Depends(require_dps),
):
    """Daftar survei awal yang dapat dibuka petugas DPS untuk verifikasi."""
    with get_db() as connection:
        # DPS melihat survei dari seluruh desa dalam kecamatan pada akun DPS.
        # Penugasan desa manual yang sudah ada juga tetap dihormati.
        base_query = """SELECT DISTINCT s.*
            FROM surveys s
            JOIN officers o ON o.member_id = %s AND o.active = 1
            WHERE (
                EXISTS (
                    SELECT 1 FROM officer_villages ov
                    WHERE ov.officer_id = o.id AND ov.village_id = s.village_id
                )
                OR EXISTS (
                    SELECT 1 FROM villages v
                    WHERE v.active = 1
                      AND lower(trim(regexp_replace(coalesce(v.district, ''), '^(kecamatan|kec\\.?)[[:space:]]+', '', 'i'))) =
                          lower(trim(regexp_replace(regexp_replace(coalesce(o.region, ''), '^(kecamatan|kec\\.?)[[:space:]]+', '', 'i'), '[,;].*$', '')))
                      AND (
                          v.id = s.village_id
                          OR lower(trim(regexp_replace(coalesce(v.village_name, ''), '^(desa|kelurahan)[[:space:]]+', '', 'i'))) =
                             lower(trim(regexp_replace(coalesce(s.village, ''), '^(desa|kelurahan)[[:space:]]+', '', 'i')))
                      )
                )
            )"""
        if not include_completed:
            base_query += " AND COALESCE(s.verification_status, '') NOT IN ('Sudah diverifikasi', 'Ditolak')"
        base_query += " ORDER BY s.id DESC"
        rows = connection.execute(base_query, (identity["member_id"],)).fetchall()
    return [dict(row) for row in rows]


@app.get("/dps/surveys/{survey_id}")
def get_survey_for_verification(
    survey_id: int,
    identity: dict = Depends(require_dps),
):
    with get_db() as connection:
        survey = connection.execute("SELECT * FROM surveys WHERE id = %s", (survey_id,)).fetchone()
        if survey is None:
            raise HTTPException(status_code=404, detail="Data survei tidak ditemukan.")
        ensure_dps_survey_access(connection, identity, survey_id)
        original_media = connection.execute(
            """SELECT id, media_category, content_type, original_filename, object_path, file_size, uploaded_at, storage_bucket
               FROM survey_media WHERE survey_id = %s ORDER BY id""", (survey_id,)
        ).fetchall()
        verification = connection.execute(
            "SELECT * FROM survey_verifications WHERE survey_id = %s", (survey_id,)
        ).fetchone()
        verification_media = connection.execute(
            """SELECT id, content_type, original_filename, object_path, file_size, uploaded_at, storage_bucket
               FROM verification_media WHERE survey_id = %s ORDER BY id""", (survey_id,)
        ).fetchall()

    def media_with_url(row):
        item = dict(row)
        item["signed_url"] = create_supabase_signed_url(item["object_path"], bucket=item.get("storage_bucket", "survey-media"))
        return item

    return {
        "survey": dict(survey),
        "survey_media": [media_with_url(row) for row in original_media],
        "verification": dict(verification) if verification else None,
        "verification_media": [media_with_url(row) for row in verification_media],
    }


@app.post("/dps/surveys/{survey_id}/verification-media")
async def upload_verification_media(
    survey_id: int,
    file: UploadFile = File(...),
    identity: dict = Depends(require_dps),
):
    content_type = (file.content_type or "").lower()
    extensions = {"image/jpeg": ".jpg", "image/png": ".png", "image/webp": ".webp"}
    if content_type not in extensions:
        raise HTTPException(status_code=400, detail="Bukti verifikasi harus berupa foto JPEG, PNG, atau WebP.")
    content = await file.read(20 * 1024 * 1024 + 1)
    if not content:
        raise HTTPException(status_code=400, detail="File foto kosong.")
    if not _validate_file_signature(content, content_type):
        raise HTTPException(status_code=400, detail="Isi file tidak cocok dengan format foto yang dikirim.")
    if len(content) > 20 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="Ukuran setiap foto verifikasi maksimal 20 MB.")
    supabase_url = os.environ.get("SUPABASE_URL", "").strip().rstrip("/")
    supabase_key = os.environ.get("SUPABASE_SECRET_KEY", "").strip()
    if not supabase_url or not supabase_key:
        raise HTTPException(status_code=503, detail="Konfigurasi penyimpanan Supabase belum lengkap.")
    with get_db() as connection:
        survey = connection.execute("SELECT id, village_id FROM surveys WHERE id = %s", (survey_id,)).fetchone()
        if survey is None:
            raise HTTPException(status_code=404, detail="Data survei tidak ditemukan.")
        ensure_dps_survey_access(connection, identity, survey_id)
        storage_bucket = get_village_bucket(connection, survey["village_id"])
        count = connection.execute("SELECT COUNT(*) AS total FROM verification_media WHERE survey_id = %s", (survey_id,)).fetchone()["total"]
        if count >= 5:
            raise HTTPException(status_code=400, detail="Maksimal 5 foto bukti verifikasi per survei.")

    object_path = f"verifications/{survey_id}/{uuid.uuid4().hex}{extensions[content_type]}"
    storage_url = f"{supabase_url}/storage/v1/object/{quote(storage_bucket, safe='')}/{quote(object_path, safe='/')}"
    request = urllib.request.Request(storage_url, data=content, headers={
        "Authorization": f"Bearer {supabase_key}", "apikey": supabase_key,
        "Content-Type": content_type, "x-upsert": "false",
    }, method="POST")
    try:
        with urllib.request.urlopen(request, timeout=120) as response:
            if response.status not in (200, 201):
                raise HTTPException(status_code=502, detail="Supabase gagal menyimpan foto verifikasi.")
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        print(f"Supabase verification upload error: HTTP {exc.code}; {body[:500]}", flush=True)
        raise HTTPException(status_code=502, detail=f"Supabase menolak foto verifikasi (HTTP {exc.code}).") from exc
    except (urllib.error.URLError, TimeoutError) as exc:
        raise HTTPException(status_code=502, detail="Supabase tidak dapat dijangkau saat mengunggah foto verifikasi.") from exc

    with get_db() as connection:
        connection.execute(
            """INSERT INTO verification_media (survey_id, verifier_id, object_path, content_type, original_filename, file_size, storage_bucket)
               VALUES (%s, %s, %s, %s, %s, %s, %s)""",
            (survey_id, identity["member_id"], object_path, content_type, os.path.basename(file.filename or "foto-verifikasi"), len(content), storage_bucket),
        )
    return {"message": "Foto verifikasi berhasil diunggah.", "survey_id": survey_id, "file_path": object_path}


@app.post("/dps/surveys/{survey_id}/verification")
def submit_survey_verification(
    survey_id: int,
    payload: VerificationCreate,
    identity: dict = Depends(require_dps),
):
    with get_db() as connection:
        exists = connection.execute("SELECT id FROM surveys WHERE id=%s", (survey_id,)).fetchone()
        if exists is None:
            raise HTTPException(status_code=404, detail="Data survei tidak ditemukan.")
        ensure_dps_survey_access(connection, identity, survey_id)
    decision = payload.decision.strip()
    if decision not in {"Sudah diverifikasi", "Ditolak"}:
        raise HTTPException(status_code=422, detail="Pilihan keputusan hanya 'Sudah diverifikasi' atau 'Ditolak'.")
    if decision == "Ditolak" and not payload.rejection_reason.strip():
        raise HTTPException(status_code=422, detail="Alasan penolakan wajib diisi.")
    if not payload.business_condition.strip():
        raise HTTPException(status_code=422, detail="Kondisi nyata usaha wajib diisi.")
    if not payload.actual_business_activity.strip():
        raise HTTPException(status_code=422, detail="Aktivitas usaha aktual wajib diisi.")
    if not payload.data_match:
        raise HTTPException(status_code=422, detail="Kesesuaian data survei awal wajib dipilih.")
    if not payload.needs:
        raise HTTPException(status_code=422, detail="Tambahkan minimal satu item kebutuhan UMKM yang terperinci.")
    if payload.data_match and payload.data_match not in {"Sesuai", "Sebagian sesuai", "Tidak sesuai"}:
        raise HTTPException(status_code=422, detail="Kesesuaian data harus Sesuai, Sebagian sesuai, atau Tidak sesuai.")
    needs_data = [item.dict() for item in payload.needs]
    total_needs_amount = round(sum(item["quantity"] * item["unit_price"] for item in needs_data), 2)
    with get_db() as connection:
        survey = connection.execute("SELECT id FROM surveys WHERE id = %s", (survey_id,)).fetchone()
        if survey is None:
            raise HTTPException(status_code=404, detail="Data survei tidak ditemukan.")
        ensure_dps_survey_access(connection, identity, survey_id)
        existing = connection.execute("SELECT id FROM survey_verifications WHERE survey_id = %s", (survey_id,)).fetchone()
        if existing:
            raise HTTPException(status_code=409, detail="Survei ini sudah memiliki hasil verifikasi. Data awal dan hasil verifikasi tidak ditimpa.")
        connection.execute(
            """INSERT INTO survey_verifications (
                survey_id, verifier_id, verifier_name, verifier_region, decision,
                business_condition, actual_business_activity, existing_assets, needs,
                working_capital_details, verified_turnover, verified_profit, verified_worker_count,
                data_match, discrepancy_notes, verification_notes, rejection_reason,
                latitude, longitude, accuracy, total_needs_amount
            ) VALUES (
                %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s
            )""",
            (
                survey_id, identity["member_id"], identity["member_name"], identity["region"], decision,
                payload.business_condition.strip(), payload.actual_business_activity.strip(), payload.existing_assets.strip(),
                Jsonb(needs_data), payload.working_capital_details.strip(), payload.verified_turnover,
                payload.verified_profit, payload.verified_worker_count, payload.data_match.strip(),
                payload.discrepancy_notes.strip(), payload.verification_notes.strip(), payload.rejection_reason.strip(),
                payload.latitude, payload.longitude, payload.accuracy, total_needs_amount,
            ),
        )
        connection.execute(
            """UPDATE surveys SET business_verified = %s, verification_status = %s, verification_notes = %s WHERE id = %s""",
            ("Ya" if decision == "Sudah diverifikasi" else "Tidak", decision,
             payload.verification_notes.strip() or payload.rejection_reason.strip(), survey_id),
        )
    return {
        "message": "Hasil verifikasi berhasil disimpan.",
        "survey_id": survey_id,
        "decision": decision,
        "verifier_id": identity["member_id"],
        "total_needs_amount": total_needs_amount,
    }


@app.post("/surveys/{survey_id}/media")
async def upload_survey_media(
    survey_id: int,
    member_id: str = Form(...),
    media_category: str = Form("business_photo"),
    file: UploadFile = File(...),
    identity: dict = Depends(get_request_identity),
):
    if identity["auth_type"] == "officer":
        if identity.get("role") != "DPSS":
            raise HTTPException(status_code=403, detail="Akun DPS mengunggah bukti melalui menu verifikasi.")
        if member_id.strip().upper() != identity["member_id"].upper():
            raise HTTPException(status_code=403, detail="Petugas hanya boleh mengunggah media untuk surveinya sendiri.")
        member_id = identity["member_id"]

    allowed_types = {
        "image/jpeg": ".jpg",
        "image/png": ".png",
        "image/webp": ".webp",
        "video/mp4": ".mp4",
        "video/quicktime": ".mov",
    }

    content_type = (file.content_type or "").lower()
    if content_type not in allowed_types:
        raise HTTPException(
            status_code=400,
            detail="Jenis file tidak diizinkan. Gunakan JPEG, PNG, WebP, MP4, atau MOV.",
        )

    allowed_categories = {
        "applicant_ktp": (1, "image"),
        "family_card": (1, "image"),
        "recipient_photo": (1, "image"),
        "business_photo": (5, "image"),
        "business_video": (1, "video"),
    }
    if media_category not in allowed_categories:
        raise HTTPException(status_code=400, detail="Kategori media tidak dikenal.")
    category_limit, category_kind = allowed_categories[media_category]
    is_video = content_type.startswith("video/")
    if (category_kind == "video") != is_video:
        raise HTTPException(status_code=400, detail="Jenis file tidak sesuai kategori media.")

    max_size = 200 * 1024 * 1024
    content = await file.read(max_size + 1)
    if not content:
        raise HTTPException(status_code=400, detail="File kosong.")
    if not _validate_file_signature(content, content_type):
        raise HTTPException(status_code=400, detail="Isi file tidak cocok dengan format berkas yang dikirim.")
    if len(content) > max_size:
        raise HTTPException(status_code=413, detail="Ukuran file maksimal 200 MB.")

    supabase_url = os.environ.get("SUPABASE_URL", "").strip().rstrip("/")
    supabase_key = os.environ.get("SUPABASE_SECRET_KEY", "").strip()
    if not supabase_url or not supabase_key:
        raise HTTPException(
            status_code=503,
            detail="Konfigurasi penyimpanan Supabase belum lengkap.",
        )

    member_id = member_id.strip()
    with get_db() as connection:
        survey = connection.execute(
            """
            SELECT id, member_id, village_id
            FROM surveys
            WHERE id = %s AND member_id = %s
            """,
            (survey_id, member_id),
        ).fetchone()
        category_count = connection.execute(
            "SELECT COUNT(*) AS total FROM survey_media WHERE survey_id = %s AND media_category = %s",
            (survey_id, media_category),
        ).fetchone()["total"]

    if survey is None:
        raise HTTPException(
            status_code=404,
            detail="Data survei tidak ditemukan untuk ID anggota tersebut.",
        )

    if category_count >= category_limit:
        label = {
            "applicant_ktp": "foto KTP",
            "family_card": "foto KK",
            "recipient_photo": "foto penerima bantuan",
            "business_photo": "foto dokumentasi usaha (maksimal 5)",
            "business_video": "video usaha",
        }[media_category]
        raise HTTPException(status_code=400, detail=f"Batas {label} sudah tercapai.")

    # Bucket diselesaikan dari desa di database, bukan dari input bebas aplikasi.
    with get_db() as connection:
        storage_bucket = get_village_bucket(connection, survey["village_id"])
    object_path = f"surveys/{survey_id}/{media_category}/{uuid.uuid4().hex}{allowed_types[content_type]}"
    storage_url = (
        f"{supabase_url}/storage/v1/object/{quote(storage_bucket, safe='')}/"
        f"{quote(object_path, safe='/')}"
    )
    request = urllib.request.Request(
        storage_url,
        data=content,
        headers={
            "Authorization": f"Bearer {supabase_key}",
            "apikey": supabase_key,
            "Content-Type": content_type,
            "x-upsert": "false",
        },
        method="POST",
    )

    try:
        with urllib.request.urlopen(request, timeout=240) as response:
            if response.status not in (200, 201):
                raise HTTPException(
                    status_code=502,
                    detail="Supabase gagal menyimpan file.",
                )
    except urllib.error.HTTPError as exc:
        # Log respons Supabase untuk diagnosis. Secret key tidak dicatat.
        error_body = exc.read().decode("utf-8", errors="replace")
        print(
            f"Supabase upload error: HTTP {exc.code}; "
            f"response: {error_body[:1000]}",
            flush=True,
        )
        raise HTTPException(
            status_code=502,
            detail=f"Supabase menolak unggahan (HTTP {exc.code}).",
        ) from exc
    except (urllib.error.URLError, TimeoutError) as exc:
        print(
            f"Supabase connection error: {type(exc).__name__}: {exc}",
            flush=True,
        )
        raise HTTPException(
            status_code=502,
            detail="Supabase tidak dapat dijangkau saat mengunggah file.",
        ) from exc

    with get_db() as connection:
        connection.execute(
            """
            INSERT INTO survey_media (
                survey_id,
                member_id,
                object_path,
                content_type,
                original_filename,
                file_size,
                media_category, storage_bucket
            )
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
            """,
            (
                survey_id,
                member_id,
                object_path,
                content_type,
                os.path.basename(file.filename or "file"),
                len(content),
                media_category,
                storage_bucket,
            ),
        )

    return {
        "message": "File berhasil diunggah.",
        "survey_id": survey_id,
        "member_id": member_id,
        "file_path": object_path,
        "content_type": content_type,
        "file_size": len(content),
        "media_category": media_category,
    }


@app.get("/admin/surveys/{survey_id}/media")
def list_survey_media(
    survey_id: int,
    _: bool = Depends(verify_api_key),
):
    with get_db() as connection:
        rows = connection.execute(
            """
            SELECT id, survey_id, member_id, object_path,
                   content_type, original_filename, file_size, uploaded_at,
                   media_category
            FROM survey_media
            WHERE survey_id = %s
            ORDER BY id
            """,
            (survey_id,),
        ).fetchall()
    return [dict(row) for row in rows]


# =========================================================
# HAPUS SATU SURVEI BESERTA MEDIA TERKAIT
# =========================================================

def delete_survey_record(survey_id: int):
    """Hapus media dari bucket masing-masing, lalu hapus catatan survei."""
    with get_db() as connection:
        survey = connection.execute("SELECT id FROM surveys WHERE id=%s", (survey_id,)).fetchone()
        if survey is None:
            raise HTTPException(status_code=404, detail="Data survei tidak ditemukan.")
        media_rows = connection.execute("SELECT object_path, COALESCE(storage_bucket, 'survey-media') AS storage_bucket FROM survey_media WHERE survey_id=%s", (survey_id,)).fetchall()
        verification_rows = connection.execute("SELECT object_path, COALESCE(storage_bucket, 'survey-media') AS storage_bucket FROM verification_media WHERE survey_id=%s", (survey_id,)).fetchall()
    groups = {}
    for row in list(media_rows) + list(verification_rows):
        groups.setdefault(row["storage_bucket"] or "survey-media", []).append(row["object_path"])
    if groups:
        supabase_url = os.environ.get("SUPABASE_URL", "").strip().rstrip("/")
        supabase_key = os.environ.get("SUPABASE_SECRET_KEY", "").strip()
        if not supabase_url or not supabase_key:
            raise HTTPException(status_code=503, detail="Konfigurasi Supabase belum lengkap. Survei tidak dihapus.")
        for bucket, paths in groups.items():
            url = f"{supabase_url}/storage/v1/object/{quote(bucket, safe='')}"
            request = urllib.request.Request(url, data=json.dumps({"prefixes": paths}).encode("utf-8"), headers={
                "Authorization": f"Bearer {supabase_key}", "apikey": supabase_key, "Content-Type": "application/json"}, method="DELETE")
            try:
                with urllib.request.urlopen(request, timeout=45) as response:
                    if response.status not in (200, 204):
                        raise HTTPException(status_code=502, detail=f"Supabase tidak mengonfirmasi penghapusan bucket {bucket}.")
            except urllib.error.HTTPError as exc:
                body = exc.read().decode("utf-8", errors="replace")
                print(f"Supabase delete error bucket={bucket} HTTP={exc.code}: {body[:300]}", flush=True)
                raise HTTPException(status_code=502, detail=f"Gagal menghapus media dari bucket {bucket}. Survei belum dihapus dari database.") from exc
            except (urllib.error.URLError, TimeoutError) as exc:
                raise HTTPException(status_code=502, detail=f"Supabase tidak dapat dijangkau untuk bucket {bucket}. Survei belum dihapus dari database.") from exc
    with get_db() as connection:
        connection.execute("DELETE FROM survey_media WHERE survey_id=%s", (survey_id,))
        connection.execute("DELETE FROM verification_media WHERE survey_id=%s", (survey_id,))
        result = connection.execute("DELETE FROM surveys WHERE id=%s", (survey_id,))
        if result.rowcount == 0:
            raise HTTPException(status_code=404, detail="Data survei tidak ditemukan.")
    return {"message": "Survei dan media terkait berhasil dihapus.", "survey_id": survey_id, "media_deleted": sum(len(v) for v in groups.values())}


@app.delete("/admin/surveys/{survey_id}")
def delete_survey_api(
    survey_id: int,
    _: bool = Depends(verify_api_key),
):
    return delete_survey_record(survey_id)




# =========================================================
# DASHBOARD WEB ADMIN - DILINDUNGI LOGIN
# =========================================================

admin_security = HTTPBasic()


def verify_admin(
    credentials: HTTPBasicCredentials = Depends(admin_security),
):
    admin_username = os.environ.get("ADMIN_USERNAME", "")
    admin_password = os.environ.get("ADMIN_PASSWORD", "")

    if not admin_username or not admin_password:
        raise HTTPException(
            status_code=503,
            detail="Login admin belum dikonfigurasi di server."
        )

    username_ok = secrets.compare_digest(
        credentials.username.encode("utf-8"),
        admin_username.encode("utf-8"),
    )

    password_ok = secrets.compare_digest(
        credentials.password.encode("utf-8"),
        admin_password.encode("utf-8"),
    )

    if not (username_ok and password_ok):
        raise HTTPException(
            status_code=401,
            detail="Username atau kata sandi salah.",
            headers={"WWW-Authenticate": "Basic"},
        )

    return credentials.username


def create_supabase_signed_url(object_path: str, expires_in: int = 3600, bucket: str = "survey-media") -> str | None:
    """Buat URL sementara untuk menampilkan file dari bucket privat."""
    supabase_url = os.environ.get("SUPABASE_URL", "").strip().rstrip("/")
    supabase_key = os.environ.get("SUPABASE_SECRET_KEY", "").strip()
    if not supabase_url or not supabase_key:
        print("Dashboard media: konfigurasi Supabase belum lengkap.", flush=True)
        return None

    sign_url = (
        f"{supabase_url}/storage/v1/object/sign/{quote(bucket, safe='')}/"
        f"{quote(object_path, safe='/')}"
    )
    payload = json.dumps({"expiresIn": expires_in}).encode("utf-8")
    request = urllib.request.Request(
        sign_url,
        data=payload,
        headers={
            "Authorization": f"Bearer {supabase_key}",
            "apikey": supabase_key,
            "Content-Type": "application/json",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=15) as response:
            result = json.loads(response.read().decode("utf-8"))
        signed_path = result.get("signedURL") or result.get("signedUrl")
        if not signed_path:
            print("Dashboard media: respons URL bertanda tangan tidak berisi signedURL.", flush=True)
            return None
        if signed_path.startswith("http://") or signed_path.startswith("https://"):
            return signed_path
        if signed_path.startswith("/"):
            # Supabase biasanya mengembalikan path /object/sign/...
            if signed_path.startswith("/storage/v1/"):
                return f"{supabase_url}{signed_path}"
            return f"{supabase_url}/storage/v1{signed_path}"
        return f"{supabase_url}/storage/v1/{signed_path.lstrip('/')}"
    except (urllib.error.HTTPError, urllib.error.URLError, TimeoutError, ValueError) as exc:
        status = getattr(exc, "code", "connection")
        print(f"Dashboard media: gagal membuat signed URL ({status}).", flush=True)
        return None


@app.post("/admin/surveys/{survey_id}/delete")
def delete_survey_from_admin(
    survey_id: int,
    username: str = Depends(verify_admin),
):
    try:
        delete_survey_record(survey_id)
    except HTTPException as exc:
        message = escape(str(exc.detail))
        return HTMLResponse(
            content=(
                "<h2>Penghapusan gagal</h2>"
                f"<p>{message}</p>"
                '<p><a href="/admin">Kembali ke dashboard</a></p>'
            ),
            status_code=exc.status_code,
        )
    return RedirectResponse(url="/admin", status_code=303)


@app.get("/admin/kelola-akun", response_class=HTMLResponse)
def account_management_dashboard(credentials: HTTPBasicCredentials = Depends(admin_security)):
    # Halaman ini dapat dibuka admin maupun pimpinan yang kredensialnya dikonfigurasi.
    username = verify_pimpinan_or_admin(credentials)
    return HTMLResponse(r"""<!doctype html>
<html lang="id"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Kelola Akses Aplikasi | Triple 555</title>
<style>
:root{--maroon:#5b1018;--gold:#d8b15a;--ink:#292323;--muted:#6b7280;--line:#e8e1d8;--paper:#f6f3ee;--red:#b42318;--green:#17603b}
*{box-sizing:border-box}body{margin:0;background:var(--paper);color:var(--ink);font:15px/1.5 system-ui,Arial,sans-serif}.wrap{max-width:1180px;margin:auto;padding:22px 15px 45px}.hero{background:linear-gradient(120deg,#31090e,var(--maroon));color:#fff;padding:23px;border-radius:16px}.hero h1{margin:0 0 5px}.hero p{margin:0;color:#f5e8e8}.panel{background:#fff;border:1px solid var(--line);border-radius:14px;padding:18px;margin:15px 0}.notice{background:#fff8e8;border-left:4px solid var(--gold);padding:12px 14px;border-radius:8px;color:#59451d}.toolbar{display:flex;gap:10px;align-items:end;flex-wrap:wrap}.field{display:flex;flex-direction:column;gap:5px;min-width:220px;flex:1}label{font-weight:650;font-size:13px}input,select{font:inherit;min-height:42px;border:1px solid #d4d0ca;border-radius:8px;padding:8px 10px;width:100%}button,.btn{display:inline-block;border:0;border-radius:8px;padding:10px 14px;background:var(--maroon);color:#fff;text-decoration:none;font:inherit;font-weight:700;cursor:pointer}button.secondary,.btn.secondary{background:#eee8e1;color:#332b2b}button.danger{background:var(--red)}button:disabled{opacity:.55;cursor:wait}.status{display:none;padding:11px 13px;border-radius:8px;margin:12px 0;white-space:pre-wrap}.status.show{display:block}.status.ok{background:#eaf7ef;color:var(--green)}.status.err{background:#fff0ef;color:#8a1c13}.tablewrap{overflow:auto;margin-top:13px}table{border-collapse:collapse;width:100%;min-width:850px}th,td{text-align:left;padding:11px 9px;border-bottom:1px solid var(--line);vertical-align:top}th{font-size:12px;text-transform:uppercase;color:#625b55;background:#faf8f5}.pill{display:inline-block;padding:3px 9px;border-radius:20px;font-size:12px;font-weight:700;background:#e8f4ec;color:#17603b}.pill.off{background:#fce8e6;color:#9b1c14}.subtle{color:var(--muted);font-size:13px}.actions{display:flex;gap:7px;flex-wrap:wrap}.empty{text-align:center;padding:25px;color:var(--muted)}dialog{border:0;border-radius:14px;padding:0;width:min(520px,calc(100% - 28px));box-shadow:0 18px 70px #0004}dialog::backdrop{background:#17090bcc}.modalhead{background:var(--maroon);color:#fff;padding:17px 20px}.modalbody{padding:18px 20px}.modalbody .field{margin:12px 0}.modalfoot{display:flex;justify-content:flex-end;gap:8px;padding:0 20px 20px}.links{margin:14px 0;display:flex;gap:8px;flex-wrap:wrap}@media(max-width:650px){.wrap{padding:12px}.panel{padding:13px}.hero{padding:18px}.toolbar .field{min-width:100%}}
</style></head><body><main class="wrap">
<header class="hero"><h1>Kelola Akses Aplikasi</h1><p>Ikatan Persaudaraan Triple 555 · Admin/Pimpinan: __USERNAME__</p></header>
<div class="links"><a class="btn secondary" href="/admin/dashboard">Dashboard Utama Admin</a><a class="btn secondary" href="/admin/pendaftaran-anggota">Verifikasi Pendaftaran Anggota</a></div>
<section class="panel"><div class="notice"><b>Prinsip pengelolaan:</b> daftar anggota diambil dari pendaftaran yang sudah disetujui. Halaman ini tidak menginput ulang identitas anggota. Admin/pimpinan hanya memilih anggota yang ditugaskan, lalu memberikan atau mencabut akses aplikasi.</div>
<div id="status" class="status" role="status" aria-live="polite"></div>
<div class="toolbar" style="margin-top:16px"><div class="field"><label for="filter">Cari anggota</label><input id="filter" placeholder="Nama, NIK, jabatan, wilayah, atau ID login"></div><div class="field" style="max-width:220px"><label for="scope">Filter akses</label><select id="scope"><option value="all">Semua anggota disetujui</option><option value="none">Belum diberi akses</option><option value="active">Akses aktif</option><option value="inactive">Akses nonaktif</option></select></div><button id="refreshBtn" class="secondary" type="button">Muat Ulang</button></div>
<div class="tablewrap"><table><thead><tr><th>Anggota terverifikasi</th><th>Jabatan / Struktur</th><th>Wilayah</th><th>Akses aplikasi</th><th>Aksi</th></tr></thead><tbody id="membersBody"><tr><td colspan="5" class="empty">Memuat data anggota...</td></tr></tbody></table></div>
<p class="subtle">Data identitas bersumber dari pendaftaran anggota yang berstatus Disetujui. Pemberian akses tidak mengubah status keanggotaan. Menonaktifkan akses tidak menghapus data anggota maupun riwayat survei.</p></section>
</main>
<dialog id="grantDialog"><form id="grantForm"><div class="modalhead"><h2 style="margin:0;font-size:19px">Berikan Akses Aplikasi</h2><div id="selectedName" class="subtle" style="color:#f3e6e6"></div></div><div class="modalbody"><input type="hidden" id="registrationId"><div class="field"><label>ID Member otomatis (juga menjadi ID Login)</label><div id="memberCode" style="padding:11px;border:1px solid #d4d0ca;border-radius:8px;background:#f8f4ec;font-weight:800;color:#5b1018">-</div></div><div class="field"><label for="initialPassword">Kata sandi awal yang dibuat admin</label><input id="initialPassword" type="password" required minlength="10" maxlength="200" autocomplete="new-password" placeholder="Minimal 10 karakter"></div><p class="subtle">Peran DPS/DPSS dan wilayah tugas diambil otomatis dari data pendaftaran anggota yang telah disetujui.</p></div><div class="modalfoot"><button type="button" class="secondary" id="cancelGrant">Batal</button><button type="submit">Setujui & Berikan Akses</button></div></form></dialog>
<script>
(function(){const $=id=>document.getElementById(id);let members=[];
function msg(text,ok=false){const e=$('status');e.textContent=text;e.className='status show '+(ok?'ok':'err')}
async function api(path,options={}){let r;try{r=await fetch(path,Object.assign({credentials:'same-origin'},options))}catch(e){throw Error('Koneksi ke server gagal. Periksa internet lalu coba lagi.')}let d={};try{d=await r.json()}catch(e){}if(!r.ok)throw Error(d.detail||d.message||('Permintaan gagal ('+r.status+').'));return d}
function esc(s){return String(s??'').replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]))}
function state(m){if(!m.login_member_id)return '<span class="pill off">Belum diberi akses</span>';return Number(m.account_active)===1?'<span class="pill">Akses aktif</span>':'<span class="pill off">Akses nonaktif</span>'}
function render(){const q=$('filter').value.trim().toLowerCase(),scope=$('scope').value;let rows=members.filter(m=>[m.full_name,m.nik,m.phone,m.org_position,m.org_level,m.region,m.village_name,m.member_code,m.login_member_id].join(' ').toLowerCase().includes(q));rows=rows.filter(m=>scope==='all'||(scope==='none'&&!m.login_member_id)||(scope==='active'&&m.login_member_id&&Number(m.account_active)===1)||(scope==='inactive'&&m.login_member_id&&Number(m.account_active)!==1));$('membersBody').innerHTML=rows.length?rows.map(m=>`<tr><td><b>${esc(m.full_name)}</b><br><span class="subtle">NIK: ${esc(m.nik)}</span><br><span class="subtle">HP: ${esc(m.phone)}</span></td><td>${esc(m.org_position)}<br><span class="subtle">${esc(m.org_level)}</span></td><td>${esc(m.region||'-')}${m.village_name?'<br><span class="subtle">Desa: '+esc(m.village_name)+'</span>':''}</td><td>${state(m)}<br><span class="subtle">ID Member: ${esc(m.member_code||'-')}</span>${m.login_member_id?'<br><span class="subtle">Akses dengan ID: '+esc(m.login_member_id)+'</span>':''}</td><td><div class="actions">${!m.login_member_id?`<button type="button" data-grant="${m.id}">Beri Akses</button>`:(Number(m.account_active)===1?`<button type="button" class="danger" data-disable="${m.id}">Cabut Akses</button>`:`<button type="button" data-enable="${m.id}">Aktifkan Akses</button>`)}</div></td></tr>`).join(''):'<tr><td colspan="5" class="empty">Tidak ada anggota yang cocok. Pastikan pendaftaran sudah berstatus Disetujui.</td></tr>'}
async function load(){ $('membersBody').innerHTML='<tr><td colspan="5" class="empty">Memuat data anggota terverifikasi...</td></tr>';try{const d=await api('/admin/kelola-akun/data');members=Array.isArray(d)?d:[];render();msg('Data dimuat: '+members.length+' anggota terverifikasi.',true)}catch(e){msg(e.message);$('membersBody').innerHTML='<tr><td colspan="5" class="empty">Gagal memuat data.</td></tr>'}}
$('refreshBtn').addEventListener('click',load);$('filter').addEventListener('input',render);$('scope').addEventListener('change',render);
$('membersBody').addEventListener('click',async e=>{const g=e.target.closest('[data-grant]'),d=e.target.closest('[data-disable]'),a=e.target.closest('[data-enable]');if(g){const m=members.find(x=>String(x.id)===g.dataset.grant);if(!m)return;$('registrationId').value=m.id;$('selectedName').textContent=m.full_name+' · '+m.org_level+' · '+m.org_position;$('memberCode').textContent=m.member_code||'ID belum tersedia';$('initialPassword').value='';$('grantDialog').showModal();return}if(d){const m=members.find(x=>String(x.id)===d.dataset.disable);if(!m||!confirm('Cabut akses aplikasi untuk '+m.full_name+'? Data anggota dan riwayat survei tetap disimpan.'))return;d.disabled=true;try{const r=await api('/admin/kelola-akun/'+m.id+'/nonaktifkan',{method:'POST'});msg(r.message,true);await load()}catch(err){msg(err.message);d.disabled=false}return}if(a){const m=members.find(x=>String(x.id)===a.dataset.enable);if(!m||!confirm('Aktifkan kembali akses aplikasi untuk '+m.full_name+'?'))return;a.disabled=true;try{const r=await api('/admin/kelola-akun/'+m.id+'/aktifkan',{method:'POST'});msg(r.message,true);await load()}catch(err){msg(err.message);a.disabled=false}}});
$('cancelGrant').addEventListener('click',()=>$('grantDialog').close());$('grantForm').addEventListener('submit',async e=>{e.preventDefault();const b=e.submitter;b.disabled=true;try{const id=$('registrationId').value;const r=await api('/admin/kelola-akun/'+id+'/akses',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({password:$('initialPassword').value})});$('grantDialog').close();msg(r.message+' ID Member/Login: '+r.member_id,true);await load()}catch(err){msg(err.message)}finally{b.disabled=false}});
load();})();
</script></body></html>""".replace("__USERNAME__", escape(username)))


@app.get("/admin/dashboard", response_class=HTMLResponse)
def admin_main_dashboard(username: str = Depends(verify_admin)):
    """Halaman utama khusus admin untuk mengakses seluruh fungsi administrasi."""
    with get_db() as connection:
        stats = connection.execute("""
            SELECT
                (SELECT COUNT(*) FROM organization_member_registrations) AS registrations_total,
                (SELECT COUNT(*) FROM organization_member_registrations WHERE status='pending') AS registrations_pending,
                (SELECT COUNT(*) FROM organization_member_registrations WHERE status='revision') AS registrations_revision,
                (SELECT COUNT(*) FROM organization_member_registrations WHERE status='approved') AS registrations_approved,
                (SELECT COUNT(*) FROM organization_member_registrations WHERE status='rejected') AS registrations_rejected,
                (SELECT COUNT(*) FROM officers) AS accounts_total,
                (SELECT COUNT(*) FROM officers WHERE active=1) AS accounts_active,
                (SELECT COUNT(*) FROM officers WHERE active<>1 OR active IS NULL) AS accounts_inactive,
                (SELECT COUNT(*) FROM surveys) AS surveys_total,
                (SELECT COUNT(*) FROM survey_verifications) AS verifications_total
        """).fetchone()
        recent = connection.execute("""
            SELECT full_name, registration_code, org_level, org_position, region, status, created_at
            FROM organization_member_registrations
            WHERE status IN ('pending','revision')
            ORDER BY created_at DESC
            LIMIT 6
        """).fetchall()

    def stat(key):
        return int((stats or {}).get(key) or 0)

    recent_rows = []
    for row in recent:
        status = str(row.get('status') or '')
        status_label = 'Perlu perbaikan' if status == 'revision' else 'Menunggu verifikasi'
        status_class = 'revision' if status == 'revision' else 'pending'
        recent_rows.append(
            "<tr>"
            f"<td><b>{escape(str(row.get('full_name') or '-'))}</b><br><small>{escape(str(row.get('registration_code') or '-'))}</small></td>"
            f"<td>{escape(str(row.get('org_level') or '-'))} · {escape(str(row.get('org_position') or '-'))}</td>"
            f"<td>{escape(str(row.get('region') or '-'))}</td>"
            f"<td><span class='badge {status_class}'>{status_label}</span></td>"
            "</tr>"
        )
    recent_html = ''.join(recent_rows) or "<tr><td colspan='4' class='empty'>Tidak ada pendaftaran yang menunggu tindakan.</td></tr>"

    return HTMLResponse(f"""<!doctype html>
<html lang='id'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>
<title>Dashboard Utama Admin | Triple 555</title>
<style>
:root{{--maroon:#650b18;--deep:#17070b;--gold:#e0ad35;--paper:#f4f0eb;--ink:#f8f1e8;--muted:#b8a8a0;--line:#503337;--green:#39a66a;--red:#b82b3b}}
*{{box-sizing:border-box}}body{{margin:0;background:radial-gradient(ellipse at top right,#35151b 0,#16090c 55%,#100709 100%);color:var(--ink);font:14px/1.5 Inter,system-ui,Arial,sans-serif}}
.wrap{{width:100%;max-width:none;margin:0;padding:0 0 36px}}.hero{{min-height:104px;background:linear-gradient(110deg,#19080c,#500813 55%,#26090e);color:white;padding:16px 28px;border-bottom:2px solid var(--gold);display:flex;align-items:center;justify-content:space-between;gap:18px;flex-wrap:wrap;border-radius:0;box-shadow:0 8px 30px #0006}}
.brand{{display:flex;align-items:center;gap:16px}}.brand img{{width:66px;height:66px;object-fit:contain;border-radius:50%;background:#2b1114;padding:3px;border:2px solid var(--gold)}}h1{{margin:0 0 4px;font-size:clamp(22px,3vw,30px);color:#f7d77f}}.hero p{{margin:0;color:#e8d8c7}}.user{{background:#ffffff12;border:1px solid #d8b15a66;padding:9px 13px;border-radius:10px;font-size:13px;color:#f6e8cf}}
.layout{{display:grid;grid-template-columns:224px minmax(0,1fr);gap:22px;max-width:1700px;margin:0 auto;padding:22px 24px}}.sidebar{{background:linear-gradient(180deg,#26090f,#14080b);border:1px solid #6e3a3d;border-radius:16px;padding:16px 12px;min-height:calc(100vh - 150px);display:flex;flex-direction:column;gap:7px;box-shadow:0 12px 28px #0005}}.sidebar .side-label{{font-size:10px;letter-spacing:1.8px;color:#c9a65a;padding:8px 12px 10px;text-transform:uppercase}}.sidebar a{{display:flex;align-items:center;gap:11px;color:#f5e9dc;text-decoration:none;padding:12px 13px;border-radius:9px;font-weight:650;border:1px solid transparent}}.sidebar a:hover,.sidebar a.active{{background:linear-gradient(100deg,#e9b83f,#b87a16);color:#28130b}}
.main-content{{min-width:0;padding:0 2px 20px}}.section-title{{margin:8px 0 13px;font-size:18px;color:#f2cf78}}.stats{{display:grid;grid-template-columns:repeat(5,minmax(0,1fr));gap:12px}}.stat,.card,.panel{{background:linear-gradient(145deg,#26191b,#1c1114);border:1px solid #5b3939;border-radius:13px;box-shadow:0 8px 20px #0003;color:#f7eee5}}.stat{{padding:16px;position:relative;overflow:hidden}}.stat span{{display:block;color:#c9b8ad;font-size:12px;font-weight:650}}.stat strong{{display:block;font-size:29px;color:#f2c34f;margin-top:5px;line-height:1.2}}.stat small{{color:#a99b93}}
.menu{{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:14px}}.card{{padding:18px;display:flex;gap:14px;align-items:flex-start}}.icon{{width:45px;height:45px;flex:0 0 45px;border-radius:12px;background:linear-gradient(145deg,#e9bd51,#a96d16);color:#2a1609;display:grid;place-items:center;font-size:22px}}.card h3{{margin:0 0 5px;color:#f1cb72;font-size:16px}}.card p{{margin:0 0 13px;color:#c0b2aa;font-size:13px}}.btn{{display:inline-block;background:linear-gradient(110deg,#f0c85e,#bd821b);color:#24130a;text-decoration:none;padding:9px 13px;border-radius:8px;font-weight:800;font-size:13px}}
.panel{{padding:17px;margin-top:15px;overflow:hidden}}.panel h2{{font-size:17px;color:#f1cb72;margin:0 0 12px}}.tablewrap{{overflow:auto}}table{{width:100%;border-collapse:collapse;min-width:650px}}th,td{{text-align:left;padding:11px 9px;border-bottom:1px solid #493235;vertical-align:top}}th{{font-size:11px;text-transform:uppercase;color:#e7c46e;background:#301b20}}td{{color:#eee3d9}}small{{color:#b8a8a0}}.badge{{display:inline-block;border-radius:20px;padding:4px 9px;font-size:11px;font-weight:750}}.pending{{background:#604316;color:#ffdf8a}}.revision{{background:#62202a;color:#ffc0c6}}.empty{{text-align:center;color:#b8a8a0;padding:22px}}.foot{{margin-top:22px;text-align:center;color:#988780;font-size:12px}}
@media(max-width:1200px){{.stats{{grid-template-columns:repeat(3,minmax(0,1fr))}}}}@media(max-width:800px){{.layout{{grid-template-columns:1fr;padding:14px}}.sidebar{{min-height:0;flex-direction:row;overflow-x:auto;white-space:nowrap}}.sidebar .side-label{{display:none}}.sidebar a{{flex:0 0 auto}}}}@media(max-width:620px){{.hero{{padding:15px}}.brand img{{width:54px;height:54px}}.stats{{grid-template-columns:repeat(2,minmax(0,1fr));gap:8px}}.stat{{padding:12px}}.stat strong{{font-size:23px}}.menu{{grid-template-columns:1fr}}.card{{padding:14px}}.layout{{padding:10px}}}}
</style></head><body><main class='wrap'>
<header class='hero'><div class='brand'><img src='/logo.png' alt='Logo 555'><div><h1>555 · ADMIN CONTROL CENTER</h1><p>Ikatan Persaudaraan Triple 555 · Pusat administrasi organisasi</p></div></div><div class='user'>Administrator: <b>{escape(username)}</b></div></header>
<div class='layout'><aside class='sidebar'><div class='side-label'>Menu Administrasi</div><a class='active' href='/admin/dashboard'>▦ &nbsp; Dashboard</a><a href='/admin/pendaftaran-anggota'>✎ &nbsp; Verifikasi Anggota</a><a href='/admin/data-anggota'>▤ &nbsp; Data Anggota</a><a href='/admin/data-personel'>♙ &nbsp; Data Personel</a><a href='/admin/kelola-akun'>⚿ &nbsp; Kelola Akses</a><a href='/admin'>⌖ &nbsp; Data Survei</a></aside><div class='main-content'>
<h2 class='section-title'>Ringkasan Sistem</h2>
<section class='stats'>
<div class='stat'><span>Total Pendaftaran Anggota</span><strong>{stat('registrations_total')}</strong><small>Seluruh status</small></div>
<div class='stat'><span>Menunggu Verifikasi</span><strong>{stat('registrations_pending')}</strong><small>Perlu diperiksa admin</small></div>
<div class='stat'><span>Perlu Perbaikan</span><strong>{stat('registrations_revision')}</strong><small>Menunggu pembaruan data</small></div>
<div class='stat'><span>Anggota Disetujui</span><strong>{stat('registrations_approved')}</strong><small>Masuk daftar anggota terverifikasi</small></div>
<div class='stat'><span>Pendaftaran Ditolak</span><strong>{stat('registrations_rejected')}</strong><small>Sudah diputuskan</small></div>
<div class='stat'><span>Total Akun Aplikasi</span><strong>{stat('accounts_total')}</strong><small>Petugas yang pernah diberi akses</small></div>
<div class='stat'><span>Akun Aktif</span><strong>{stat('accounts_active')}</strong><small>Dapat menggunakan aplikasi</small></div>
<div class='stat'><span>Akun Nonaktif</span><strong>{stat('accounts_inactive')}</strong><small>Akses sedang dicabut</small></div>
<div class='stat'><span>Total Survei</span><strong>{stat('surveys_total')}</strong><small>Data survei tersimpan</small></div>
<div class='stat'><span>Hasil Verifikasi DPS</span><strong>{stat('verifications_total')}</strong><small>Verifikasi tersimpan</small></div>
</section>
<h2 class='section-title'>Menu Administrasi</h2>
<section class='menu'>
<article class='card'><div class='icon'>📝</div><div><h3>Verifikasi Pendaftaran Anggota</h3><p>Periksa pendaftaran, edit data yang keliru, setujui, minta perbaikan, atau tolak pendaftaran.</p><a class='btn' href='/admin/pendaftaran-anggota'>Buka Verifikasi</a></div></article>
<article class='card'><div class='icon'>🗂️</div><div><h3>Daftar Seluruh Anggota</h3><p>Lihat seluruh anggota DPS dan DPSS yang sudah disetujui, termasuk foto profil dan dokumen identitas.</p><a class='btn' href='/admin/data-anggota'>Buka Data Anggota</a></div></article>
<article class='card'><div class='icon'>👥</div><div><h3>Data Personel Petugas</h3><p>Lihat dan kelola data profil petugas yang sudah memiliki akun aplikasi.</p><a class='btn' href='/admin/data-personel'>Buka Data Personel</a></div></article>
<article class='card'><div class='icon'>🔐</div><div><h3>Kelola Akses Aplikasi</h3><p>Pilih anggota terverifikasi yang ditugaskan, berikan akses, aktifkan kembali, atau cabut akses aplikasi.</p><a class='btn' href='/admin/kelola-akun'>Kelola Akses</a></div></article>
<article class='card'><div class='icon'>📍</div><div><h3>Data Survei dan Media</h3><p>Pantau data survei, lokasi GPS, dokumen, foto, video, dan hasil verifikasi yang tersimpan.</p><a class='btn' href='/admin'>Buka Data Survei</a></div></article>
</section>
<section class='panel'><h2>Pendaftaran yang Memerlukan Tindakan</h2><div class='tablewrap'><table><thead><tr><th>Nama / Kode</th><th>Struktur / Jabatan</th><th>Wilayah</th><th>Status</th></tr></thead><tbody>{recent_html}</tbody></table></div><p style='margin:14px 0 0'><a class='btn' href='/admin/pendaftaran-anggota?status=pending'>Periksa Pendaftaran</a></p></section>
<div class='foot'>Dashboard Administrasi Triple 555 · Data ditampilkan langsung dari server</div>
</div></div></main></body></html>""")


@app.get("/admin", response_class=HTMLResponse)
def admin_dashboard(
    username: str = Depends(verify_admin),
):
    with get_db() as connection:
        rows = connection.execute(
            """
            SELECT *
            FROM surveys
            ORDER BY id DESC
            """
        ).fetchall()
        media_rows = connection.execute(
            """
            SELECT id, survey_id, member_id, object_path, content_type,
                   original_filename, file_size, uploaded_at, media_category, storage_bucket
            FROM survey_media
            ORDER BY id DESC
            """
        ).fetchall()
        verification_rows = connection.execute(
            "SELECT * FROM survey_verifications ORDER BY id DESC"
        ).fetchall()
        verification_media_rows = connection.execute(
            "SELECT * FROM verification_media ORDER BY id DESC"
        ).fetchall()

    verification_by_survey = {row["survey_id"]: row for row in verification_rows}
    verification_media_by_survey = {}
    for media in verification_media_rows:
        verification_media_by_survey.setdefault(media["survey_id"], []).append(media)

    media_by_survey = {}
    for media in media_rows:
        media_by_survey.setdefault(media["survey_id"], []).append(media)

    table_rows = ""
    for row in rows:
        survey_id = row["id"]
        lat = row["latitude"]
        lon = row["longitude"]
        map_url = escape(f"https://www.google.com/maps?q={lat},{lon}", quote=True)

        media_html = []
        for media in media_by_survey.get(survey_id, []):
            signed_url = create_supabase_signed_url(media["object_path"], bucket=media.get("storage_bucket") or "survey-media")
            filename = escape(str(media["original_filename"] or "Media"))
            media_category = escape(str(media.get("media_category") or "business_photo"))
            category_labels = {
                "applicant_ktp": "KTP Pemohon",
                "family_card": "KK Pemohon",
                "recipient_photo": "Foto Penerima",
                "business_photo": "Foto Usaha",
                "business_video": "Video Usaha",
            }
            category_label = category_labels.get(media_category, media_category)
            content_type = str(media["content_type"] or "")
            size_mb = media["file_size"] / (1024 * 1024)
            if signed_url:
                safe_url = escape(signed_url, quote=True)
                if content_type.startswith("image/"):
                    media_html.append(
                        f'<a href="{safe_url}" target="_blank" rel="noopener noreferrer">'
                        f'<img class="media-thumb" src="{safe_url}" alt="{filename}"></a>'
                        f'<div><strong>{category_label}</strong><br><a href="{safe_url}" target="_blank" rel="noopener noreferrer">{filename}</a>'
                        f'<small>{size_mb:.2f} MB</small></div>'
                    )
                elif content_type.startswith("video/"):
                    media_html.append(
                        f'<video class="media-video" controls preload="metadata" src="{safe_url}"></video>'
                        f'<div><strong>{category_label}</strong><br><a href="{safe_url}" target="_blank" rel="noopener noreferrer">{filename}</a>'
                        f'<small>{size_mb:.2f} MB</small></div>'
                    )
                else:
                    media_html.append(
                        f'<a href="{safe_url}" target="_blank" rel="noopener noreferrer">{filename}</a>'
                    )
            else:
                media_html.append(
                    f'<span>{filename} (tersimpan, pratinjau gagal dibuat)</span>'
                )
        for media in verification_media_by_survey.get(survey_id, []):
            signed_url = create_supabase_signed_url(media["object_path"], bucket=media.get("storage_bucket") or "survey-media")
            filename = escape(str(media.get("original_filename") or "Foto verifikasi"))
            if signed_url:
                safe_url = escape(signed_url, quote=True)
                media_html.append(
                    f'<a href="{safe_url}" target="_blank" rel="noopener noreferrer">'
                    f'<img class="media-thumb" src="{safe_url}" alt="{filename}"></a>'
                    f'<div><strong>Foto Verifikasi DPS</strong><br>'
                    f'<a href="{safe_url}" target="_blank" rel="noopener noreferrer">{filename}</a></div>'
                )
            else:
                media_html.append(f'<span>{filename} (pratinjau gagal dibuat)</span>')
        media_cell = "<br>".join(media_html) if media_html else '<span class="muted">Belum ada media</span>'
        verification_row = verification_by_survey.get(survey_id)
        needs_summary = "-"
        if verification_row:
            needs = verification_row.get("needs") or []
            if isinstance(needs, str):
                try:
                    needs = json.loads(needs)
                except ValueError:
                    needs = []
            needs_summary = "\n".join(
                f"- {item.get('item_name', 'Item')}: {item.get('quantity', 0)} {item.get('unit', 'unit')} x Rp {item.get('unit_price', 0):,.0f}"
                for item in needs
            ) or "Tidak ada item kebutuhan yang dicatat"

        detail_text = "\n".join([
            f"Pemohon: {row.get('applicant_name') or '-'}",
            f"NIK: {row.get('applicant_nik') or '-'}",
            f"Telepon: {row.get('applicant_phone') or '-'}",
            f"Alamat: {row.get('applicant_address') or '-'}",
            f"Desa: {row.get('village') or '-'}",
            f"Usaha: {row.get('business_name') or '-'}",
            f"Jenis usaha: {row.get('business_type') or '-'}",
            f"Mulai usaha: {row.get('business_since') or '-'}",
            f"Lokasi usaha: {row.get('business_location') or '-'}",
            f"Kepemilikan: {row.get('ownership_status') or '-'}",
            f"NIB: {row.get('nib') or '-'}",
            f"Modal: {row.get('capital_amount') if row.get('capital_amount') is not None else '-'}",
            f"Omzet: {row.get('turnover_amount') if row.get('turnover_amount') is not None else '-'}",
            f"Laba bersih: {row.get('net_profit_amount') if row.get('net_profit_amount') is not None else '-'}",
            f"Pekerja: {row.get('worker_count') if row.get('worker_count') is not None else '-'}",
            f"Status usaha: {row.get('business_active') or '-'}",
            f"Kendala: {row.get('business_issues') or '-'}",
            f"Bantuan diminta: {row.get('assistance_types') or '-'}",
            f"Perkiraan bantuan: {row.get('assistance_amount') if row.get('assistance_amount') is not None else '-'}",
            f"Alasan: {row.get('assistance_reason') or '-'}",
            f"Usaha ditemukan: {row.get('business_verified') or '-'}",
            f"Status verifikasi: {row.get('verification_status') or '-'}",
            f"Catatan verifikasi: {row.get('verification_notes') or '-'}",
            f"Catatan survei: {row.get('notes') or '-'}",
            f"\n=== HASIL VERIFIKASI DPS ===",
            f"Keputusan: {verification_row.get('decision') if verification_row else '-'}",
            f"Petugas DPS: {verification_row.get('verifier_name') if verification_row else '-'} ({verification_row.get('verifier_id') if verification_row else '-'})",
            f"Wilayah DPS: {verification_row.get('verifier_region') if verification_row else '-'}",
            f"Waktu verifikasi: {verification_row.get('verified_at') if verification_row else '-'}",
            f"Kondisi usaha: {verification_row.get('business_condition') if verification_row else '-'}",
            f"Aktivitas usaha aktual: {verification_row.get('actual_business_activity') if verification_row else '-'}",
            f"Aset yang sudah dimiliki: {verification_row.get('existing_assets') if verification_row else '-'}",
            f"Rincian kebutuhan: {needs_summary}",
            f"Total estimasi kebutuhan: Rp {verification_row.get('total_needs_amount', 0):,.0f}" if verification_row else "Total estimasi kebutuhan: -",
            f"Rincian modal kerja: {verification_row.get('working_capital_details') if verification_row else '-'}",
            f"Omzet terverifikasi: {verification_row.get('verified_turnover') if verification_row else '-'}",
            f"Laba terverifikasi: {verification_row.get('verified_profit') if verification_row else '-'}",
            f"Jumlah pekerja terverifikasi: {verification_row.get('verified_worker_count') if verification_row else '-'}",
            f"Kesesuaian data: {verification_row.get('data_match') if verification_row else '-'}",
            f"Perbedaan data: {verification_row.get('discrepancy_notes') if verification_row else '-'}",
            f"Catatan verifikasi: {verification_row.get('verification_notes') if verification_row else '-'}",
            f"Alasan penolakan: {verification_row.get('rejection_reason') if verification_row else '-'}",
            f"GPS verifikasi: {str(verification_row.get('latitude')) + ', ' + str(verification_row.get('longitude')) if verification_row else '-'}",
        ])
        table_rows += f"""
        <tr>
            <td>{survey_id}</td>
            <td>{escape(str(row['member_id']))}</td>
            <td class="notes">{escape(detail_text)}</td>
            <td>{lat}<br>{lon}<br><a href="{map_url}" target="_blank" rel="noopener noreferrer">Lihat Peta</a></td>
            <td>{escape(str(row['surveyed_at']))}</td>
            <td class="media-cell">{media_cell}</td>
            <td>
                <form method="post" action="/admin/surveys/{survey_id}/delete"
                      onsubmit="return confirm('Hapus permanen survei ID {survey_id} beserta foto dan videonya?');">
                    <button type="submit" class="delete-button">Delete</button>
                </form>
            </td>
        </tr>
        """

    if not table_rows:
        table_rows = '<tr><td colspan="7">Belum ada data survei.</td></tr>'

    page = f"""
    <!DOCTYPE html>
    <html lang="id">
    <head>
        <meta charset="UTF-8">
        <meta name="viewport" content="width=device-width, initial-scale=1">
        <meta http-equiv="refresh" content="60">
        <title>Dashboard Survei GPS</title>
        <style>
            body {{ font-family: Arial, sans-serif; margin: 0; padding: 24px; background: #f1f5f9; color: #1e293b; }}
            h1 {{ margin-bottom: 8px; }}
            .info {{ margin-bottom: 20px; color: #475569; }}
            .card {{ background: white; padding: 18px; border-radius: 10px; margin-bottom: 20px; }}
            .table-wrap {{ overflow-x: auto; }}
            table {{ border-collapse: collapse; width: 100%; background: white; font-size: 14px; }}
            th, td {{ border: 1px solid #e2e8f0; padding: 10px; text-align: left; vertical-align: top; }}
            th {{ background: #e2e8f0; white-space: nowrap; }}
            tr:nth-child(even) {{ background: #f8fafc; }}
            a {{ color: #0369a1; font-weight: bold; }}
            .notes {{ min-width: 280px; max-width: 440px; white-space: pre-wrap; }}
            .media-cell {{ min-width: 180px; max-width: 260px; overflow-wrap: anywhere; }}
            .media-thumb {{ display: block; width: auto; max-width: 180px; max-height: 150px; border-radius: 6px; margin: 4px 0; }}
            .media-video {{ display: block; width: 220px; max-width: 100%; margin: 4px 0; }}
            small {{ display: block; color: #64748b; margin-bottom: 8px; }}
            .muted {{ color: #64748b; }}
            .delete-button {{ display: inline-flex; align-items: center; justify-content: center; min-width: 58px; height: 30px; box-sizing: border-box; padding: 4px 9px; background: #b91c1c; color: #fff; border: 0; border-radius: 9px; font-size: 12px; line-height: 1.2; font-weight: 700; cursor: pointer; white-space: nowrap; }}
            .delete-button:hover {{ background: #991b1b; }}
            form {{ margin: 0; }}
        </style>
    </head>
    <body>
        <h1>Dashboard Survei GPS</h1>
        <div class="info">Login admin berhasil. Data diperbarui otomatis setiap 60 detik. Pratinjau media menggunakan tautan sementara untuk bucket privat.</div>
        <div class="info"><a href="/admin/dashboard">Dashboard Utama Admin</a> &nbsp; | &nbsp; <a href="/admin/pendaftaran-anggota">Verifikasi Pendaftaran Anggota</a> &nbsp; | &nbsp; <a href="/admin/kelola-akun">Kelola Akses Aplikasi</a></div>
        <div class="card"><strong>Total data survei:</strong> {len(rows)} &nbsp; | &nbsp; <strong>Total file media:</strong> {len(media_rows) + len(verification_media_rows)} &nbsp; | &nbsp; <strong>Total hasil verifikasi:</strong> {len(verification_rows)}</div>
        <div class="table-wrap">
            <table>
                <thead>
                    <tr>
                        <th>ID</th><th>ID Anggota</th><th>Data Pemohon, Usaha & Bantuan</th>
                        <th>Lokasi GPS</th><th>Waktu Survei</th><th>Dokumen / Foto / Video</th><th>Aksi</th>
                    </tr>
                </thead>
                <tbody>{table_rows}</tbody>
            </table>
        </div>
    </body>
    </html>
    """
    return HTMLResponse(content=page)

# =========================================================
# DASHBOARD KHUSUS PIMPINAN: REKAP TERPISAH SURVEI DPSS DAN VERIFIKASI DPS
# =========================================================
from io import BytesIO
from fastapi.responses import StreamingResponse
from fastapi.security import HTTPBasicCredentials
from urllib.parse import urlencode


def verify_pimpinan_or_admin(credentials: HTTPBasicCredentials = Depends(admin_security)):
    """Akses dashboard untuk admin lama atau akun pimpinan dari environment."""
    allowed_pairs = [
        (os.environ.get("ADMIN_USERNAME", ""), os.environ.get("ADMIN_PASSWORD", "")),
        (os.environ.get("PIMPINAN_USERNAME", ""), os.environ.get("PIMPINAN_PASSWORD", "")),
    ]
    supplied_user = credentials.username.encode("utf-8")
    supplied_password = credentials.password.encode("utf-8")
    for expected_user, expected_password in allowed_pairs:
        if expected_user and expected_password and secrets.compare_digest(supplied_user, expected_user.encode("utf-8")) and secrets.compare_digest(supplied_password, expected_password.encode("utf-8")):
            return credentials.username
    raise HTTPException(status_code=401, detail="Username atau kata sandi admin/pimpinan salah atau belum dikonfigurasi.", headers={"WWW-Authenticate": "Basic"})



class OfficerAccessGrant(BaseModel):
    password: str = Field(min_length=10, max_length=200)


@app.get("/admin/kelola-akun/data")
def list_verified_members_for_app_access(username: str = Depends(verify_pimpinan_or_admin)):
    """Daftar anggota yang sudah disetujui; identitas tetap bersumber dari pendaftaran."""
    with get_db() as connection:
        rows = connection.execute("""
            SELECT r.id, r.full_name, r.nik, r.phone, r.org_level, r.org_position,
                   r.region, r.member_code, r.linked_member_id AS login_member_id,
                   v.village_name, o.active AS account_active
            FROM organization_member_registrations r
            LEFT JOIN villages v ON v.id = r.village_id
            LEFT JOIN officers o ON o.member_id = r.linked_member_id
            WHERE r.status = 'approved'
            ORDER BY r.org_level, r.region, r.full_name
        """).fetchall()
    return [dict(row) for row in rows]


@app.post("/admin/kelola-akun/{registration_id}/akses")
def grant_app_access_to_verified_member(
    registration_id: int, payload: OfficerAccessGrant,
    username: str = Depends(verify_pimpinan_or_admin),
):
    """Buat/aktifkan kredensial aplikasi hanya untuk anggota yang sudah disetujui."""
    with get_db() as connection:
        registration = connection.execute(
            "SELECT * FROM organization_member_registrations WHERE id=%s AND status='approved'",
            (registration_id,),
        ).fetchone()
        if not registration:
            raise HTTPException(status_code=404, detail="Anggota terverifikasi tidak ditemukan. Pastikan status pendaftarannya Disetujui.")
        if registration["org_level"] not in {"DPS", "DPSS"}:
            raise HTTPException(status_code=422, detail="Tingkat kepengurusan anggota tidak valid untuk akses aplikasi survei.")
        if registration["org_level"] == "DPSS" and not registration["village_id"]:
            raise HTTPException(status_code=422, detail="Data anggota DPSS belum memiliki desa penugasan. Perbaiki data pendaftarannya terlebih dahulu.")
        member_id = (registration.get("linked_member_id") or registration.get("member_code") or "").strip().upper()
        if not member_id:
            raise HTTPException(status_code=409, detail="ID anggota belum terbentuk. Periksa data wilayah pendaftaran.")
        linked_id = (registration.get("linked_member_id") or "").strip().upper()
        if linked_id and linked_id != member_id:
            raise HTTPException(status_code=409, detail=f"Anggota ini sudah terhubung dengan ID login {linked_id}. Gunakan ID tersebut atau periksa tautan data.")
        other_link = connection.execute(
            "SELECT id FROM organization_member_registrations WHERE linked_member_id=%s AND id<>%s",
            (member_id, registration_id),
        ).fetchone()
        if other_link:
            raise HTTPException(status_code=409, detail="ID login ini sudah ditautkan ke pendaftaran anggota lain.")
        officer = connection.execute("SELECT id FROM officers WHERE member_id=%s", (member_id,)).fetchone()
        if officer and not linked_id:
            raise HTTPException(status_code=409, detail="ID login sudah digunakan oleh akun lain. Pilih ID login yang berbeda.")
        salt = secrets.token_hex(16)
        hashed = _password_hash(payload.password, salt)
        now = datetime.now(timezone.utc).isoformat()
        connection.execute(
            "INSERT INTO members (member_id,member_name,active,created_at) VALUES (%s,%s,1,%s) "
            "ON CONFLICT (member_id) DO UPDATE SET member_name=EXCLUDED.member_name,active=1",
            (member_id, registration["full_name"], now),
        )
        if officer:
            connection.execute("""UPDATE officers SET member_name=%s,password_salt=%s,password_hash=%s,
                role=%s,region=%s,active=1,updated_at=NOW() WHERE id=%s""",
                (registration["full_name"], salt, hashed, registration["org_level"], registration["region"], officer["id"]))
            officer_id = officer["id"]
        else:
            officer_row = connection.execute("""INSERT INTO officers
                (member_id,member_name,password_salt,password_hash,role,region,active)
                VALUES (%s,%s,%s,%s,%s,%s,1) RETURNING id""",
                (member_id, registration["full_name"], salt, hashed, registration["org_level"], registration["region"])).fetchone()
            officer_id = officer_row["id"]
        connection.execute("""INSERT INTO officer_profiles
            (officer_id,nik,phone,address,birth_place,birth_date,profile_photo_path,ktp_path,updated_at)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,NOW())
            ON CONFLICT (officer_id) DO UPDATE SET nik=EXCLUDED.nik,phone=EXCLUDED.phone,
            address=EXCLUDED.address,birth_place=EXCLUDED.birth_place,birth_date=EXCLUDED.birth_date,
            profile_photo_path=CASE WHEN EXCLUDED.profile_photo_path='' THEN officer_profiles.profile_photo_path ELSE EXCLUDED.profile_photo_path END,
            ktp_path=CASE WHEN EXCLUDED.ktp_path='' THEN officer_profiles.ktp_path ELSE EXCLUDED.ktp_path END,
            updated_at=NOW()""",
            (officer_id, registration["nik"], registration["phone"], registration["address"],
             registration["birth_place"], registration["birth_date"], registration["profile_photo_path"], registration["ktp_path"]))
        connection.execute("UPDATE organization_member_registrations SET linked_member_id=%s,updated_at=NOW() WHERE id=%s", (member_id, registration_id))
        # Penugasan desa berasal dari data anggota terverifikasi. DPS tanpa desa khusus
        # mengikuti cakupan kecamatan melalui filter wilayah yang sudah ada.
        connection.execute("DELETE FROM officer_villages WHERE officer_id=%s", (officer_id,))
        if registration["village_id"]:
            connection.execute("INSERT INTO officer_villages (officer_id,village_id) VALUES (%s,%s) ON CONFLICT DO NOTHING", (officer_id, registration["village_id"]))
        connection.execute("DELETE FROM officer_sessions WHERE officer_id=%s", (officer_id,))
    return {"message": "Akses aplikasi berhasil diberikan berdasarkan data anggota terverifikasi.", "member_id": member_id, "role": registration["org_level"], "region": registration["region"]}


@app.post("/admin/kelola-akun/{registration_id}/nonaktifkan")
def revoke_app_access_from_verified_member(registration_id: int, username: str = Depends(verify_pimpinan_or_admin)):
    with get_db() as connection:
        row = connection.execute("""SELECT r.linked_member_id,o.id AS officer_id
            FROM organization_member_registrations r
            LEFT JOIN officers o ON o.member_id=r.linked_member_id
            WHERE r.id=%s AND r.status='approved'""", (registration_id,)).fetchone()
        if not row or not row["linked_member_id"] or not row["officer_id"]:
            raise HTTPException(status_code=404, detail="Akun aplikasi anggota ini belum tersedia.")
        connection.execute("UPDATE officers SET active=0,updated_at=NOW() WHERE id=%s", (row["officer_id"],))
        connection.execute("UPDATE members SET active=0 WHERE member_id=%s", (row["linked_member_id"],))
        connection.execute("DELETE FROM officer_sessions WHERE officer_id=%s", (row["officer_id"],))
    return {"message": "Akses aplikasi berhasil dicabut. Data anggota dan riwayat survei tetap disimpan."}


@app.post("/admin/kelola-akun/{registration_id}/aktifkan")
def reactivate_app_access_for_verified_member(registration_id: int, username: str = Depends(verify_pimpinan_or_admin)):
    with get_db() as connection:
        row = connection.execute("""SELECT r.linked_member_id,o.id AS officer_id
            FROM organization_member_registrations r
            LEFT JOIN officers o ON o.member_id=r.linked_member_id
            WHERE r.id=%s AND r.status='approved'""", (registration_id,)).fetchone()
        if not row or not row["linked_member_id"] or not row["officer_id"]:
            raise HTTPException(status_code=404, detail="Akun aplikasi anggota ini belum tersedia.")
        connection.execute("UPDATE officers SET active=1,updated_at=NOW() WHERE id=%s", (row["officer_id"],))
        connection.execute("UPDATE members SET active=1 WHERE member_id=%s", (row["linked_member_id"],))
    return {"message": "Akses aplikasi berhasil diaktifkan kembali."}

def _lead_filters(request):
    return {k: request.query_params.get(k, "").strip() for k in ("tanggal_mulai", "tanggal_selesai", "wilayah", "petugas")}


def _lead_where(filters, alias="s"):
    clauses, params = [], []
    if filters.get("tanggal_mulai"):
        clauses.append(f"NULLIF({alias}.surveyed_at, '')::timestamptz::date >= %s::date")
        params.append(filters["tanggal_mulai"])
    if filters.get("tanggal_selesai"):
        clauses.append(f"NULLIF({alias}.surveyed_at, '')::timestamptz::date <= %s::date")
        params.append(filters["tanggal_selesai"])
    if filters.get("wilayah"):
        clauses.append(f"COALESCE({alias}.village, '') ILIKE %s")
        params.append("%" + filters["wilayah"] + "%")
    if filters.get("petugas"):
        clauses.append(f"(COALESCE({alias}.member_id, '') ILIKE %s OR COALESCE(o.member_name, '') ILIKE %s)")
        params.extend(["%" + filters["petugas"] + "%", "%" + filters["petugas"] + "%"])
    return (" WHERE " + " AND ".join(clauses) if clauses else ""), tuple(params)


def _lead_survey_rows(filters):
    where_sql, params = _lead_where(filters, "s")
    sql = """SELECT s.id AS survey_id, s.member_id, COALESCE(o.member_name,s.member_id) AS officer_name,
        s.surveyed_at, s.received_at, s.village, s.applicant_name, s.applicant_phone, s.applicant_address,
        s.business_name, s.business_type, s.business_location, s.capital_amount, s.turnover_amount,
        s.net_profit_amount, s.worker_count, s.assistance_types, s.assistance_amount, s.notes,
        s.latitude, s.longitude, s.accuracy
        FROM surveys s LEFT JOIN officers o ON o.member_id=s.member_id""" + where_sql + " ORDER BY s.id DESC"
    with get_db() as connection:
        return [dict(r) for r in connection.execute(sql, params).fetchall()]


def _lead_verification_rows(filters):
    where_sql, params = _lead_where(filters, "s")
    sql = """SELECT s.id AS survey_id, s.member_id, COALESCE(o.member_name,s.member_id) AS officer_name,
        s.surveyed_at, s.village, s.applicant_name, s.business_name, s.business_type,
        v.decision, v.verifier_id, v.verifier_name, v.verifier_region, v.verified_at,
        v.business_condition, v.actual_business_activity, v.existing_assets, v.needs,
        v.working_capital_details, v.verified_turnover, v.verified_profit, v.verified_worker_count,
        v.data_match, v.discrepancy_notes, v.verification_notes, v.rejection_reason,
        v.latitude AS verification_latitude, v.longitude AS verification_longitude, v.accuracy AS verification_accuracy,
        v.total_needs_amount
        FROM surveys s LEFT JOIN officers o ON o.member_id=s.member_id
        JOIN survey_verifications v ON v.survey_id=s.id""" + where_sql + " ORDER BY v.verified_at DESC NULLS LAST, s.id DESC"
    with get_db() as connection:
        return [dict(r) for r in connection.execute(sql, params).fetchall()]


def _lead_media(survey_ids, verification=False, include_identity=False):
    if not survey_ids:
        return {}
    table = "verification_media" if verification else "survey_media"
    with get_db() as connection:
        rows = connection.execute(
            f"SELECT survey_id, object_path, content_type, original_filename, file_size, storage_bucket" +
            (", NULL AS media_category" if verification else ", COALESCE(media_category, 'business_photo') AS media_category") +
            f" FROM {table} WHERE survey_id = ANY(%s) ORDER BY id", (survey_ids,)
        ).fetchall()
    grouped = {}
    for item in rows:
        item = dict(item)
        category = str(item.get("media_category") or "business_photo").lower()
        # Dokumen identitas hanya boleh muncul pada halaman dashboard survei.
        # Tetap dikecualikan dari laporan verifikasi dan ekspor.
        is_identity = any(word in category for word in ("ktp", "kk", "identity", "face", "selfie", "document", "dokumen"))
        if is_identity and (verification or not include_identity):
            continue
        url = create_supabase_signed_url(item["object_path"], bucket=item.get("storage_bucket") or "survey-media")
        item["signed_url"] = url
        grouped.setdefault(item["survey_id"], []).append(item)
    return grouped


def _lead_e(value):
    return escape("" if value is None else str(value), quote=True)


def _lead_qs(filters):
    return urlencode({k: v for k, v in filters.items() if v})


def _lead_page(title, username, content):
    return f"""<!doctype html><html lang='id'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>
<title>{_lead_e(title)} | Ikatan Persaudaraan Triple 555</title><style>
:root{{--maroon:#650b18;--deep:#17070b;--gold:#e0ad35;--paper:#f4f0eb;--ink:#f7eee5;--line:#503337;--muted:#b8a8a0}}
*{{box-sizing:border-box}}body{{margin:0;background:radial-gradient(ellipse at top right,#35151b 0,#16090c 55%,#100709 100%);color:var(--ink);font:14px/1.5 Inter,system-ui,Arial,sans-serif}}
header{{background:linear-gradient(110deg,#19080c,#500813 55%,#26090e);color:white;padding:16px 26px;border-bottom:2px solid var(--gold);box-shadow:0 8px 30px #0006}}.brand{{display:flex;align-items:center;gap:14px}}.brand img{{width:64px;height:64px;object-fit:contain;border-radius:50%;background:#2b1114;padding:3px;border:2px solid var(--gold)}}header h1{{margin:0;font-size:clamp(22px,3vw,30px);color:#f7d77f}}header p{{margin:4px 0 0;color:#e8d8c7}}.user{{margin-top:5px;color:#d9c3a7;font-size:12px}}
.shell{{display:grid;grid-template-columns:230px minmax(0,1fr);gap:22px;max-width:1740px;margin:auto;padding:22px 24px 40px}}.nav{{display:flex;flex-direction:column;gap:7px;background:linear-gradient(180deg,#26090f,#14080b);border:1px solid #6e3a3d;border-radius:16px;padding:15px 11px;align-self:start;min-height:calc(100vh - 160px);box-shadow:0 12px 28px #0005}}.nav:before{{content:'MENU PIMPINAN';font-size:10px;letter-spacing:1.8px;color:#c9a65a;padding:8px 12px}}
a.btn,button{{display:inline-block;border:1px solid transparent;border-radius:9px;background:linear-gradient(110deg,#e9bd51,#b87a16);color:#24130a;padding:10px 13px;font:inherit;font-weight:750;text-decoration:none;cursor:pointer}}a.gold{{background:linear-gradient(110deg,#f0c85e,#bd821b);color:#24130a}}a.light{{background:#321b20;color:#f3dfb1;border:1px solid #644349}}
/* Layout media laporan: berada di bawah data teks dan memenuhi lebar kolom laporan. */
.lead-report-panel{{min-width:0;overflow:hidden}}
.lead-report-tablewrap{{overflow-x:auto}}
.lead-report-table{{width:100%;min-width:900px;table-layout:auto}}
.lead-report-table .lead-report-data-row td{{background:#211417}}
.lead-report-table .lead-report-media-row td{{padding:13px 12px 17px;background:#180d10;border-bottom:2px solid #765047}}
.lead-media-gallery{{display:flex;flex-wrap:wrap;align-items:flex-start;gap:12px;width:100%;max-width:100%}}
.lead-media-section{{flex:1 1 100%;min-width:0;margin:0;padding:0}}
.lead-media-heading{{margin:0 0 8px;color:#e7c46e;font-size:12px;letter-spacing:.3px}}
.lead-identity-grid,.lead-photo-grid,.lead-video-grid,.lead-media-unavailable-list{{display:flex;flex-wrap:wrap;align-items:flex-start;gap:10px;width:100%;max-width:100%}}
.lead-identity-item,.lead-photo-item,.lead-video-item,.lead-media-unavailable{{flex:0 1 150px;min-width:0;max-width:150px;padding:7px;border:1px solid #5b3939;border-radius:9px;background:#26191b;overflow:hidden}}
.lead-media-label{{font-size:11px;font-weight:750;color:#e7c46e;margin-bottom:5px;overflow-wrap:anywhere}}
.lead-identity-image,.lead-photo-image{{display:block;width:100%;height:105px;object-fit:contain;background:#10090b;border-radius:6px}}
.lead-video-item video{{display:block;width:100%;height:105px;object-fit:contain;background:#10090b;border-radius:6px}}
.lead-media-name{{font-size:10px;color:#c9b8ad;margin-top:5px;overflow-wrap:anywhere}}
.lead-media-empty{{font-size:12px;color:#b8a8a0;padding:7px 0}}
.lead-media-unavailable p{{font-size:11px;margin:4px 0;color:#d8c9c0;overflow-wrap:anywhere}}
main{{min-width:0;padding:0}}.panel,.metric{{background:linear-gradient(145deg,#26191b,#1c1114);border:1px solid #5b3939;border-radius:13px;padding:16px;margin:12px 0;box-shadow:0 8px 20px #0003;color:#f7eee5}}.metrics{{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:12px}}.metric strong{{display:block;font-size:27px;color:#f2c34f}}.metric span{{color:#c9b8ad}}
.filters{{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:10px}}label{{display:block;font-size:12px;font-weight:700;margin-bottom:5px;color:#e8cc8b}}input,select,textarea{{width:100%;min-height:40px;padding:9px;border:1px solid #624449;border-radius:8px;background:#180e11;color:#f8eee3;font:inherit}}.actions{{display:flex;gap:8px;flex-wrap:wrap;margin-top:12px}}.tablewrap{{overflow-x:auto}}table{{width:100%;border-collapse:collapse;min-width:850px}}th,td{{padding:10px;border-bottom:1px solid #493235;text-align:left;vertical-align:top}}th{{background:#301b20;color:#e7c46e;font-size:12px;white-space:nowrap}}td{{color:#eee3d9}}td small{{color:#b8a8a0}}.media{{display:flex;gap:8px;flex-wrap:wrap;max-width:300px}}.media img{{width:100px;height:78px;object-fit:cover;border-radius:7px;border:1px solid #624449}}.media video{{width:180px;max-width:100%;border-radius:7px}}.muted{{color:#b8a8a0}}
.cards{{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:14px}}.report-card{{background:linear-gradient(145deg,#26191b,#1c1114);border:1px solid #5b3939;border-top:3px solid var(--gold);border-radius:14px;padding:20px}}.report-card h2{{margin:0 0 8px;color:#f1cb72}}
.people-grid{{display:grid;grid-template-columns:repeat(auto-fill,minmax(270px,1fr));gap:12px}}.person-card{{display:flex;gap:14px;align-items:center;background:#211518;border:1px solid #5b3939;border-left:4px solid var(--gold);border-radius:13px;padding:14px}}.person-photo{{width:82px;height:82px;flex:0 0 82px;border-radius:12px;object-fit:cover;background:#302225}}.no-photo{{display:flex;align-items:center;justify-content:center;text-align:center;color:#b8a8a0;font-size:11px}}.person-info h3{{margin:0 0 3px;color:#f1cb72}}.person-info p{{margin:6px 0}}.status{{display:inline-block;padding:4px 9px;border-radius:20px;font-size:11px;font-weight:800}}.status.active{{background:#174d32;color:#a7efbf}}.status.inactive{{background:#604316;color:#ffdf8a}}footer{{text-align:center;color:#988780;font-size:12px;padding:20px}}
@media(max-width:900px){{.shell{{grid-template-columns:1fr;padding:14px}}.nav{{min-height:0;flex-direction:row;overflow-x:auto;white-space:nowrap}}.nav:before{{display:none}}.nav a{{flex:0 0 auto}}}}@media(max-width:800px){{.metrics{{grid-template-columns:repeat(2,minmax(0,1fr))}}.filters{{grid-template-columns:repeat(2,minmax(0,1fr))}}}}@media(max-width:520px){{.cards,.filters{{grid-template-columns:1fr}}header{{padding:18px 16px}}.shell{{padding:10px}}}}
</style></head><body><header><div class='brand'><img src='/logo.png' alt='Logo Triple 555'><div><h1>555 · EXECUTIVE COMMAND CENTER</h1><p>Ikatan Persaudaraan Triple 555 · Sistem Survei Bantuan UMKM</p><div class='user'>Masuk sebagai: {_lead_e(username)} · Pimpinan/Admin</div></div></div></header><main><div class='shell'><nav class='nav'><a class='btn light' href='/pimpinan'>▦ &nbsp; Ringkasan</a><a class='btn light' href='/pimpinan/survei'>▤ &nbsp; Laporan Survei DPSS</a><a class='btn light' href='/pimpinan/verifikasi'>✓ &nbsp; Laporan Verifikasi DPS</a><a class='btn light' href='/pimpinan/peta'>⌖ &nbsp; Peta Lokasi GPS</a><a class='btn light' href='/pimpinan/personel'>♙ &nbsp; Data Personel</a><a class='btn light' href='/admin/kelola-akun'>⚿ &nbsp; Kelola Akses</a></nav><main>{content}</main></div><footer>© 2026 Ikatan Persaudaraan Triple 555 · Data digunakan sesuai kewenangan organisasi.</footer></body></html>"""



# =========================================================
# MODUL PROFIL PERSONEL DPS / DPSS
# =========================================================

def _officer_private_bucket():
    return os.environ.get("SUPABASE_OFFICER_BUCKET", "officer-private").strip() or "officer-private"


def _upload_officer_private(file: UploadFile | None, member_id: str, kind: str) -> str:
    if file is None or not file.filename:
        return ""
    content_type = (file.content_type or "").lower().split(";", 1)[0].strip()
    allowed_types = {"image/jpeg": ".jpg", "image/png": ".png", "image/webp": ".webp"}
    if content_type not in allowed_types:
        raise HTTPException(status_code=400, detail=f"{kind}: format harus JPG, PNG, atau WEBP.")
    data = file.file.read()
    if not data:
        raise HTTPException(status_code=400, detail=f"{kind}: berkas kosong.")
    if len(data) > 5 * 1024 * 1024:
        raise HTTPException(status_code=413, detail=f"{kind}: ukuran maksimal 5 MB.")
    supabase_url = os.environ.get("SUPABASE_URL", "").strip().rstrip("/")
    supabase_key = os.environ.get("SUPABASE_SECRET_KEY", "").strip()
    if not supabase_url or not supabase_key:
        raise HTTPException(status_code=503, detail="Penyimpanan dokumen personel belum dikonfigurasi.")
    ext = allowed_types[content_type]
    safe_id = re.sub(r"[^A-Za-z0-9_-]", "_", member_id)
    object_path = f"{safe_id}/{kind}-{uuid.uuid4().hex}{ext}"
    bucket = _officer_private_bucket()
    upload_url = f"{supabase_url}/storage/v1/object/{quote(bucket, safe='')}/{quote(object_path, safe='/')}"
    req = urllib.request.Request(
        upload_url, data=data,
        headers={
            "Authorization": f"Bearer {supabase_key}",
            "apikey": supabase_key,
            "Content-Type": content_type or "application/octet-stream",
            "x-upsert": "false",
        }, method="POST"
    )
    try:
        with urllib.request.urlopen(req, timeout=30) as response:
            response.read()
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", "ignore")[:500]
        raise HTTPException(status_code=502, detail=f"Gagal mengunggah {kind} ke penyimpanan privat: {detail}")
    except (urllib.error.URLError, TimeoutError) as exc:
        raise HTTPException(status_code=502, detail=f"Koneksi penyimpanan {kind} gagal: {exc}")
    return object_path


def _officer_profile_rows(search_text=""):
    with get_db() as connection:
        return connection.execute("""
            SELECT o.id AS officer_id, o.member_id, o.member_name, o.role, o.region, o.active,
                   COALESCE(p.nik,'') AS nik, COALESCE(p.phone,'') AS phone,
                   COALESCE(p.address,'') AS address,
                   COALESCE(p.birth_place,'') AS birth_place, p.birth_date,
                   COALESCE(p.profile_photo_path,'') AS profile_photo_path,
                   COALESCE(p.ktp_path,'') AS ktp_path,
                   COALESCE(string_agg(DISTINCT v.village_name, ', '), '') AS villages
            FROM officers o
            LEFT JOIN officer_profiles p ON p.officer_id=o.id
            LEFT JOIN officer_villages ov ON ov.officer_id=o.id
            LEFT JOIN villages v ON v.id=ov.village_id
            WHERE (%s = '' OR o.member_id ILIKE %s OR o.member_name ILIKE %s
                   OR COALESCE(p.phone,'') ILIKE %s OR COALESCE(o.region,'') ILIKE %s
                   OR COALESCE(v.village_name,'') ILIKE %s)
            GROUP BY o.id, p.officer_id
            ORDER BY CASE WHEN o.role='DPS' THEN 0 ELSE 1 END, o.member_name
        """, (search_text, f"%{search_text}%", f"%{search_text}%",
              f"%{search_text}%", f"%{search_text}%", f"%{search_text}%")).fetchall()


@app.get("/admin/data-anggota/{registration_id}/dokumen-identitas")
def admin_view_member_identity_document(registration_id: int, username: str = Depends(verify_admin)):
    """Buat signed URL baru saat admin membuka dokumen identitas anggota."""
    with get_db() as connection:
        row = connection.execute("SELECT ktp_path FROM organization_member_registrations WHERE id=%s AND status='approved'", (registration_id,)).fetchone()
    if not row or not row.get("ktp_path"):
        raise HTTPException(status_code=404, detail="Dokumen identitas anggota tidak ditemukan.")
    signed_url = create_supabase_signed_url(row["ktp_path"], expires_in=300, bucket=_officer_private_bucket())
    if not signed_url:
        raise HTTPException(status_code=502, detail="Dokumen identitas tidak dapat dibuka dari penyimpanan privat.")
    return RedirectResponse(url=signed_url, status_code=302)


@app.get("/admin/data-anggota", response_class=HTMLResponse)
def admin_data_anggota(q: str = "", tingkat: str = "SEMUA", username: str = Depends(verify_admin)):
    """Direktori seluruh anggota yang telah disetujui, terpisah dari daftar akun aplikasi."""
    tingkat = tingkat.strip().upper()
    if tingkat not in {"SEMUA", "DPS", "DPSS"}:
        tingkat = "SEMUA"
    term = q.strip()
    with get_db() as connection:
        rows = connection.execute("""SELECT r.*,v.village_name
            FROM organization_member_registrations r
            LEFT JOIN villages v ON v.id=r.village_id
            WHERE r.status='approved' AND (%s='SEMUA' OR r.org_level=%s)
              AND (%s='' OR r.full_name ILIKE %s OR r.member_code ILIKE %s OR r.nik ILIKE %s
                   OR r.phone ILIKE %s OR r.org_position ILIKE %s OR r.region ILIKE %s
                   OR COALESCE(v.village_name,'') ILIKE %s)
            ORDER BY r.org_level,r.region,v.village_name,r.full_name""",
            (tingkat,tingkat,term,*([f"%{term}%"]*7))).fetchall()
    cards=[]
    for r in rows:
        photo_url=create_supabase_signed_url(r['profile_photo_path'],1800,_officer_private_bucket()) if r['profile_photo_path'] else None
        has_ktp=bool(r['ktp_path'])
        photo_html=f"<img class='photo' src='{_lead_e(photo_url)}' alt='Foto {_lead_e(r['full_name'])}'>" if photo_url else "<div class='photo no-photo'>Foto belum tersedia</div>"
        doc_html=f"<a class='btn doc' href='/admin/data-anggota/{int(r['id'])}/dokumen-identitas' target='_blank' rel='noopener'>Lihat Dokumen Identitas</a>" if has_ktp else "<span class='muted'>Dokumen identitas belum tersedia</span>"
        birth_date=r['birth_date'].isoformat() if r['birth_date'] else '-'
        village=r['village_name'] or ('Seluruh Desa' if r['org_level']=='DPS' else '-')
        cards.append(f"""<article class='member-card'>{photo_html}<div class='member-data'><div class='topline'><div><h2>{_lead_e(r['full_name'])}</h2><span class='id'>ID Member: {_lead_e(r['member_code'] or '-')}</span></div><span class='level'>{_lead_e(r['org_level'])}</span></div>
        <div class='details'><div><b>NIK</b><span>{_lead_e(r['nik'])}</span></div><div><b>Jenis Kelamin</b><span>{_lead_e(r['gender'] or '-')}</span></div><div><b>Tempat, Tanggal Lahir</b><span>{_lead_e(r['birth_place'] or '-')}, {_lead_e(birth_date)}</span></div><div><b>Nomor HP / WhatsApp</b><span>{_lead_e(r['phone'] or '-')}</span></div><div><b>Email</b><span>{_lead_e(r['email'] or '-')}</span></div><div><b>Jabatan</b><span>{_lead_e(r['org_position'])}</span></div><div><b>Wilayah</b><span>{_lead_e(r['region'] or '-')} · {_lead_e(village)}</span></div><div class='wide'><b>Alamat</b><span>{_lead_e(r['address'] or '-')}</span></div></div><div class='doc-row'>{doc_html}<a class='btn light' href='/admin/pendaftaran-anggota/{int(r['id'])}/edit'>Edit Data</a></div></div></article>""")
    cards_html=''.join(cards) or "<div class='empty'>Belum ada anggota terverifikasi yang sesuai dengan pencarian.</div>"
    html=f"""<!doctype html><html lang='id'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>Daftar Seluruh Anggota | Triple 555</title>
    <style>*{{box-sizing:border-box}}body{{margin:0;background:#f5f1eb;color:#292323;font:14px/1.5 system-ui,Arial,sans-serif}}header{{background:linear-gradient(120deg,#31090e,#5b1018);color:#fff;padding:22px;border-bottom:4px solid #d8b15a}}main{{max-width:1250px;margin:auto;padding:18px}}.nav,.filters,.doc-row{{display:flex;gap:9px;flex-wrap:wrap;align-items:center}}.nav{{margin-bottom:15px}}.btn,button{{display:inline-block;background:#5b1018;color:#fff;text-decoration:none;border:0;border-radius:8px;padding:9px 12px;font-weight:700;cursor:pointer}}.btn.light{{background:#eee8e1;color:#382e2e}}.panel{{background:#fff;padding:16px;border-radius:13px;margin-bottom:14px}}input,select{{padding:10px;border:1px solid #d6cec6;border-radius:8px;font:inherit}}.filters input{{flex:1;min-width:220px}}.count{{color:#6c625b;margin:10px 0}}.grid{{display:grid;gap:13px}}.member-card{{display:grid;grid-template-columns:145px 1fr;gap:17px;background:#fff;border:1px solid #e7ded4;border-radius:14px;padding:16px;box-shadow:0 3px 12px #32100b08}}.photo{{width:145px;height:170px;border-radius:10px;object-fit:cover;background:#f0ece7}}.no-photo{{display:grid;place-items:center;text-align:center;color:#777;padding:10px}}.topline{{display:flex;justify-content:space-between;gap:12px;align-items:flex-start}}h2{{margin:0 0 4px;color:#5b1018;font-size:19px}}.id{{font-weight:800;color:#775313}}.level{{background:#f5ead5;color:#5b1018;border-radius:20px;padding:5px 11px;font-weight:800}}.details{{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:10px;margin:14px 0}}.details div{{min-width:0}}.details b,.details span{{display:block}}.details b{{font-size:11px;color:#786e65;text-transform:uppercase}}.details span{{overflow-wrap:anywhere}}.details .wide{{grid-column:1/-1}}.doc-row{{border-top:1px solid #eee5dc;padding-top:12px}}.doc{{background:#8b641c}}.muted,.empty{{color:#766d65}}.empty{{background:#fff;padding:28px;text-align:center;border-radius:12px}}@media(max-width:700px){{main{{padding:11px}}.member-card{{grid-template-columns:90px 1fr;padding:11px;gap:11px}}.photo{{width:90px;height:115px}}.details{{grid-template-columns:repeat(2,minmax(0,1fr));gap:8px}}.topline{{flex-direction:column}}}}</style></head><body><header><h1 style='margin:0'>Daftar Seluruh Anggota DPS dan DPSS</h1><div>Data anggota terverifikasi · Administrasi Organisasi Triple 555</div></header><main>
    <nav class='nav'><a class='btn light' href='/admin/dashboard'>Dashboard Utama Admin</a><a class='btn light' href='/admin/kelola-akun'>Kelola Akses Aplikasi</a><a class='btn light' href='/admin/pendaftaran-anggota'>Verifikasi Pendaftaran</a></nav>
    <section class='panel'><form class='filters' method='get'><input name='q' value='{_lead_e(term)}' placeholder='Cari nama, ID Member, NIK, HP, jabatan, atau wilayah'><select name='tingkat'><option value='SEMUA' {'selected' if tingkat=='SEMUA' else ''}>Semua tingkat</option><option value='DPS' {'selected' if tingkat=='DPS' else ''}>DPS</option><option value='DPSS' {'selected' if tingkat=='DPSS' else ''}>DPSS</option></select><button type='submit'>Cari / Filter</button><a class='btn light' href='/admin/data-anggota'>Reset</a></form><p class='count'>Jumlah anggota terverifikasi: <b>{len(rows)}</b></p></section><section class='grid'>{cards_html}</section></main></body></html>"""
    return HTMLResponse(html)


@app.get("/admin/data-personel", response_class=HTMLResponse)
def admin_data_personel(request: Request, q: str = "", username: str = Depends(verify_admin)):
    rows = _officer_profile_rows(q.strip())
    body = []
    for r in rows:
        status = "Aktif" if r["active"] == 1 else "Nonaktif"
        photo = create_supabase_signed_url(r["profile_photo_path"], bucket=_officer_private_bucket()) if r["profile_photo_path"] else None
        photo_html = f"<img class='avatar' src='{_lead_e(photo)}'>" if photo else "<div class='avatar blank'>Foto<br>belum ada</div>"
        territory = r["villages"] or r["region"] or "-"
        body.append(f"""<tr><td>{photo_html}</td><td><b>{_lead_e(r['member_name'])}</b><br><small>{_lead_e(r['member_id'])} · {_lead_e(r['role'])}</small></td><td>{_lead_e(r['phone'] or '-')}</td><td>{_lead_e(territory)}</td><td>{status}</td><td><div style='display:flex;gap:7px;flex-wrap:wrap'><a class='btn' href='/admin/data-personel/{r['officer_id']}'>Isi / Edit Data</a><form method='post' action='/admin/data-personel/{r['officer_id']}/hapus' onsubmit="return confirm('Hapus data personel {_lead_e(r['member_name'])} dengan ID {_lead_e(r['member_id'])}? Akun dan pendaftaran terkait akan dihapus. Data survei/verifikasi tidak akan dihapus.');"><button class='btn danger' type='submit'>Hapus</button></form></div></td></tr>""")
    if not body:
        body = ["<tr><td colspan='6'>Belum ada akun petugas.</td></tr>"]
    deleted_notice = "<div style='background:#eaf7ef;color:#14532d;padding:11px 14px;border-radius:9px;margin:12px 0'>Data personel dan pendaftaran terkait berhasil dihapus.</div>" if request.query_params.get("deleted") == "1" else ""
    html = f"""<!doctype html><html lang='id'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>Data Personel DPS/DPSS</title>
    <style>body{{font:14px system-ui;background:#f6f3ee;margin:0;color:#292323}}header{{background:linear-gradient(120deg,#31090e,#5b1018);color:white;padding:22px}}main{{max-width:1200px;margin:auto;padding:20px}}.panel{{background:white;padding:18px;border-radius:13px;margin:12px 0;overflow:auto}}table{{width:100%;border-collapse:collapse;min-width:800px}}td,th{{padding:10px;border-bottom:1px solid #e5e0d8;text-align:left}}th{{color:#5b1018;background:#f8f4ec}}input{{padding:10px;border:1px solid #ddd;border-radius:8px}}.btn,button{{background:#5b1018;color:white;border:0;border-radius:8px;padding:9px 12px;text-decoration:none;font-weight:bold;cursor:pointer}}.danger{{background:#a32020}}form{{margin:0}}.avatar{{width:58px;height:58px;border-radius:50%;object-fit:cover;background:#eee}}.blank{{display:flex;align-items:center;justify-content:center;text-align:center;font-size:10px;color:#777}}small{{color:#777}}</style></head><body><header><h1>Data Personel DPS dan DPSS</h1><p>Dashboard pendataan identitas petugas · Masuk sebagai {_lead_e(username)}</p></header><main><p><a class='btn' href='/admin/dashboard'>Dashboard Utama Admin</a> <a class='btn' href='/admin/pendaftaran-anggota'>Verifikasi Pendaftaran Mandiri</a> <a class='btn' href='/pendaftaran-anggota' target='_blank' rel='noopener'>Buka Formulir Anggota</a></p>{deleted_notice}<section class='panel'><form method='get'><input name='q' value='{_lead_e(q)}' placeholder='Cari nama, ID, nomor HP, wilayah' style='width:min(440px,75%)'><button>Cari</button></form><h2>Daftar akun petugas ({len(rows)})</h2><table><thead><tr><th>Foto</th><th>Nama / ID / Peran</th><th>Nomor HP</th><th>Wilayah / Desa</th><th>Status akun</th><th>Kelola</th></tr></thead><tbody>{''.join(body)}</tbody></table></section></main></body></html>"""
    return HTMLResponse(html)


@app.get("/admin/data-personel/{officer_id}", response_class=HTMLResponse)
def admin_edit_data_personel(officer_id: int, request: Request, username: str = Depends(verify_admin)):
    with get_db() as connection:
        row = connection.execute("""
            SELECT o.id AS officer_id,o.member_id,o.member_name,o.role,o.region,o.active,
                   p.nik,p.phone,p.address,p.birth_place,p.birth_date,p.profile_photo_path,p.ktp_path
            FROM officers o LEFT JOIN officer_profiles p ON p.officer_id=o.id WHERE o.id=%s
        """, (officer_id,)).fetchone()
        village_rows = connection.execute(
            "SELECT id, village_name, district FROM villages WHERE active=1 ORDER BY district, village_name"
        ).fetchall()
        assigned_rows = connection.execute(
            "SELECT village_id FROM officer_villages WHERE officer_id=%s", (officer_id,)
        ).fetchall()
    if not row:
        raise HTTPException(status_code=404, detail="Akun petugas tidak ditemukan.")
    assigned_ids = {v["village_id"] for v in assigned_rows}
    village_controls = "".join(
        f"<label class='village'><input type='checkbox' name='village_ids' value='{v['id']}' {'checked' if v['id'] in assigned_ids else ''}><span>{_lead_e(v['village_name'])} <small>· {_lead_e(v['district'])}</small></span></label>"
        for v in village_rows
    ) or "<p>Belum ada data desa aktif. Tambahkan melalui pengelolaan wilayah terlebih dahulu.</p>"
    photo = create_supabase_signed_url(row["profile_photo_path"], bucket=_officer_private_bucket()) if row["profile_photo_path"] else None
    ktp = create_supabase_signed_url(row["ktp_path"], bucket=_officer_private_bucket()) if row["ktp_path"] else None
    saved = "<div class='notice'>Perubahan profil berhasil disimpan.</div>" if request.query_params.get("saved") == "1" else ""
    html = f"""<!doctype html><html lang='id'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>Edit Profil Personel</title><style>
    body{{font:14px system-ui;background:#f6f3ee;margin:0;color:#292323}}header{{background:#31090e;color:white;padding:20px}}main{{max-width:900px;margin:auto;padding:20px}}form,.panel{{background:white;padding:20px;border-radius:12px;margin-bottom:14px}}label{{display:block;font-weight:bold;margin:12px 0 5px}}input,textarea,select{{width:100%;padding:10px;border:1px solid #ddd;border-radius:8px;box-sizing:border-box}}input[type=checkbox]{{width:18px}}.grid{{display:grid;grid-template-columns:1fr 1fr;gap:12px}}button,.btn{{display:inline-block;background:#5b1018;color:white;padding:10px 14px;border:0;border-radius:8px;text-decoration:none;margin-top:16px;cursor:pointer}}img{{max-width:180px;max-height:180px;object-fit:contain;margin:8px}}small,.muted{{color:#777}}.village-list{{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:8px;margin-top:8px}}.village{{display:flex;align-items:center;gap:8px;border:1px solid #e4ded7;border-radius:8px;padding:8px;margin:0}}.notice{{background:#eaf7ef;color:#14532d;padding:10px;border-radius:8px;margin-bottom:12px}}@media(max-width:650px){{.grid,.village-list{{grid-template-columns:1fr}}}}
    </style></head><body><header><h1>Pengelolaan Data Personel</h1><div>{_lead_e(row['member_name'])} · {_lead_e(row['member_id'])} · {_lead_e(row['role'])}</div></header><main><p><a class='btn' href='/admin/data-personel'>Kembali ke daftar personel</a></p>{saved}<form method='post' enctype='multipart/form-data' action='/admin/data-personel/{officer_id}'>
    <h2>1. Data akun dan penugasan</h2><div class='grid'><div><label>Nama lengkap</label><input name='member_name' value='{_lead_e(row['member_name'])}' maxlength='100' required></div><div><label>ID anggota / petugas</label><input value='{_lead_e(row['member_id'])}' disabled></div>
    <div><label>Peran</label><select name='role'><option value='DPSS' {'selected' if row['role']=='DPSS' else ''}>DPSS · Desa/Kelurahan</option><option value='DPS' {'selected' if row['role']=='DPS' else ''}>DPS · Kecamatan</option></select></div>
    <div><label>Wilayah / Kecamatan</label><input name='region' value='{_lead_e(row['region'] or '')}' maxlength='150' placeholder='Contoh: Tambang'></div>
    <div><label>Status akun</label><select name='active'><option value='1' {'selected' if row['active']==1 else ''}>Aktif</option><option value='0' {'selected' if row['active']!=1 else ''}>Nonaktif</option></select></div></div>
    <label>Desa penugasan (khusus DPSS)</label><div class='village-list'>{village_controls}</div><p class='muted'>Pilih desa yang menjadi wilayah tugas petugas DPSS. Untuk akun DPS, daftar desa tidak digunakan.</p>
    <h2 style='margin-top:28px'>2. Identitas dan kontak</h2><div class='grid'><div><label>NIK</label><input name='nik' value='{_lead_e(row['nik'] or '')}' maxlength='30' inputmode='numeric'></div><div><label>Nomor HP</label><input name='phone' value='{_lead_e(row['phone'] or '')}' maxlength='30' inputmode='tel'></div><div><label>Tempat lahir</label><input name='birth_place' value='{_lead_e(row['birth_place'] or '')}' maxlength='100'></div><div><label>Tanggal lahir</label><input type='date' name='birth_date' value='{_lead_e(row['birth_date'].isoformat() if row['birth_date'] else '')}'></div></div>
    <label>Alamat lengkap</label><textarea name='address' rows='3'>{_lead_e(row['address'] or '')}</textarea>
    <h2 style='margin-top:28px'>3. Dokumen personel</h2><label>Foto profil (JPG/PNG/WEBP, maksimal 5 MB)</label><input type='file' name='profile_photo' accept='image/jpeg,image/png,image/webp'>{f'<img src="{_lead_e(photo)}" alt="Foto profil saat ini">' if photo else '<small>Belum ada foto profil.</small>'}
    <label>Foto / scan KTP (JPG/PNG/WEBP, maksimal 5 MB)</label><input type='file' name='ktp_photo' accept='image/jpeg,image/png,image/webp'>{f'<img src="{_lead_e(ktp)}" alt="KTP tersimpan">' if ktp else '<small>Belum ada dokumen KTP.</small>'}
    <p><button type='submit'>Simpan Seluruh Data Personel</button></p></form></main></body></html>"""
    return HTMLResponse(html)


@app.post("/admin/data-personel/{officer_id}")
def admin_save_data_personel(
    officer_id: int,
    member_name: str = Form(...), role: str = Form(...), region: str = Form(""),
    active: int = Form(1), village_ids: list[int] = Form([]),
    nik: str = Form(""), phone: str = Form(""), address: str = Form(""),
    birth_place: str = Form(""), birth_date: str = Form(""),
    profile_photo: UploadFile | None = File(None), ktp_photo: UploadFile | None = File(None),
    username: str = Depends(verify_admin),
):
    member_name = member_name.strip()
    role = role.strip().upper()
    region = region.strip()
    if len(member_name) < 2:
        raise HTTPException(status_code=422, detail="Nama lengkap minimal 2 karakter.")
    if role not in {"DPS", "DPSS"}:
        raise HTTPException(status_code=422, detail="Peran hanya boleh DPS atau DPSS.")
    if active not in (0, 1):
        raise HTTPException(status_code=422, detail="Status akun tidak valid.")
    if birth_date.strip():
        try:
            datetime.strptime(birth_date.strip(), "%Y-%m-%d")
        except ValueError as exc:
            raise HTTPException(status_code=422, detail="Tanggal lahir tidak valid.") from exc
    with get_db() as connection:
        officer = connection.execute("SELECT id,member_id FROM officers WHERE id=%s", (officer_id,)).fetchone()
        if not officer:
            raise HTTPException(status_code=404, detail="Akun petugas tidak ditemukan.")
        if role == "DPSS" and village_ids:
            valid_villages = connection.execute(
                "SELECT id FROM villages WHERE active=1 AND id=ANY(%s)", (village_ids,)
            ).fetchall()
            if len({v["id"] for v in valid_villages}) != len(set(village_ids)):
                raise HTTPException(status_code=422, detail="Ada desa penugasan yang tidak valid atau tidak aktif.")
    photo_path = _upload_officer_private(profile_photo, officer["member_id"], "profile") if profile_photo and profile_photo.filename else ""
    ktp_path = _upload_officer_private(ktp_photo, officer["member_id"], "ktp") if ktp_photo and ktp_photo.filename else ""
    with get_db() as connection:
        connection.execute(
            "UPDATE officers SET member_name=%s, role=%s, region=%s, active=%s, updated_at=NOW() WHERE id=%s",
            (member_name, role, region, active, officer_id),
        )
        connection.execute(
            "UPDATE members SET member_name=%s, active=%s WHERE member_id=%s",
            (member_name, active, officer["member_id"]),
        )
        if role == "DPSS":
            connection.execute("DELETE FROM officer_villages WHERE officer_id=%s", (officer_id,))
            for village_id in sorted(set(village_ids)):
                connection.execute(
                    "INSERT INTO officer_villages (officer_id,village_id) VALUES (%s,%s) ON CONFLICT DO NOTHING",
                    (officer_id, village_id),
                )
        else:
            connection.execute("DELETE FROM officer_villages WHERE officer_id=%s", (officer_id,))
        if active == 0:
            connection.execute("DELETE FROM officer_sessions WHERE officer_id=%s", (officer_id,))
        connection.execute("""
            INSERT INTO officer_profiles (officer_id,nik,phone,address,birth_place,birth_date,profile_photo_path,ktp_path,updated_at)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,NOW())
            ON CONFLICT (officer_id) DO UPDATE SET
              nik=EXCLUDED.nik, phone=EXCLUDED.phone, address=EXCLUDED.address,
              birth_place=EXCLUDED.birth_place, birth_date=EXCLUDED.birth_date,
              profile_photo_path=CASE WHEN EXCLUDED.profile_photo_path='' THEN officer_profiles.profile_photo_path ELSE EXCLUDED.profile_photo_path END,
              ktp_path=CASE WHEN EXCLUDED.ktp_path='' THEN officer_profiles.ktp_path ELSE EXCLUDED.ktp_path END,
              updated_at=NOW()
        """, (officer_id, nik.strip(), phone.strip(), address.strip(), birth_place.strip(),
              birth_date.strip() or None, photo_path, ktp_path))
    return RedirectResponse(url=f"/admin/data-personel/{officer_id}?saved=1", status_code=303)


@app.post("/admin/data-personel/{officer_id}/hapus")
def admin_delete_data_personel(officer_id: int, username: str = Depends(verify_admin)):
    """Hapus akun/personel dan pendaftaran terkait hanya jika belum memiliki aktivitas survei/verifikasi."""
    with get_db() as connection:
        officer = connection.execute(
            "SELECT id, member_id, member_name FROM officers WHERE id=%s FOR UPDATE",
            (officer_id,),
        ).fetchone()
        if not officer:
            raise HTTPException(status_code=404, detail="Akun personel tidak ditemukan.")

        member_id = officer["member_id"]
        survey_count = connection.execute(
            "SELECT COUNT(*) AS n FROM surveys WHERE member_id=%s", (member_id,)
        ).fetchone()["n"]
        media_count = connection.execute(
            "SELECT COUNT(*) AS n FROM survey_media WHERE member_id=%s", (member_id,)
        ).fetchone()["n"]
        verification_count = connection.execute(
            "SELECT COUNT(*) AS n FROM survey_verifications WHERE verifier_id=%s", (member_id,)
        ).fetchone()["n"]
        verification_media_count = connection.execute(
            "SELECT COUNT(*) AS n FROM verification_media WHERE verifier_id=%s", (member_id,)
        ).fetchone()["n"]
        if survey_count or media_count or verification_count or verification_media_count:
            raise HTTPException(
                status_code=409,
                detail=(
                    "Data tidak dapat dihapus karena akun sudah memiliki aktivitas: "
                    f"{survey_count} survei, {media_count} media survei, "
                    f"{verification_count} verifikasi, dan {verification_media_count} media verifikasi. "
                    "Nonaktifkan akun jika akses perlu dihentikan."
                ),
            )

        # Hanya pendaftaran yang tertaut ke ID akun ini yang ikut dibersihkan.
        connection.execute(
            "DELETE FROM organization_member_registrations WHERE linked_member_id=%s OR member_code=%s",
            (member_id, member_id),
        )
        # FK ON DELETE CASCADE membersihkan sesi, profil, dan penugasan desa milik officer.
        connection.execute("DELETE FROM officers WHERE id=%s", (officer_id,))
        connection.execute("DELETE FROM members WHERE member_id=%s", (member_id,))

    return RedirectResponse(url="/admin/data-personel?deleted=1", status_code=303)


# =========================================================
# PENDAFTARAN MANDIRI ANGGOTA ORGANISASI 555
# =========================================================

@app.get("/pendaftaran-anggota", response_class=HTMLResponse)
def public_member_registration_form():
    with get_db() as connection:
        villages = connection.execute(
            "SELECT id, village_name, district FROM villages WHERE active=1 ORDER BY village_name"
        ).fetchall()
    village_options = "".join(
        f"<option value='{int(v['id'])}'>{_lead_e(v['village_name'])} · {_lead_e(v['district'])}</option>"
        for v in villages
    )
    html = f"""<!doctype html><html lang='id'><head><meta charset='utf-8'>
    <meta name='viewport' content='width=device-width,initial-scale=1'><title>Pendataan Anggota 555</title>
    <style>
    :root{{--maroon:#650b18;--deep:#17070b;--gold:#e0ad35;--paper:#f4f0eb;--ink:#f7eee5;--line:#503337;--muted:#c1b2aa}}
*{{box-sizing:border-box}}body{{margin:0;background:radial-gradient(ellipse at top right,#35151b 0,#16090c 58%,#100709 100%);color:var(--ink);font:15px/1.55 Inter,system-ui,Arial,sans-serif}}
header{{background:linear-gradient(110deg,#19080c,#500813 55%,#26090e);color:#fff;padding:20px 18px;border-bottom:2px solid var(--gold);box-shadow:0 8px 30px #0006}}header .brand{{max-width:1120px;margin:auto;display:flex;align-items:center;gap:16px}}header .brand img{{width:66px;height:66px;object-fit:contain;border-radius:50%;background:#2b1114;padding:3px;border:2px solid var(--gold)}}header .brand h1{{margin:0;color:#f7d77f;font-size:clamp(21px,3vw,29px)}}header .brand p{{margin:3px 0 0;color:#e8d8c7}}
main{{max-width:1120px;margin:24px auto;padding:0 16px 45px}}.panel{{background:linear-gradient(145deg,#26191b,#1c1114);border:1px solid #5b3939;border-radius:15px;padding:22px;margin-bottom:15px;box-shadow:0 10px 26px #0004}}h2{{font-size:17px;margin:0 0 13px;color:#f1cb72}}.grid{{display:grid;grid-template-columns:1fr 1fr;gap:12px}}label{{display:block;font-weight:700;margin:10px 0 5px;color:#e8cc8b}}input,select,textarea{{width:100%;padding:11px;border:1px solid #624449;border-radius:8px;font:inherit;background:#180e11;color:#f8eee3}}input[type=file]{{padding:8px}}.full{{grid-column:1/-1}}.hint{{font-size:13px;color:#b8a8a0;margin:5px 0}}button{{background:linear-gradient(110deg,#f0c85e,#bd821b);color:#24130a;border:0;border-radius:9px;padding:13px 20px;font:inherit;font-weight:800;cursor:pointer;box-shadow:0 5px 16px #0004}}.notice{{background:#3b2a17;border:1px solid #8b682c;border-left:4px solid var(--gold);color:#f5e4bf;padding:13px;border-radius:8px;font-size:13px}}
@media(max-width:620px){{header .brand img{{width:56px;height:56px}}.grid{{grid-template-columns:1fr}}.full{{grid-column:auto}}main{{padding:0 12px 32px;margin:14px auto}}.panel{{padding:15px}}}}
</style></head><body><header><div class='brand'><img src='/logo.png' alt='Logo Ikatan Persaudaraan Triple 555'><div><h1>Pendataan Mandiri Anggota</h1><p>Ikatan Persaudaraan Triple 555</p></div></div></header>
    <main><div class='panel'><div class='notice'><b>Petunjuk:</b> Formulir ini diisi oleh masing-masing anggota. Data akan diperiksa admin sebelum ditetapkan sebagai data resmi organisasi. Pastikan NIK dan nomor HP benar. Foto KTP hanya dapat diakses admin yang berwenang.</div>
    <form method='post' action='/pendaftaran-anggota' enctype='multipart/form-data'>
    <h2 style='margin-top:18px'>1. Identitas anggota</h2><div class='grid'>
    <div class='full'><label>Nama lengkap sesuai KTP *</label><input name='full_name' required maxlength='150' autocomplete='name'></div>
    <div><label>NIK *</label><input name='nik' required minlength='16' maxlength='16' inputmode='numeric' pattern='[0-9]{{16}}' title='NIK harus 16 digit'></div>
    <div><label>Jenis kelamin *</label><select name='gender' required><option value=''>Pilih</option><option value='Laki-laki'>Laki-laki</option><option value='Perempuan'>Perempuan</option></select></div>
    <div><label>Tempat lahir</label><input name='birth_place' maxlength='100'></div><div><label>Tanggal lahir</label><input type='date' name='birth_date'></div>
    <div class='full'><label>Alamat lengkap *</label><textarea name='address' rows='3' required maxlength='1500'></textarea></div>
    <div><label>Nomor HP / WhatsApp *</label><input name='phone' required maxlength='30' inputmode='tel' autocomplete='tel'></div>
    <div><label>Email (opsional)</label><input type='email' name='email' maxlength='150' autocomplete='email'></div></div></div>
    <div class='panel'><h2>2. Kedudukan dalam organisasi</h2><div class='grid'>
    <div><label>Tingkat kepengurusan *</label><select name='org_level' id='org_level' required><option value=''>Pilih</option><option value='DPS'>DPS · Kecamatan</option><option value='DPSS'>DPSS · Desa/Kelurahan</option></select></div>
    <div><label>Jabatan organisasi sesuai SK *</label><input name='org_position' required maxlength='150' placeholder='Isi jabatan sesuai SK pengangkatan'></div>
    <div><label>Wilayah / Kecamatan *</label><input name='region' required maxlength='150' value='Tambang'></div>
    <div><label id='village_label'>Cakupan wilayah desa</label><select name='village_id' id='village_id'><option value='all'>Seluruh Desa (cakupan DPS)</option>{village_options}</select><p class='hint' id='village_hint'>Untuk DPS, pilih Seluruh Desa atau pilih desa tertentu bila penugasan hanya mencakup desa tersebut.</p></div>
    </div><p class='hint'>Isi jabatan persis seperti yang tercantum dalam SK. Admin akan memeriksa kesesuaian tingkat dan wilayah.</p></div>
    <div class='panel'><h2>3. Foto dan dokumen</h2>
    <label>Foto profil *</label><input type='file' name='profile_photo' accept='image/jpeg,image/png,image/webp' required><p class='hint'>JPG, PNG, atau WEBP. Maksimal 5 MB.</p>
    <label>Foto / scan KTP *</label><input type='file' name='ktp_photo' accept='image/jpeg,image/png,image/webp' required><p class='hint'>Pastikan tulisan pada KTP terbaca. File disimpan di penyimpanan privat.</p>
    <label>Dokumen SK (opsional, PDF/JPG/PNG belum diaktifkan pada tahap ini)</label><p class='hint'>Untuk sementara, dokumen SK dapat diperiksa admin secara terpisah. Formulir ini hanya menerima foto profil dan KTP.</p>
    </div><div class='panel'><div class='notice'>Dengan mengirim formulir ini, anggota menyatakan data yang diisi benar dan menyetujui penggunaannya untuk administrasi internal organisasi 555.</div><p style='margin:16px 0 0'><button type='submit'>Kirim Data Anggota</button></p></div>
    </form><script>
    (function(){{
      const level=document.getElementById('org_level');
      const village=document.getElementById('village_id');
      const label=document.getElementById('village_label');
      const hint=document.getElementById('village_hint');
      function syncVillage(){{
        if(level.value==='DPSS'){{
          label.textContent='Desa / Kelurahan DPSS *';
          hint.textContent='Untuk DPSS, wajib pilih desa tempat kepengurusan.';
          village.required=true;
          if(village.value==='all') village.value='';
          village.options[0].textContent='Pilih desa';
          village.options[0].value='';
        }}else{{
          label.textContent='Cakupan wilayah desa';
          hint.textContent='Untuk DPS, pilih Seluruh Desa atau pilih desa tertentu bila penugasan hanya mencakup desa tersebut.';
          village.required=false;
          village.options[0].textContent='Seluruh Desa (cakupan DPS)';
          village.options[0].value='all';
          if(!village.value) village.value='all';
        }}
      }}
      level.addEventListener('change',syncVillage);
      syncVillage();
    }})();
    </script></main></body></html>"""
    return HTMLResponse(html)


@app.post("/pendaftaran-anggota", response_class=HTMLResponse)
def submit_public_member_registration(
    full_name: str = Form(...), nik: str = Form(...), gender: str = Form(...),
    birth_place: str = Form(""), birth_date: str = Form(""), address: str = Form(...),
    phone: str = Form(...), email: str = Form(""), org_level: str = Form(...),
    org_position: str = Form(...), region: str = Form(...), village_id: str = Form("all"),
    profile_photo: UploadFile | None = File(None),
    ktp_photo: UploadFile | None = File(None),
):
    full_name, nik, gender = full_name.strip(), re.sub(r"\D", "", nik), gender.strip()
    address, phone, email = address.strip(), phone.strip(), email.strip()
    org_level, org_position, region = org_level.strip().upper(), org_position.strip(), region.strip()
    village_value = (village_id or "").strip().lower()
    if village_value in {"", "all", "semua"}:
        village_id = None
    else:
        try:
            village_id = int(village_value)
        except (TypeError, ValueError) as exc:
            raise HTTPException(status_code=422, detail="Pilihan desa tidak valid.") from exc
    if len(full_name) < 3 or len(nik) != 16 or gender not in {"Laki-laki", "Perempuan"}:
        raise HTTPException(status_code=422, detail="Periksa nama, NIK 16 digit, dan jenis kelamin.")
    if not address or len(phone) < 8 or org_level not in {"DPS", "DPSS"} or len(org_position) < 2 or not region:
        raise HTTPException(status_code=422, detail="Lengkapi alamat, nomor HP, tingkat kepengurusan, jabatan, dan wilayah.")
    if birth_date.strip():
        try:
            datetime.strptime(birth_date.strip(), "%Y-%m-%d")
        except ValueError as exc:
            raise HTTPException(status_code=422, detail="Tanggal lahir tidak valid.") from exc
    if not profile_photo or not profile_photo.filename or not ktp_photo or not ktp_photo.filename:
        raise HTTPException(status_code=422, detail="Foto profil dan foto KTP wajib diunggah.")
    code = "REG-" + uuid.uuid4().hex[:16].upper()
    with get_db() as connection:
        duplicate = connection.execute(
            "SELECT id,status FROM organization_member_registrations WHERE nik=%s ORDER BY id DESC LIMIT 1", (nik,)
        ).fetchone()
        if duplicate:
            raise HTTPException(status_code=409, detail="NIK ini sudah pernah didaftarkan. Hubungi admin untuk memperbaiki atau memeriksa pendaftaran sebelumnya.")
        if village_id is not None:
            village = connection.execute("SELECT id FROM villages WHERE id=%s AND active=1", (village_id,)).fetchone()
            if not village:
                raise HTTPException(status_code=422, detail="Desa yang dipilih tidak tersedia.")
        if org_level == "DPSS" and village_id is None:
            raise HTTPException(status_code=422, detail="Pilih desa untuk kepengurusan DPSS.")
    photo_path = _upload_officer_private(profile_photo, code, "profile")
    ktp_path = _upload_officer_private(ktp_photo, code, "ktp")
    with get_db() as connection:
        village_name = ""
        if village_id is not None:
            village_row = connection.execute("SELECT village_name FROM villages WHERE id=%s", (village_id,)).fetchone()
            village_name = village_row["village_name"] if village_row else ""
        # Nomor dialokasikan dan baris pendaftaran disimpan dalam transaksi yang sama.
        member_code = _allocate_member_code(connection, org_level, region, village_name)
        connection.execute("""INSERT INTO organization_member_registrations
            (registration_code,member_code,full_name,nik,birth_place,birth_date,gender,address,phone,email,
             org_level,org_position,region,village_id,profile_photo_path,ktp_path)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
            (code,member_code,full_name,nik,birth_place.strip(),birth_date.strip() or None,gender,address,phone,email,
             org_level,org_position,region,village_id,photo_path,ktp_path))
    return HTMLResponse(f"""<!doctype html><html lang='id'><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>
    <title>Pendaftaran diterima</title><body style='font:16px system-ui;background:#f4f0eb;padding:24px'><main style='max-width:600px;margin:8vh auto;background:white;padding:28px;border-radius:14px;border-top:6px solid #d8b15a'><h2 style='color:#5b1018'>Data berhasil dikirim</h2><p>Terima kasih. Pendaftaran anggota telah diterima dan menunggu pemeriksaan admin.</p><p><b>ID Member: {escape(member_code)}</b></p><p><b>Kode pendaftaran: {escape(code)}</b></p><p>Simpan ID Member dan kode pendaftaran ini untuk keperluan administrasi.</p><p><a href='/pendaftaran-anggota' style='display:inline-block;background:#5b1018;color:#fff;text-decoration:none;padding:11px 16px;border-radius:8px'>Kembali ke Formulir</a></p></main></body></html>""")


@app.get("/admin/pendaftaran-anggota", response_class=HTMLResponse)
def admin_member_registration_list(status: str = "pending", q: str = "", username: str = Depends(verify_admin)):
    allowed = {"pending", "revision", "approved", "rejected", "all"}
    if status not in allowed:
        status = "pending"
    with get_db() as connection:
        rows = connection.execute("""SELECT r.*,v.village_name FROM organization_member_registrations r
            LEFT JOIN villages v ON v.id=r.village_id
            WHERE (%s='all' OR r.status=%s) AND (%s='' OR r.full_name ILIKE %s OR r.nik ILIKE %s OR r.phone ILIKE %s OR r.registration_code ILIKE %s)
            ORDER BY r.created_at DESC LIMIT 300""",
            (status,status,q.strip(),f"%{q.strip()}%",f"%{q.strip()}%",f"%{q.strip()}%",f"%{q.strip()}%")).fetchall()
    cards=[]
    for r in rows:
        photo=create_supabase_signed_url(r['profile_photo_path'],1800,_officer_private_bucket()) if r['profile_photo_path'] else None
        ktp=create_supabase_signed_url(r['ktp_path'],1800,_officer_private_bucket()) if r['ktp_path'] else None
        photo_html=f"<img src='{_lead_e(photo)}' style='width:100px;height:100px;object-fit:cover;border-radius:10px'>" if photo else ''
        ktp_html=f"<a target='_blank' rel='noopener' href='{_lead_e(ktp)}'>Lihat KTP</a>" if ktp else 'Tidak ada'
        cards.append(f"""<article class='card'><div class='head'>{photo_html}<div><b>{_lead_e(r['full_name'])}</b><div>{_lead_e(r['registration_code'])} · ID Member: {_lead_e(r.get('member_code') or '-')} · {_lead_e(r['status'].upper())}</div><div>{_lead_e(r['org_level'])} · {_lead_e(r['org_position'])}</div></div></div>
        <p><b>NIK:</b> {_lead_e(r['nik'])} · <b>HP:</b> {_lead_e(r['phone'])}</p><p><b>Wilayah:</b> {_lead_e(r['region'])} · {_lead_e(('Seluruh Desa' if r['org_level']=='DPS' and not r['village_id'] else (r['village_name'] or '-')))}</p>
        <p><b>Alamat:</b> {_lead_e(r['address'])}</p><p><b>Dokumen:</b> {ktp_html}</p>
        <p style='display:flex;gap:8px;flex-wrap:wrap'><a href='/admin/pendaftaran-anggota/{int(r['id'])}/edit' style='display:inline-block;background:#8a5b16;color:white;padding:9px 12px;border-radius:7px;text-decoration:none'>Edit Data</a></p>
        <form method='post' action='/admin/pendaftaran-anggota/{int(r['id'])}/hapus' onsubmit="return confirm('Hapus permanen pendaftaran ini? Akun personel dan riwayat survei tidak akan dihapus.')"><button type='submit' style='background:#9b2226'>Hapus Pendaftaran</button></form>
        <form method='post' action='/admin/pendaftaran-anggota/{int(r['id'])}'><label>Status pemeriksaan</label><select name='new_status'><option value='pending' {'selected' if r['status']=='pending' else ''}>Menunggu</option><option value='revision' {'selected' if r['status']=='revision' else ''}>Perlu perbaikan</option><option value='approved' {'selected' if r['status']=='approved' else ''}>Disetujui</option><option value='rejected' {'selected' if r['status']=='rejected' else ''}>Ditolak</option></select>
        <label>Catatan admin</label><textarea name='admin_note' rows='2'>{_lead_e(r['admin_note'])}</textarea><p><b>ID Member otomatis:</b> {_lead_e(r.get('member_code') or '-')}<br><small>ID login aplikasi akan menggunakan ID Member ini setelah admin memberikan akses.</small></p><button type='submit'>Simpan Pemeriksaan</button></form></article>""")
    content=f"""<!doctype html><html lang='id'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>Verifikasi Pendaftaran Anggota</title><style>
    body{{font:14px system-ui;background:#f4f0eb;margin:0;color:#292323}}header{{background:#31090e;color:white;padding:16px 20px;border-bottom:4px solid #d8b15a}}header .brand{{max-width:1100px;margin:auto;display:flex;align-items:center;gap:14px}}header .brand img{{width:62px;height:62px;object-fit:contain;border-radius:50%;background:#fff;padding:3px}}header h1{{margin:0 0 4px}}main{{max-width:1100px;margin:auto;padding:18px}}.card{{background:white;padding:17px;border-radius:12px;margin:12px 0;border:1px solid #e5ddd4}}.head{{display:flex;gap:14px;align-items:center}}label{{display:block;font-weight:bold;margin:10px 0 5px}}input,select,textarea{{width:100%;padding:10px;border:1px solid #d5cec7;border-radius:8px;box-sizing:border-box;font:inherit}}button{{background:#5b1018;color:white;border:0;padding:11px 15px;border-radius:8px;margin-top:12px}}.filters{{display:flex;gap:8px;flex-wrap:wrap}}.filters>*{{padding:9px;border-radius:8px}}a{{color:#5b1018}}small{{color:#777}}
    </style></head><body><header><div class='brand'><img src='/logo.png' alt='Logo Ikatan Persaudaraan Triple 555'><div><h1>Verifikasi Pendaftaran Anggota 555</h1><div>Admin organisasi · dokumen identitas hanya ditampilkan di halaman ini</div></div></div></header><main>
    <form class='filters' method='get'><select name='status'><option value='pending' {'selected' if status=='pending' else ''}>Menunggu</option><option value='revision' {'selected' if status=='revision' else ''}>Perlu perbaikan</option><option value='approved' {'selected' if status=='approved' else ''}>Disetujui</option><option value='rejected' {'selected' if status=='rejected' else ''}>Ditolak</option><option value='all' {'selected' if status=='all' else ''}>Semua</option></select><input name='q' value='{_lead_e(q)}' placeholder='Cari nama, NIK, HP, kode'><button type='submit'>Tampilkan</button></form>
    <p><a href='/admin/dashboard'>Dashboard Utama Admin</a> &nbsp; | &nbsp; <a href='/admin/data-personel'>Data Personel</a></p>{''.join(cards) if cards else '<p>Belum ada pendaftaran pada kategori ini.</p>'}</main></body></html>"""
    return HTMLResponse(content)


@app.get("/admin/pendaftaran-anggota/{registration_id}/edit", response_class=HTMLResponse)
def admin_edit_member_registration_page(registration_id: int, username: str = Depends(verify_admin)):
    with get_db() as connection:
        row = connection.execute("SELECT * FROM organization_member_registrations WHERE id=%s", (registration_id,)).fetchone()
        villages = connection.execute("SELECT id,village_name FROM villages WHERE active=1 ORDER BY village_name").fetchall()
    if not row:
        raise HTTPException(status_code=404, detail="Pendaftaran tidak ditemukan.")
    village_options = "<option value=''>Seluruh Desa (cakupan DPS)</option>" + "".join(
        f"<option value='{int(v['id'])}' {'selected' if row['village_id']==v['id'] else ''}>{_lead_e(v['village_name'])}</option>" for v in villages
    )
    photo = create_supabase_signed_url(row['profile_photo_path'], 1800, _officer_private_bucket()) if row['profile_photo_path'] else None
    ktp = create_supabase_signed_url(row['ktp_path'], 1800, _officer_private_bucket()) if row['ktp_path'] else None
    html = f"""<!doctype html><html lang='id'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>Edit Pendaftaran Anggota</title>
    <style>body{{font:15px system-ui;background:#f4f0eb;margin:0;color:#292323}}header{{background:#31090e;color:white;padding:20px;border-bottom:4px solid #d8b15a}}main{{max-width:850px;margin:22px auto;padding:0 14px}}form,.box{{background:white;padding:20px;border-radius:12px;margin:12px 0}}label{{display:block;font-weight:700;margin:12px 0 5px}}input,select,textarea{{width:100%;box-sizing:border-box;padding:10px;border:1px solid #d5cec7;border-radius:8px;font:inherit}}button,a.btn{{display:inline-block;background:#5b1018;color:white;border:0;border-radius:8px;padding:11px 15px;text-decoration:none;margin-top:14px;cursor:pointer}}.grid{{display:grid;grid-template-columns:1fr 1fr;gap:12px}}img{{max-width:150px;max-height:150px;object-fit:contain}}small{{color:#666}}@media(max-width:600px){{.grid{{grid-template-columns:1fr}}}}</style></head><body><header><h1>Edit Data Pendaftaran Anggota</h1><div>{_lead_e(row['registration_code'])} · Status: {_lead_e(row['status'])}</div></header><main>
    <form method='post' enctype='multipart/form-data' action='/admin/pendaftaran-anggota/{registration_id}/edit'>
    <div class='grid'><div><label>Nama lengkap *</label><input name='full_name' required maxlength='150' value='{_lead_e(row['full_name'])}'></div><div><label>NIK *</label><input name='nik' required minlength='16' maxlength='16' pattern='[0-9]{{16}}' value='{_lead_e(row['nik'])}'></div>
    <div><label>Jenis kelamin *</label><select name='gender' required><option value='Laki-laki' {'selected' if row['gender']=='Laki-laki' else ''}>Laki-laki</option><option value='Perempuan' {'selected' if row['gender']=='Perempuan' else ''}>Perempuan</option></select></div><div><label>Tempat lahir</label><input name='birth_place' maxlength='100' value='{_lead_e(row['birth_place'])}'></div>
    <div><label>Tanggal lahir</label><input type='date' name='birth_date' value='{_lead_e(row['birth_date'].isoformat() if row['birth_date'] else '')}'></div><div><label>Nomor HP *</label><input name='phone' required maxlength='30' value='{_lead_e(row['phone'])}'></div>
    <div><label>Email</label><input type='email' name='email' maxlength='150' value='{_lead_e(row['email'])}'></div><div><label>Tingkat organisasi *</label><select name='org_level' id='org_level'><option value='DPS' {'selected' if row['org_level']=='DPS' else ''}>DPS</option><option value='DPSS' {'selected' if row['org_level']=='DPSS' else ''}>DPSS</option></select></div>
    <div><label>Jabatan *</label><input name='org_position' required maxlength='150' value='{_lead_e(row['org_position'])}'></div><div><label>Wilayah / Kecamatan *</label><input name='region' required maxlength='150' value='{_lead_e(row['region'])}'></div>
    <div><label>Desa / Cakupan DPS</label><select name='village_id' id='village_id'>{village_options}</select><small>Untuk DPS, kosongkan jika mencakup seluruh desa. Untuk DPSS, wajib pilih desa.</small></div></div>
    <label>Alamat lengkap *</label><textarea name='address' required rows='3' maxlength='1500'>{_lead_e(row['address'])}</textarea>
    <div class='grid'><div><label>Foto profil saat ini</label>{f"<img src='{_lead_e(photo)}'>" if photo else 'Belum ada'}<label>Ganti foto profil (opsional)</label><input type='file' name='profile_photo' accept='image/jpeg,image/png,image/webp'></div>
    <div><label>Foto KTP saat ini</label>{f"<img src='{_lead_e(ktp)}'>" if ktp else 'Belum ada'}<label>Ganti foto KTP (opsional)</label><input type='file' name='ktp_photo' accept='image/jpeg,image/png,image/webp'></div></div>
    <button type='submit'>Simpan Perubahan</button> <a class='btn' href='/admin/pendaftaran-anggota'>Batal</a></form></main>
    <script>const level=document.getElementById('org_level'), village=document.getElementById('village_id');function syncVillage(){{village.required=level.value==='DPSS';if(level.value==='DPSS'&&!village.value){{}}}}level.addEventListener('change',syncVillage);syncVillage();</script></body></html>"""
    return HTMLResponse(html)


@app.post("/admin/pendaftaran-anggota/{registration_id}/edit")
def admin_edit_member_registration(
    registration_id: int, full_name: str = Form(...), nik: str = Form(...), gender: str = Form(...),
    birth_place: str = Form(""), birth_date: str = Form(""), address: str = Form(...), phone: str = Form(...),
    email: str = Form(""), org_level: str = Form(...), org_position: str = Form(...), region: str = Form(...),
    village_id: str = Form(""), profile_photo: UploadFile | None = File(None), ktp_photo: UploadFile | None = File(None),
    username: str = Depends(verify_admin),
):
    full_name, nik, gender = full_name.strip(), re.sub(r"\D", "", nik), gender.strip()
    address, phone, email = address.strip(), phone.strip(), email.strip()
    org_level, org_position, region = org_level.strip().upper(), org_position.strip(), region.strip()
    if len(full_name) < 3 or len(nik) != 16 or gender not in {"Laki-laki", "Perempuan"} or not address or len(phone) < 8 or org_level not in {"DPS", "DPSS"} or len(org_position) < 2 or not region:
        raise HTTPException(status_code=422, detail="Periksa kembali seluruh data wajib.")
    if birth_date.strip():
        try: datetime.strptime(birth_date.strip(), "%Y-%m-%d")
        except ValueError as exc: raise HTTPException(status_code=422, detail="Tanggal lahir tidak valid.") from exc
    village_value = (village_id or "").strip()
    try:
        village_id_value = None if village_value in {"", "all"} else int(village_value)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail="Pilihan desa tidak valid.") from exc
    if org_level == "DPSS" and village_id_value is None:
        raise HTTPException(status_code=422, detail="DPSS wajib memilih desa.")
    with get_db() as connection:
        row = connection.execute("SELECT * FROM organization_member_registrations WHERE id=%s", (registration_id,)).fetchone()
        if not row: raise HTTPException(status_code=404, detail="Pendaftaran tidak ditemukan.")
        duplicate = connection.execute("SELECT id FROM organization_member_registrations WHERE nik=%s AND id<>%s", (nik, registration_id)).fetchone()
        if duplicate: raise HTTPException(status_code=409, detail="NIK sudah digunakan pada pendaftaran lain.")
        if village_id_value is not None:
            valid_village = connection.execute("SELECT id FROM villages WHERE id=%s AND active=1", (village_id_value,)).fetchone()
            if not valid_village: raise HTTPException(status_code=422, detail="Desa tidak tersedia.")
    new_photo = _upload_officer_private(profile_photo, row['registration_code'], "profile") if profile_photo and profile_photo.filename else row['profile_photo_path']
    new_ktp = _upload_officer_private(ktp_photo, row['registration_code'], "ktp") if ktp_photo and ktp_photo.filename else row['ktp_path']
    with get_db() as connection:
        new_member_code = row.get('member_code') or ''
        if not row.get('linked_member_id'):
            new_village_name = ''
            if village_id_value is not None:
                new_village_row = connection.execute("SELECT village_name FROM villages WHERE id=%s", (village_id_value,)).fetchone()
                new_village_name = new_village_row['village_name'] if new_village_row else ''
            old_village_name = ''
            if row.get('village_id'):
                old_village_row = connection.execute("SELECT village_name FROM villages WHERE id=%s", (row['village_id'],)).fetchone()
                old_village_name = old_village_row['village_name'] if old_village_row else ''
            if _member_code_prefix(row['org_level'], row['region'], old_village_name) != _member_code_prefix(org_level, region, new_village_name):
                new_member_code = _allocate_member_code(connection, org_level, region, new_village_name)
        connection.execute("""UPDATE organization_member_registrations SET member_code=%s,full_name=%s,nik=%s,gender=%s,birth_place=%s,birth_date=%s,address=%s,phone=%s,email=%s,org_level=%s,org_position=%s,region=%s,village_id=%s,profile_photo_path=%s,ktp_path=%s,updated_at=NOW() WHERE id=%s""",
            (new_member_code,full_name,nik,gender,birth_place.strip(),birth_date.strip() or None,address,phone,email,org_level,org_position,region,village_id_value,new_photo,new_ktp,registration_id))
        if row['linked_member_id'] and row['status']=='approved':
            officer = connection.execute("SELECT id FROM officers WHERE member_id=%s", (row['linked_member_id'],)).fetchone()
            if officer:
                connection.execute("UPDATE officers SET member_name=%s,role=%s,region=%s,updated_at=NOW() WHERE id=%s", (full_name,org_level,region,officer['id']))
                connection.execute("UPDATE members SET member_name=%s WHERE member_id=%s", (full_name,row['linked_member_id']))
                connection.execute("""UPDATE officer_profiles SET nik=%s,phone=%s,address=%s,birth_place=%s,birth_date=%s,profile_photo_path=%s,ktp_path=%s,updated_at=NOW() WHERE officer_id=%s""", (nik,phone,address,birth_place.strip(),birth_date.strip() or None,new_photo,new_ktp,officer['id']))
    return RedirectResponse(url="/admin/pendaftaran-anggota?status=all", status_code=303)


@app.post("/admin/pendaftaran-anggota/{registration_id}/hapus")
def admin_delete_member_registration(registration_id: int, username: str = Depends(verify_admin)):
    with get_db() as connection:
        row = connection.execute("SELECT id,profile_photo_path,ktp_path,linked_member_id,status FROM organization_member_registrations WHERE id=%s", (registration_id,)).fetchone()
    if not row: raise HTTPException(status_code=404, detail="Pendaftaran tidak ditemukan.")
    # Jika sudah disetujui dan ditautkan, file yang sama mungkin sedang dipakai profil personel.
    # Karena itu hapus hanya catatan pendaftarannya, jangan hapus objek dokumen personel.
    keep_personnel_files = bool(row['linked_member_id'] and row['status'] == 'approved')
    paths = [] if keep_personnel_files else [path for path in (row['profile_photo_path'], row['ktp_path']) if path]
    if paths:
        supabase_url = os.environ.get("SUPABASE_URL", "").strip().rstrip("/")
        supabase_key = os.environ.get("SUPABASE_SECRET_KEY", "").strip()
        if not supabase_url or not supabase_key:
            raise HTTPException(status_code=503, detail="Konfigurasi Supabase belum lengkap. Data tidak dihapus.")
        bucket = _officer_private_bucket()
        request = urllib.request.Request(f"{supabase_url}/storage/v1/object/{quote(bucket, safe='')}", data=json.dumps({"prefixes": paths}).encode("utf-8"), headers={"Authorization": f"Bearer {supabase_key}", "apikey": supabase_key, "Content-Type": "application/json"}, method="DELETE")
        try:
            with urllib.request.urlopen(request, timeout=45) as response:
                if response.status not in (200, 204): raise HTTPException(status_code=502, detail="Supabase tidak mengonfirmasi penghapusan dokumen. Data belum dihapus.")
        except urllib.error.HTTPError as exc:
            raise HTTPException(status_code=502, detail="Gagal menghapus dokumen privat. Data belum dihapus.") from exc
        except (urllib.error.URLError, TimeoutError) as exc:
            raise HTTPException(status_code=502, detail="Penyimpanan dokumen tidak dapat dijangkau. Data belum dihapus.") from exc
    with get_db() as connection:
        connection.execute("DELETE FROM organization_member_registrations WHERE id=%s", (registration_id,))
    return RedirectResponse(url="/admin/pendaftaran-anggota?status=all", status_code=303)


@app.post("/admin/pendaftaran-anggota/{registration_id}")
def admin_review_member_registration(
    registration_id: int, new_status: str = Form(...), admin_note: str = Form(""),
    username: str = Depends(verify_admin),
):
    if new_status not in {"pending", "revision", "approved", "rejected"}:
        raise HTTPException(status_code=422, detail="Status pemeriksaan tidak valid.")
    with get_db() as connection:
        row = connection.execute("SELECT * FROM organization_member_registrations WHERE id=%s", (registration_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Pendaftaran tidak ditemukan.")
        # Jika anggota lama sudah terhubung ke akun, sinkronkan profilnya. Pendaftaran baru
        # tidak meminta admin mengetik ID; ID login ditetapkan otomatis saat akses diberikan.
        linked_member_id = (row.get("linked_member_id") or "").strip().upper()
        if linked_member_id and new_status == "approved":
            officer = connection.execute("SELECT id FROM officers WHERE member_id=%s", (linked_member_id,)).fetchone()
            if officer:
                connection.execute("""INSERT INTO officer_profiles
                    (officer_id,nik,phone,address,birth_place,birth_date,profile_photo_path,ktp_path,updated_at)
                    VALUES (%s,%s,%s,%s,%s,%s,%s,%s,NOW()) ON CONFLICT (officer_id) DO UPDATE SET
                    nik=EXCLUDED.nik,phone=EXCLUDED.phone,address=EXCLUDED.address,birth_place=EXCLUDED.birth_place,
                    birth_date=EXCLUDED.birth_date,
                    profile_photo_path=CASE WHEN EXCLUDED.profile_photo_path='' THEN officer_profiles.profile_photo_path ELSE EXCLUDED.profile_photo_path END,
                    ktp_path=CASE WHEN EXCLUDED.ktp_path='' THEN officer_profiles.ktp_path ELSE EXCLUDED.ktp_path END,updated_at=NOW()""",
                    (officer['id'],row['nik'],row['phone'],row['address'],row['birth_place'],row['birth_date'],row['profile_photo_path'],row['ktp_path']))
                connection.execute("UPDATE officers SET member_name=%s,role=%s,region=%s,updated_at=NOW() WHERE id=%s",
                    (row['full_name'],row['org_level'],row['region'],officer['id']))
                connection.execute("UPDATE members SET member_name=%s WHERE member_id=%s", (row['full_name'],linked_member_id))
        connection.execute("""UPDATE organization_member_registrations SET status=%s,admin_note=%s,
            updated_at=NOW(),reviewed_at=NOW() WHERE id=%s""",
            (new_status,admin_note.strip(),registration_id))
    return RedirectResponse(url="/admin/pendaftaran-anggota?status="+new_status, status_code=303)


@app.get("/pimpinan/personel", response_class=HTMLResponse)
def pimpinan_data_personel(q: str = "", username: str = Depends(verify_pimpinan_or_admin)):
    with get_db() as connection:
        rows = connection.execute("""
            SELECT r.id, r.full_name, r.registration_code, r.org_level, r.org_position,
                   r.region, r.phone, r.profile_photo_path, r.linked_member_id,
                   o.active AS account_active, v.village_name
            FROM organization_member_registrations r
            LEFT JOIN officers o ON o.member_id=r.linked_member_id
            LEFT JOIN villages v ON v.id=r.village_id
            WHERE r.status='approved'
              AND (%s='' OR r.full_name ILIKE %s OR r.nik ILIKE %s
                   OR r.phone ILIKE %s OR r.org_position ILIKE %s
                   OR r.region ILIKE %s OR COALESCE(v.village_name,'') ILIKE %s)
            ORDER BY r.org_level, r.region, r.full_name
        """, (q.strip(), f"%{q.strip()}%", f"%{q.strip()}%", f"%{q.strip()}%",
              f"%{q.strip()}%", f"%{q.strip()}%", f"%{q.strip()}%")).fetchall()
    cards = []
    for r in rows:
        photo = create_supabase_signed_url(r["profile_photo_path"], expires_in=1800, bucket=_officer_private_bucket()) if r["profile_photo_path"] else None
        avatar = f"<img class='person-photo' src='{_lead_e(photo)}' alt='Foto {_lead_e(r['full_name'])}'>" if photo else "<div class='person-photo no-photo'>Foto<br>belum ada</div>"
        territory = r["village_name"] or r["region"] or "-"
        if not r["linked_member_id"]:
            access = "<span class='status inactive'>BELUM DIBERI AKSES</span>"
        elif r["account_active"] == 1:
            access = "<span class='status active'>AKSES AKTIF</span>"
        else:
            access = "<span class='status inactive'>AKSES NONAKTIF</span>"
        cards.append(f"<article class='person-card'>{avatar}<div class='person-info'><h3>{_lead_e(r['full_name'])}</h3><div class='muted'>{_lead_e(r['org_level'])} · {_lead_e(r['org_position'])}</div><p>☎ {_lead_e(r['phone'] or 'Nomor HP belum diisi')}</p><p>⌖ {_lead_e(territory)}</p>{access}</div></article>")
    if not cards:
        cards = ["<div class='panel'>Belum ada anggota dengan pendaftaran berstatus Disetujui.</div>"]
    content = f"""<section class='panel'><form method='get'><label>Cari anggota</label><input name='q' value='{_lead_e(q)}' placeholder='Nama, nomor HP, jabatan, wilayah'><button type='submit'>Cari</button> <a class='btn light' href='/pimpinan/personel'>Reset</a></form><p class='muted'>Menampilkan {len(rows)} anggota yang pendaftarannya sudah disetujui. Status akses aplikasi ditampilkan terpisah dari status keanggotaan.</p><p><a class='btn' href='/admin/kelola-akun'>Kelola Akses Aplikasi</a></p></section><section class='people-grid'>{''.join(cards)}</section>"""
    return HTMLResponse(_lead_page("Data Anggota Terverifikasi", username, content))


def _lead_map_embed(surveys, verifications, map_id="lead-dashboard-map"):
    markers = []
    for row in surveys:
        try:
            lat, lon = float(row.get("latitude")), float(row.get("longitude"))
            if not (-90 <= lat <= 90 and -180 <= lon <= 180):
                continue
        except (TypeError, ValueError):
            continue
        markers.append({"kind":"survey","lat":lat,"lon":lon,"survey_id":row.get("survey_id"),
            "applicant":row.get("applicant_name") or "-","business":row.get("business_name") or "-",
            "village":row.get("village") or "-","officer":row.get("officer_name") or row.get("member_id") or "-",
            "date":row.get("surveyed_at") or "-","accuracy":row.get("accuracy")})
    for row in verifications:
        try:
            lat, lon = float(row.get("verification_latitude")), float(row.get("verification_longitude"))
            if not (-90 <= lat <= 90 and -180 <= lon <= 180):
                continue
        except (TypeError, ValueError):
            continue
        markers.append({"kind":"verification","lat":lat,"lon":lon,"survey_id":row.get("survey_id"),
            "applicant":row.get("applicant_name") or "-","business":row.get("business_name") or "-",
            "village":row.get("village") or "-","officer":row.get("verifier_name") or row.get("verifier_id") or "-",
            "date":row.get("verified_at") or "-","accuracy":row.get("verification_accuracy"),
            "decision":row.get("decision") or "-"})
    data = json.dumps(markers, ensure_ascii=False, default=str).replace("</", "<\\/")
    survey_count = sum(1 for m in markers if m["kind"] == "survey")
    verify_count = len(markers) - survey_count
    return f"""<section class='panel'>
      <h2 style='margin-top:0;color:#5b1018'>🗺️ Peta Sebaran GPS Survei dan Verifikasi</h2>
      <p class='muted'>Peta lokasi tersimpan. Merah = survei DPSS, emas = verifikasi DPS. Pilih titik untuk melihat detail dan akurasi GPS.</p>
      <div style='display:flex;gap:18px;flex-wrap:wrap;margin:10px 0;font-weight:700'><span>🔴 Survei DPSS: {survey_count}</span><span>🟡 Verifikasi DPS: {verify_count}</span><span>📍 Total: {len(markers)}</span><a class='btn light' href='/pimpinan/peta'>Buka Peta & Filter Lengkap</a></div>
      <div id='{map_id}' style='height:440px;min-height:320px;border:1px solid #ddd;border-radius:12px;z-index:1'></div>
      <p id='{map_id}-status' class='muted' aria-live='polite' style='margin-bottom:0'>Memuat peta…</p>
    </section>
    <link rel='stylesheet' href='https://unpkg.com/leaflet@1.9.4/dist/leaflet.css' crossorigin=''/>
    <script src='https://unpkg.com/leaflet@1.9.4/dist/leaflet.js' crossorigin=''></script>
    <script type='application/json' id='{map_id}-data'>{data}</script>
    <script>(function(){{const el=document.getElementById('{map_id}'),st=document.getElementById('{map_id}-status');const data=JSON.parse(document.getElementById('{map_id}-data').textContent||'[]');if(!window.L){{st.textContent='Peta gagal dimuat. Periksa koneksi internet browser.';return;}}const map=L.map(el).setView([0.5,101.4],8);L.tileLayer('https://{{s}}.tile.openstreetmap.org/{{z}}/{{x}}/{{y}}.png',{{maxZoom:19,attribution:'&copy; OpenStreetMap contributors'}}).addTo(map);const bounds=[];const safe=v=>String(v===null||v===undefined||v===''?'-':v).replace(/[&<>\"']/g,c=>({{'&':'&amp;','<':'&lt;','>':'&gt;','\"':'&quot;',"'":'&#39;'}}[c]));data.forEach(x=>{{const survey=x.kind==='survey';const m=L.circleMarker([x.lat,x.lon],{{radius:8,color:survey?'#8f1825':'#9a6b12',weight:2,fillColor:survey?'#d94a55':'#e5b94f',fillOpacity:.9}}).addTo(map);const acc=x.accuracy===null||x.accuracy===undefined?'-':safe(x.accuracy)+' m';m.bindPopup(`<b>${{survey?'SURVEI DPSS':'VERIFIKASI DPS'}}</b><br>ID: ${{safe(x.survey_id)}}<br>Pemohon: ${{safe(x.applicant)}}<br>Usaha: ${{safe(x.business)}}<br>Desa: ${{safe(x.village)}}<br>Petugas: ${{safe(x.officer)}}<br>Waktu: ${{safe(x.date)}}<br>Koordinat: ${{x.lat.toFixed(6)}}, ${{x.lon.toFixed(6)}}<br>Akurasi: ${{acc}}${{x.decision?'<br>Keputusan: '+safe(x.decision):''}}`);bounds.push([x.lat,x.lon]);}});if(bounds.length){{map.fitBounds(bounds,{{padding:[28,28],maxZoom:15}});st.textContent=data.length+' titik GPS ditampilkan. Klik penanda untuk detail.';}}else{{st.textContent='Belum ada data koordinat GPS yang dapat dipetakan.';}}setTimeout(()=>map.invalidateSize(),180);}})();</script>"""

@app.get("/pimpinan", response_class=HTMLResponse)
def pimpinan_dashboard(request: Request, username: str = Depends(verify_pimpinan_or_admin)):
    filters = _lead_filters(request)
    surveys = _lead_survey_rows(filters)
    verifications = _lead_verification_rows(filters)
    pending = max(0, len(surveys) - len(verifications))
    content = f"""<section class='metrics'><div class='metric'><span>Total survei DPSS</span><strong>{len(surveys)}</strong></div><div class='metric'><span>Sudah diverifikasi DPS</span><strong>{len(verifications)}</strong></div><div class='metric'><span>Menunggu verifikasi</span><strong>{pending}</strong></div><div class='metric'><span>Wilayah dalam hasil filter</span><strong>{len(set(str(r.get('village') or '') for r in surveys if r.get('village')))}</strong></div></section>
{_lead_map_embed(surveys, verifications)}
<section class='cards'><article class='report-card'><h2>📋 Laporan Survei DPSS</h2><p>Data awal pemohon dan usaha yang dikirim petugas desa. Foto usaha dan video survei ditampilkan pada halaman ini. Data identitas pribadi tidak dimasukkan ke laporan ekspor.</p><a class='btn' href='/pimpinan/survei'>Buka Laporan Survei</a></article><article class='report-card'><h2>🔎 Laporan Verifikasi DPS</h2><p>Hasil pemeriksaan DPS, keputusan, catatan verifikasi, rincian kebutuhan, serta media verifikasi. Laporan ini terpisah dari data survei awal.</p><a class='btn' href='/pimpinan/verifikasi'>Buka Laporan Verifikasi</a></article></section>
<section class='cards'><article class='report-card'><h2>👥 Data Personel DPS dan DPSS</h2><p>Daftar anggota terverifikasi dan status penugasan akses aplikasi. Pemberian akses hanya untuk anggota yang ditunjuk dan disetujui pimpinan/admin.</p><a class='btn' href='/pimpinan/personel'>Lihat Data Personel</a></article><article class='report-card'><h2>🗂️ Pendataan Identitas Petugas</h2><p>Identitas anggota bersumber dari pendaftaran mandiri yang telah diverifikasi. Pengaturan akun hanya untuk memberikan atau mencabut akses aplikasi.</p><a class='btn' href='/admin/kelola-akun'>Kelola Akses Aplikasi</a></article></section>
<section class='panel'><h2>Catatan akses</h2><p class='muted'>Dashboard ini hanya menampilkan rekap. Fungsi input dan pengiriman dari aplikasi HP tidak diubah. Ekspor laporan survei dan verifikasi tersedia masing-masing pada halaman laporan.</p></section>"""
    return HTMLResponse(_lead_page("Dashboard Pimpinan", username, content))



@app.get("/pimpinan/peta", response_class=HTMLResponse)
def pimpinan_peta_gps(request: Request, username: str = Depends(verify_pimpinan_or_admin)):
    filters = _lead_filters(request)
    surveys = _lead_survey_rows(filters)
    verifications = _lead_verification_rows(filters)
    markers = []
    villages = sorted({str(r.get("village") or "").strip() for r in surveys if str(r.get("village") or "").strip()}, key=str.casefold)
    for row in surveys:
        try:
            lat, lon = float(row.get("latitude")), float(row.get("longitude"))
            if not (-90 <= lat <= 90 and -180 <= lon <= 180):
                continue
        except (TypeError, ValueError):
            continue
        markers.append({
            "kind": "survey", "lat": lat, "lon": lon,
            "survey_id": row.get("survey_id"), "applicant": row.get("applicant_name") or "-",
            "business": row.get("business_name") or "-", "village": row.get("village") or "-",
            "officer": row.get("officer_name") or row.get("member_id") or "-",
            "date": row.get("surveyed_at") or "-", "accuracy": row.get("accuracy"),
        })
    for row in verifications:
        try:
            lat, lon = float(row.get("verification_latitude")), float(row.get("verification_longitude"))
            if not (-90 <= lat <= 90 and -180 <= lon <= 180):
                continue
        except (TypeError, ValueError):
            continue
        markers.append({
            "kind": "verification", "lat": lat, "lon": lon,
            "survey_id": row.get("survey_id"), "applicant": row.get("applicant_name") or "-",
            "business": row.get("business_name") or "-", "village": row.get("village") or "-",
            "officer": row.get("verifier_name") or row.get("verifier_id") or "-",
            "date": row.get("verified_at") or "-", "accuracy": row.get("verification_accuracy"),
            "decision": row.get("decision") or "-",
        })
    marker_json = json.dumps(markers, ensure_ascii=False, default=str).replace("</", "<\\/")
    village_options = "".join(f"<option value='{_lead_e(v)}' {'selected' if filters.get('wilayah') == v else ''}>{_lead_e(v)}</option>" for v in villages)
    content = f"""<section class='panel'>
      <h2 style='margin-top:0;color:#5b1018'>Peta Sebaran Lokasi Survei dan Verifikasi</h2>
      <p class='muted'>Titik merah menunjukkan GPS survei DPSS, sedangkan titik emas menunjukkan GPS verifikasi DPS. Peta menggunakan koordinat yang tersimpan di database, bukan lokasi perangkat yang sedang membuka dashboard.</p>
      <form method='get' action='/pimpinan/peta' class='filters'>
        <div><label>TANGGAL MULAI</label><input type='date' name='tanggal_mulai' value='{_lead_e(filters.get('tanggal_mulai'))}'></div>
        <div><label>TANGGAL SELESAI</label><input type='date' name='tanggal_selesai' value='{_lead_e(filters.get('tanggal_selesai'))}'></div>
        <div><label>DESA</label><select name='wilayah' style='width:100%;min-height:40px;padding:8px;border:1px solid #d7d2ca;border-radius:8px'><option value=''>Semua desa</option>{village_options}</select></div>
        <div><label>PETUGAS DPSS</label><input name='petugas' value='{_lead_e(filters.get('petugas'))}' placeholder='ID atau nama petugas'></div>
        <div class='actions' style='grid-column:1/-1'><button type='submit'>Terapkan Filter</button><a class='btn light' href='/pimpinan/peta'>Reset</a></div>
      </form>
      <div class='map-legend' style='display:flex;gap:18px;flex-wrap:wrap;margin:14px 0;font-weight:700'><span>🔴 GPS Survei DPSS: {sum(1 for m in markers if m['kind']=='survey')}</span><span>🟡 GPS Verifikasi DPS: {sum(1 for m in markers if m['kind']=='verification')}</span><span>📍 Total titik: {len(markers)}</span></div>
      <div id='lead-gps-map' style='height:560px;min-height:360px;border:1px solid #ddd;border-radius:12px;z-index:1'></div>
      <p id='lead-map-status' class='muted' aria-live='polite' style='margin-bottom:0'>Memuat peta…</p>
    </section>
    <script type='application/json' id='lead-map-data'>{marker_json}</script>
    <link rel='stylesheet' href='https://unpkg.com/leaflet@1.9.4/dist/leaflet.css' crossorigin=''/>
    <script src='https://unpkg.com/leaflet@1.9.4/dist/leaflet.js' crossorigin=''></script>
    <script>
    (function(){{
      const el=document.getElementById('lead-gps-map'), status=document.getElementById('lead-map-status');
      const data=JSON.parse(document.getElementById('lead-map-data').textContent||'[]');
      if(!window.L){{status.textContent='Pustaka peta tidak berhasil dimuat. Periksa koneksi internet browser.';return;}}
      const map=L.map(el).setView([0.5,101.4],8);
      L.tileLayer('https://{{s}}.tile.openstreetmap.org/{{z}}/{{x}}/{{y}}.png',{{maxZoom:19,attribution:'&copy; OpenStreetMap contributors'}}).addTo(map);
      const bounds=[];
      data.forEach(item=>{{
        const isSurvey=item.kind==='survey';
        const marker=L.circleMarker([item.lat,item.lon],{{radius:8,color:isSurvey?'#8f1825':'#9a6b12',weight:2,fillColor:isSurvey?'#d94a55':'#e5b94f',fillOpacity:.9}}).addTo(map);
        const safe=v=>String(v===null||v===undefined||v===''?'-':v).replace(/[&<>\"']/g,c=>({{'&':'&amp;','<':'&lt;','>':'&gt;','\"':'&quot;',"'":'&#39;'}}[c]));
        const acc=item.accuracy===null||item.accuracy===undefined?'-':safe(item.accuracy)+' m';
        marker.bindPopup(`<b>${{isSurvey?'SURVEI DPSS':'VERIFIKASI DPS'}}</b><br>ID Survei: ${{safe(item.survey_id)}}<br>Pemohon: ${{safe(item.applicant)}}<br>Usaha: ${{safe(item.business)}}<br>Desa: ${{safe(item.village)}}<br>Petugas: ${{safe(item.officer)}}<br>Waktu: ${{safe(item.date)}}<br>Koordinat: ${{item.lat.toFixed(6)}}, ${{item.lon.toFixed(6)}}<br>Akurasi GPS: ${{acc}}${{item.decision?'<br>Keputusan: '+safe(item.decision):''}}`);
        bounds.push([item.lat,item.lon]);
      }});
      if(bounds.length){{map.fitBounds(bounds,{{padding:[30,30],maxZoom:15}});status.textContent=`${{data.length}} titik GPS ditampilkan. Pilih penanda untuk melihat detail.`;}}
      else{{status.textContent='Belum ada koordinat GPS yang cocok dengan filter ini.';}}
      setTimeout(()=>map.invalidateSize(),150);
    }})();
    </script>"""
    return HTMLResponse(_lead_page("Peta Lokasi GPS", username, content))

def _lead_filters_form(filters, base, export_prefix):
    qs = _lead_qs(filters)
    exp = "?" + qs if qs else ""
    return f"""<section class='panel'><form method='get' action='{base}'><div class='filters'>
    <div><label>TANGGAL MULAI</label><input type='date' name='tanggal_mulai' value='{_lead_e(filters.get('tanggal_mulai'))}'></div>
    <div><label>TANGGAL SELESAI</label><input type='date' name='tanggal_selesai' value='{_lead_e(filters.get('tanggal_selesai'))}'></div>
    <div><label>WILAYAH / DESA</label><input name='wilayah' value='{_lead_e(filters.get('wilayah'))}' placeholder='Cari wilayah'></div>
    <div><label>PETUGAS DPSS (ID/NAMA)</label><input name='petugas' value='{_lead_e(filters.get('petugas'))}' placeholder='Cari petugas'></div></div>
    <div class='actions'><button type='submit'>Terapkan Filter</button><a class='btn light' href='{base}'>Reset</a><a class='btn gold' href='{export_prefix}/excel{exp}'>Ekspor Excel</a><a class='btn' target='_blank' href='{export_prefix}/pdf{exp}'>Ekspor PDF</a></div></form></section>"""


def _lead_media_html(media_rows):
    identity_bits = []
    photo_bits = []
    video_bits = []
    other_bits = []

    category_labels = {
        "applicant_ktp": "KTP Pemohon",
        "family_card": "Kartu Keluarga",
        "recipient_photo": "Foto Penerima",
        "business_photo": "Foto Usaha",
        "business_video": "Video Usaha",
    }

    for media in media_rows:
        url = media.get("signed_url")
        name = _lead_e(media.get("original_filename") or "Media")
        category = str(
            media.get("media_category") or "business_photo"
        ).lower()
        ctype = str(media.get("content_type") or "").lower()

        label = _lead_e(
            category_labels.get(
                category,
                category.replace("_", " ").title()
            )
        )

        if not url:
            other_bits.append(
                f"<div class='lead-media-unavailable'>"
                f"<small>{label}</small>"
                f"<p>{name} · Pratinjau tidak tersedia</p>"
                f"</div>"
            )
            continue

        safe_url = _lead_e(url)

        if ctype.startswith("image/"):
            # Dokumen identitas ditampilkan utuh, tidak dipotong.
            is_identity = category in (
                "applicant_ktp",
                "family_card",
            )

            css_class = (
                "lead-identity-item"
                if is_identity
                else "lead-photo-item"
            )

            image_class = (
                "lead-identity-image"
                if is_identity
                else "lead-photo-image"
            )

            html = (
                f"<div class='{css_class}'>"
                f"<div class='lead-media-label'>{label}</div>"
                f"<a href='{safe_url}' target='_blank' "
                f"rel='noopener noreferrer' "
                f"aria-label='Buka {label} ukuran penuh'>"
                f"<img class='{image_class}' "
                f"src='{safe_url}' alt='{name}' loading='lazy'>"
                f"</a>"
                f"<div class='lead-media-name'>{name}</div>"
                f"</div>"
            )

            if is_identity:
                identity_bits.append(html)
            else:
                photo_bits.append(html)

        elif ctype.startswith("video/"):
            video_bits.append(
                f"<div class='lead-video-item'>"
                f"<div class='lead-media-label'>{label}</div>"
                f"<video controls preload='metadata' "
                f"playsinline src='{safe_url}'></video>"
                f"<div class='lead-media-name'>{name}</div>"
                f"</div>"
            )

    sections = []

    if identity_bits:
        sections.append(
            "<section class='lead-media-section'>"
            "<h4 class='lead-media-heading'>"
            "Dokumen Identitas</h4>"
            "<div class='lead-identity-grid'>"
            + "".join(identity_bits) +
            "</div></section>"
        )

    if photo_bits:
        sections.append(
            "<section class='lead-media-section'>"
            "<h4 class='lead-media-heading'>"
            "Galeri Foto Usaha</h4>"
            "<div class='lead-photo-grid'>"
            + "".join(photo_bits) +
            "</div></section>"
        )

    if video_bits:
        sections.append(
            "<section class='lead-media-section'>"
            "<h4 class='lead-media-heading'>"
            "Video Survei</h4>"
            "<div class='lead-video-grid'>"
            + "".join(video_bits) +
            "</div></section>"
        )

    if other_bits:
        sections.append(
            "<section class='lead-media-section'>"
            "<div class='lead-media-unavailable-list'>"
            + "".join(other_bits) +
            "</div></section>"
        )

    if not sections:
        return (
            "<div class='lead-media-empty'>"
            "Belum ada foto atau video dokumentasi.</div>"
        )

    return "<div class='lead-media-gallery'>" + "".join(sections) + "</div>"


def _lead_report_html(rows, media_map, verification=False):
    body = []
    for r in rows:
        media = _lead_media_html(media_map.get(r["survey_id"], []))
        if not verification:
            detail = "<br>".join(_lead_e(x) for x in [r.get("business_type"), r.get("business_location"), "Catatan: " + str(r.get("notes") or "")] if x)
            body.append(f"<tr class='lead-report-data-row'><td>{r['survey_id']}</td><td>{_lead_e(r.get('surveyed_at'))}</td><td>{_lead_e(r.get('officer_name'))}<br><small>{_lead_e(r.get('member_id'))}</small></td><td>{_lead_e(r.get('village'))}</td><td><b>{_lead_e(r.get('applicant_name'))}</b><br>{_lead_e(r.get('applicant_phone'))}<br>{_lead_e(r.get('applicant_address'))}</td><td><b>{_lead_e(r.get('business_name'))}</b><br>{detail}</td><td>{_lead_e(r.get('capital_amount'))}<br><small>Omzet: {_lead_e(r.get('turnover_amount'))}</small></td></tr><tr class='lead-report-media-row'><td colspan='7'>{media}</td></tr>")
        else:
            needs = r.get("needs") or []
            if isinstance(needs, str):
                try: needs = json.loads(needs)
                except ValueError: needs = []
            needs_text = "<br>".join(_lead_e(f"{n.get('item_name','Item')}: {n.get('quantity',0)} {n.get('unit','unit')} × Rp {n.get('unit_price',0)}") for n in needs if isinstance(n, dict)) or "-"
            body.append(f"<tr class='lead-report-data-row'><td>{r['survey_id']}</td><td>{_lead_e(r.get('verified_at'))}</td><td>{_lead_e(r.get('verifier_name'))}<br><small>{_lead_e(r.get('verifier_id'))}</small></td><td>{_lead_e(r.get('verifier_region') or r.get('village'))}</td><td><b>{_lead_e(r.get('applicant_name'))}</b><br>{_lead_e(r.get('business_name'))}</td><td><b>{_lead_e(r.get('decision'))}</b><br>Kondisi: {_lead_e(r.get('business_condition'))}<br>Aktivitas: {_lead_e(r.get('actual_business_activity'))}<br>Kesesuaian: {_lead_e(r.get('data_match'))}<br>Catatan: {_lead_e(r.get('verification_notes'))}<br>Alasan: {_lead_e(r.get('rejection_reason'))}</td><td>{needs_text}<br><b>Total: Rp {_lead_e(r.get('total_needs_amount'))}</b></td></tr><tr class='lead-report-media-row'><td colspan='7'>{media}</td></tr>")
    if not body:
        body = ["<tr><td colspan='7' style='text-align:center'>Tidak ada data yang sesuai filter.</td></tr>"]
    if verification:
        heading = "Laporan Verifikasi DPS"
        headers = "<th>ID Survei</th><th>Waktu Verifikasi</th><th>Verifikator</th><th>Wilayah DPS</th><th>Pemohon / Usaha</th><th>Hasil Pemeriksaan</th><th>Rincian Kebutuhan</th>"
    else:
        heading = "Laporan Survei DPSS"
        headers = "<th>ID</th><th>Waktu Survei</th><th>Petugas DPSS</th><th>Desa</th><th>Pemohon</th><th>Informasi Usaha</th><th>Modal / Omzet Awal</th>"
    return f"<section class='panel lead-report-panel'><h2>{heading} · {len(rows)} data</h2><div class='tablewrap lead-report-tablewrap'><table class='lead-report-table'><thead><tr>{headers}</tr></thead><tbody>{''.join(body)}</tbody></table></div></section>"


@app.get("/pimpinan/survei", response_class=HTMLResponse)
def pimpinan_laporan_survei(request: Request, username: str = Depends(verify_pimpinan_or_admin)):
    filters = _lead_filters(request)
    rows = _lead_survey_rows(filters)
    # Halaman dashboard survei menampilkan KTP bersama foto/video survei.
    # Endpoint ekspor tidak memanggil include_identity=True, sehingga KTP tetap privat.
    media = _lead_media([r["survey_id"] for r in rows], verification=False, include_identity=True)
    content = _lead_filters_form(filters, "/pimpinan/survei", "/pimpinan/survei/export") + _lead_report_html(rows, media, False)
    return HTMLResponse(_lead_page("Laporan Survei DPSS", username, content))


@app.get("/pimpinan/verifikasi", response_class=HTMLResponse)
def pimpinan_laporan_verifikasi(request: Request, username: str = Depends(verify_pimpinan_or_admin)):
    filters = _lead_filters(request)
    rows = _lead_verification_rows(filters)
    media = _lead_media([r["survey_id"] for r in rows], verification=True)
    content = _lead_filters_form(filters, "/pimpinan/verifikasi", "/pimpinan/verifikasi/export") + _lead_report_html(rows, media, True)
    return HTMLResponse(_lead_page("Laporan Verifikasi DPS", username, content))


def _lead_export_excel(rows, verification, filename, sheet_title):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    if verification:
        columns = [("ID Survei","survey_id"),("Waktu Verifikasi","verified_at"),("ID Petugas DPSS","member_id"),("Nama Petugas DPSS","officer_name"),("Desa","village"),("Nama Pemohon","applicant_name"),("Nama Usaha","business_name"),("Keputusan DPS","decision"),("ID Verifikator","verifier_id"),("Nama Verifikator","verifier_name"),("Wilayah DPS","verifier_region"),("Kondisi Usaha","business_condition"),("Aktivitas Aktual","actual_business_activity"),("Aset","existing_assets"),("Rincian Kebutuhan","needs"),("Total Kebutuhan","total_needs_amount"),("Omzet Terverifikasi","verified_turnover"),("Laba Terverifikasi","verified_profit"),("Jumlah Pekerja","verified_worker_count"),("Kesesuaian Data","data_match"),("Perbedaan","discrepancy_notes"),("Catatan Verifikasi","verification_notes"),("Alasan Penolakan","rejection_reason")]
    else:
        columns = [("ID Survei","survey_id"),("Waktu Survei","surveyed_at"),("ID Petugas DPSS","member_id"),("Nama Petugas DPSS","officer_name"),("Desa","village"),("Nama Pemohon","applicant_name"),("Telepon","applicant_phone"),("Alamat","applicant_address"),("Nama Usaha","business_name"),("Jenis Usaha","business_type"),("Lokasi Usaha","business_location"),("Modal Awal","capital_amount"),("Omzet Awal","turnover_amount"),("Laba Awal","net_profit_amount"),("Jumlah Pekerja","worker_count"),("Jenis Bantuan","assistance_types"),("Nilai Bantuan","assistance_amount"),("Catatan Survei","notes"),("Latitude","latitude"),("Longitude","longitude")]
    wb=Workbook(); ws=wb.active; ws.title=sheet_title; ws.append([c[0] for c in columns])
    for cell in ws[1]:
        cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="5B1018"); cell.alignment=Alignment(wrap_text=True,vertical="top")
    for r in rows:
        values=[]
        for _, key in columns:
            value=r.get(key)
            if key == "needs" and not isinstance(value, str): value=json.dumps(value or [], ensure_ascii=False)
            values.append(value)
        ws.append(values)
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    for i,(label,_) in enumerate(columns,1): ws.column_dimensions[get_column_letter(i)].width=min(max(len(label)+3,14),34)
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.alignment=Alignment(wrap_text=True,vertical="top")
    out=BytesIO(); wb.save(out); out.seek(0)
    return StreamingResponse(out,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":f"attachment; filename={filename}"})


def _lead_export_pdf(rows, verification, filename, title):
    from xml.sax.saxutils import escape as xml_escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    if verification:
        keys=[("ID","survey_id"),("Waktu verifikasi","verified_at"),("Pemohon / usaha","applicant_name"),("Keputusan","decision"),("Verifikator","verifier_name"),("Kondisi","business_condition"),("Catatan","verification_notes"),("Total kebutuhan","total_needs_amount")]
    else:
        keys=[("ID","survey_id"),("Waktu survei","surveyed_at"),("Petugas","officer_name"),("Desa","village"),("Pemohon","applicant_name"),("Usaha","business_name"),("Modal awal","capital_amount"),("Omzet awal","turnover_amount")]
    out=BytesIO(); doc=SimpleDocTemplate(out,pagesize=landscape(A4),rightMargin=20,leftMargin=20,topMargin=22,bottomMargin=22,title=title)
    styles=getSampleStyleSheet(); styles.add(ParagraphStyle(name="LeadCell",parent=styles["BodyText"],fontSize=7,leading=8)); styles.add(ParagraphStyle(name="LeadHead",parent=styles["BodyText"],fontName="Helvetica-Bold",fontSize=7,textColor=colors.white,leading=8))
    story=[Paragraph(xml_escape(title),styles["Title"]),Spacer(1,8),Paragraph(f"Jumlah data: {len(rows)}. Foto dan video tidak disematkan di PDF; buka dashboard untuk melihat media.",styles["Normal"]),Spacer(1,10)]
    data=[[Paragraph(xml_escape(label),styles["LeadHead"]) for label,_ in keys]]
    for r in rows:
        vals=[]
        for _,key in keys:
            value=r.get(key)
            if key=="applicant_name": value=(str(value or "")+"\n"+str(r.get("applicant_phone") or "")).strip()
            if key=="business_name": value=(str(value or "")+"\n"+str(r.get("business_type") or "")).strip()
            vals.append(Paragraph(xml_escape("" if value is None else str(value)).replace("\n","<br/>"),styles["LeadCell"]))
        data.append(vals)
    widths=[30,75,90,65,100,100,65,65] if not verification else [30,75,100,75,90,100,150,75]
    table=Table(data,colWidths=widths,repeatRows=1,hAlign="LEFT")
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#5B1018")),("GRID",(0,0),(-1,-1),0.35,colors.HexColor("#C9C3BB")),("VALIGN",(0,0),(-1,-1),"TOP"),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F8F4EC")]),("LEFTPADDING",(0,0),(-1,-1),4),("RIGHTPADDING",(0,0),(-1,-1),4)]))
    story.append(table); doc.build(story); out.seek(0)
    return StreamingResponse(out,media_type="application/pdf",headers={"Content-Disposition":f"attachment; filename={filename}"})


@app.get("/pimpinan/survei/export/excel")
def pimpinan_survei_excel(request: Request, username: str = Depends(verify_pimpinan_or_admin)):
    return _lead_export_excel(_lead_survey_rows(_lead_filters(request)),False,"laporan_survei_dpss.xlsx","Survei DPSS")

@app.get("/pimpinan/survei/export/pdf")
def pimpinan_survei_pdf(request: Request, username: str = Depends(verify_pimpinan_or_admin)):
    return _lead_export_pdf(_lead_survey_rows(_lead_filters(request)),False,"laporan_survei_dpss.pdf","Laporan Survei DPSS")

@app.get("/pimpinan/verifikasi/export/excel")
def pimpinan_verifikasi_excel(request: Request, username: str = Depends(verify_pimpinan_or_admin)):
    return _lead_export_excel(_lead_verification_rows(_lead_filters(request)),True,"laporan_verifikasi_dps.xlsx","Verifikasi DPS")

@app.get("/pimpinan/verifikasi/export/pdf")
def pimpinan_verifikasi_pdf(request: Request, username: str = Depends(verify_pimpinan_or_admin)):
    return _lead_export_pdf(_lead_verification_rows(_lead_filters(request)),True,"laporan_verifikasi_dps.pdf","Laporan Verifikasi DPS")

# =========================================================
# MENJALANKAN SERVER SECARA LOKAL
# =========================================================

if __name__ == "__main__":
    uvicorn.run(
        app,
        host="0.0.0.0",
        port=int(os.environ.get("PORT", "8000")),
    )
